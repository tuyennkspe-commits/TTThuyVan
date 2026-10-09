"""Mưa thiết kế và úng nội đồng — một file độc lập.
Cài thư viện: python -m pip install numpy scipy matplotlib pandas openpyxl python-docx
Căn cứ: hồ sơ Hà Nội–Gia Bình Part1, trang in 14–17, 29–30, 37–38.
Htk/Hp là CAO ĐỘ mặt nước (m), Xp là lượng mưa (mm). Chỉ phục vụ tính úng nội đồng.
Không có dữ liệu trạm hoặc hệ số khí hậu mặc định; dùng dữ liệu của dự án.
"""
import csv
import io
import json
import math
import calendar
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import numpy as np
import pandas as pd
from scipy.stats import pearson3, gumbel_r, skew, norm
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk

GROWTH_HEADER='Phut;P_pct;X_toan_mm;X_coso_mm;Xi_mm;Tang_KB1_pct;Tang_KB2_pct'
POINT_HEADER='Vi_tri;Hi_m;Nam_dieu_tra;Z_dat_m'
COLUMN_TITLES={'duration':'Thời đoạn (phút)','p':'Tần suất vượt P (%)','scenario':'Kịch bản','Xp':'Mưa thiết kế Xp (mm)','Xi':'Mưa năm điều tra Xi (mm)','beta':'Hệ số β','deltaH':'Chênh mực nước ΔH (m)','location':'Vị trí','Hi':'Cao độ điều tra Hi (m)','Hp':'Mực nước thiết kế Htk (m)','days':'Thời đoạn khống chế (ngày)','depth':'Chiều sâu ngập (m)','X_full':'Mưa toàn chuỗi (mm)','X_base':'Mưa cơ sở (mm)','change_percent':'Mức thay đổi (%)','ground':'Cao độ đất (m)','year':'Năm điều tra','full':'Mưa toàn chuỗi (mm)','base':'Mưa cơ sở (mm)','xi':'Mưa năm điều tra (mm)','mid':'Thay đổi KB1 (%)','end':'Thay đổi KB2 (%)'}
SCENARIOS=('Hiện trạng','Kịch bản 1','Kịch bản 2','Bao lớn nhất hiện trạng và 2 kịch bản')

def num(value,name='Giá trị'):
    try:value=float(str(value).strip().replace(',','.'))
    except (ValueError,TypeError):raise ValueError(f'{name}: phải là số.') from None
    if not math.isfinite(value):raise ValueError(f'{name}: phải hữu hạn.')
    return value

def change_rate(base,future):
    base=num(base,'Mưa thời kỳ cơ sở');future=num(future,'Mưa kịch bản')
    if base<=0 or future<0:raise ValueError('Mưa cơ sở phải dương; mưa kịch bản không âm.')
    return (future/base-1)*100

def table(text):
    lines=[line for line in text.strip().splitlines() if line.strip()]
    if not lines:raise ValueError('Bảng dữ liệu đang trống.')
    delimiter=';' if ';' in lines[0] else '\t' if '\t' in lines[0] else ','
    rows=list(csv.reader(lines,delimiter=delimiter))
    if len(rows)<2:raise ValueError('Bảng cần dòng tiêu đề và ít nhất một dòng số liệu.')
    size=len(rows[0])
    for i,row in enumerate(rows[1:],2):
        if len(row)!=size:raise ValueError(f'Dòng {i}: cần {size} cột, nhận được {len(row)}.')
    return [[v.strip() for v in row] for row in rows]

def frequency(values,probabilities,method='Pearson III',cs=None):
    x=np.asarray(values,dtype=float);p=np.asarray(probabilities,dtype=float)
    if len(x)<4 or not np.all(np.isfinite(x)) or np.any(x<0) or x.mean()<=0:raise ValueError('Mỗi chuỗi cần ≥4 năm mưa hữu hạn, không âm, trung bình dương.')
    if np.any(~np.isfinite(p)) or np.any((p<=0)|(p>=100)):raise ValueError('P phải trong (0;100) %.')
    mean=float(x.mean());sd=float(x.std(ddof=1));cv=sd/mean
    sample_cs=float(skew(x,bias=False)) if sd>0 else 0.0
    selected_cs=sample_cs if cs is None else num(cs,'Cs nhập trực tiếp')
    if sd==0:xp=np.full(len(p),mean)
    elif method=='Gumbel':
        scale=sd*np.sqrt(6)/np.pi;loc=mean-np.euler_gamma*scale
        xp=gumbel_r.isf(p/100,loc=loc,scale=scale)
    else:xp=mean+sd*pearson3.isf(p/100,skew=selected_cs)
    if np.any(~np.isfinite(xp)) or np.any(xp<0):raise ValueError('Phân vị mưa âm/không hữu hạn: cần xem lại phương pháp hoặc tham số; không tự đổi thành 0.')
    return xp,dict(n=len(x),mean=mean,cv=cv,cs_sample=sample_cs,cs_used=selected_cs,method=method)

def parse_annual(text):
    rows=table(text);header=rows[0]
    if len(header)<2:raise ValueError('Mẫu: Nam;1440;4320;7200;10080 (thời đoạn phút).')
    durations=[num(v,'Thời đoạn phút') for v in header[1:]]
    if any(d<=0 for d in durations) or len(set(durations))!=len(durations):raise ValueError('Thời đoạn phải dương và không trùng.')
    years=[];matrix=[]
    for line,row in enumerate(rows[1:],2):
        year=num(row[0],f'Năm dòng {line}')
        if year!=int(year) or not 1800<=year<=2300 or year in years:raise ValueError(f'Dòng {line}: năm không hợp lệ hoặc trùng.')
        values=[num(v,f'Mưa dòng {line}') for v in row[1:]]
        if any(v<0 for v in values):raise ValueError(f'Dòng {line}: mưa không được âm.')
        order=np.argsort(durations)
        if any(np.diff(np.array(values)[order])<-1e-8):raise ValueError(f'Dòng {line}: tổng mưa cực đại phải không giảm theo thời đoạn.')
        years.append(int(year));matrix.append(values)
    return np.array(years),np.array(durations),np.array(matrix)

def read_rain_workbook(path,sheet_name=None):
    """Đọc bảng Nam/phút hoặc biểu Rx nhiều tầng; không đoán cột ngày/tháng là mưa."""
    import re
    import unicodedata
    import openpyxl
    workbook=openpyxl.load_workbook(path,data_only=True)
    def normalized(value):
        text=unicodedata.normalize('NFKD',str(value or '').lower())
        return ''.join(c for c in text if not unicodedata.combining(c)).replace('đ','d')
    candidates=[]
    for sheet in workbook:
        if sheet_name and sheet.title!=sheet_name:continue
        for row in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row,60)):
            first=normalized(row[0].value).strip()
            # Plain annual format: Nam;1440;4320... with no multi-level headings.
            if first in ('nam','year'):
                durations=[];columns=[]
                for cell in row[1:]:
                    if cell.value is None:continue
                    try:value=num(cell.value,'Thời đoạn')
                    except ValueError:durations=[];break
                    durations.append(value);columns.append(cell.column)
                if durations and all(v>0 for v in durations):
                    candidates.append((sheet,row[0].row,1,columns,durations));break
            # Duration row: 1 ngày / 3 ngµy / 5 ngày / 7 ngày, optionally hours/minutes.
            durations=[];columns=[]
            for cell in row:
                text=normalized(cell.value)
                match=re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*(ng\S*|day\S*|gio|hour\S*|phut|minute\S*|min)\s*',text)
                if match:
                    value=num(match.group(1),'Thời đoạn');unit=match.group(2)
                    factor=1440 if unit.startswith(('ng','day')) else 60 if unit.startswith(('gio','hour')) else 1
                    durations.append(value*factor);columns.append(cell.column)
            if len(durations)<2:continue
            # Require a recognizable year header below duration groups.
            year_col=None
            for below in sheet.iter_rows(min_row=row[0].row,max_row=min(row[0].row+5,sheet.max_row)):
                for cell in below:
                    text=normalized(cell.value).strip()
                    if text in ('nam','year') or re.fullmatch(r'n[^a-z0-9]*m',text):year_col=cell.column;break
                if year_col:break
            if year_col:
                candidates.append((sheet,row[0].row,year_col,columns,durations));break
    if not candidates:raise ValueError('Không nhận diện được bảng mưa. Cần tiêu đề năm và nhóm 1/3/5/7 ngày hoặc bảng Nam;thời đoạn phút.')
    if len(candidates)>1:raise ValueError('Có nhiều bảng mưa phù hợp: '+', '.join(item[0].title for item in candidates)+'. Hãy chọn rõ trang tính.')
    sheet,header,year_col,columns,durations=candidates[0]
    records=[];started=False
    for row_index in range(header+1,sheet.max_row+1):
        raw_year=sheet.cell(row_index,year_col).value
        if raw_year is None:continue
        try:year=num(raw_year,'Năm')
        except ValueError:
            text=normalized(raw_year).strip()
            if not started or text in ('max','min','tb','trung binh','average','mean','tong'):continue
            raise ValueError(f'{sheet.title}!{sheet.cell(row_index,year_col).coordinate}: năm không hợp lệ.')
        if not 1800<=year<=2300 or year!=int(year):
            if not started:continue
            raise ValueError(f'{sheet.title}, dòng {row_index}: năm không hợp lệ.')
        started=True;values=[]
        for column in columns:
            cell=sheet.cell(row_index,column)
            if cell.value is None:raise ValueError(f'{sheet.title}!{cell.coordinate}: thiếu mưa hoặc công thức chưa có giá trị lưu; cần mở và lưu Excel rồi nhập lại.')
            values.append(num(cell.value,f'{sheet.title}!{cell.coordinate}'))
        records.append([int(year)]+values)
    if not records:raise ValueError('Nhận diện được tiêu đề nhưng không có dòng mưa theo năm.')
    text='Nam;'+';'.join(f'{d:g}' for d in durations)+'\n'+'\n'.join(';'.join(str(v) if i==0 else f'{v:.12g}' for i,v in enumerate(row)) for row in records)
    years,_,_=parse_annual(text)  # Validate duplicates, nonfinite values, and duration consistency.
    station=''
    for row in sheet.iter_rows(min_row=1,max_row=min(header,10)):
        for cell in row:
            raw=str(cell.value or '')
            if normalized(raw).startswith('tram:'):station=raw.split(':',1)[1].strip()
    return text,dict(sheet=sheet.title,station=station,year_start=int(years.min()),year_end=int(years.max()),count=len(years),durations=durations)

def daily_to_annual(text):
    rows=table(text);dates=[];rain=[]
    for i,row in enumerate(rows[1:],2):
        if len(row)!=2:raise ValueError('Mưa ngày cần 2 cột Ngay;Mua_mm, ngày YYYY-MM-DD.')
        try:date=pd.Timestamp(row[0])
        except Exception:raise ValueError(f'Dòng {i}: ngày không hợp lệ.') from None
        if date!=date.normalize():raise ValueError('Dữ liệu phải là lượng mưa ngày, không có giờ.')
        value=num(row[1],f'Mưa dòng {i}')
        if value<0:raise ValueError('Mưa ngày không được âm.')
        dates.append(date);rain.append(value)
    series=pd.Series(rain,index=pd.DatetimeIndex(dates)).sort_index()
    if series.index.has_duplicates:raise ValueError('Có ngày trùng; cần kiểm tra nguồn dữ liệu.')
    records=[];omitted=[]
    # Chỉ dùng năm đầy đủ; rolling trong từng năm, không gán cửa sổ qua năm cho năm khác.
    for year,part in series.groupby(series.index.year):
        expected=pd.date_range(f'{year}-01-01',f'{year}-12-31')
        if not part.index.equals(expected):omitted.append(int(year));continue
        records.append([int(year)]+[float(part.rolling(days,min_periods=days).sum().max()) for days in (1,3,5,7)])
    if not records:raise ValueError('Không có năm đủ tất cả ngày; không điền mưa thiếu bằng 0.')
    text='Nam;1440;4320;7200;10080\n'+'\n'.join(';'.join(str(v) for v in row) for row in records)
    return text,omitted

def parse_design(text):
    rows=table(text)
    if len(rows[0])!=7:raise ValueError('Bảng mưa cần đúng 7 cột theo mẫu.')
    output=[];seen=set()
    for i,row in enumerate(rows[1:],2):
        d,p,x,base,xi,mid,end=row
        d=num(d,'Thời đoạn');p=num(p,'P%');x=num(x,'Mưa toàn chuỗi');xi=None if not xi else num(xi,'Mưa năm điều tra')
        if d<=0 or not 0<p<100 or x<0 or xi is not None and xi<0:raise ValueError(f'Dòng {i}: thời đoạn/P/mưa không hợp lệ.')
        if (d,p) in seen:raise ValueError(f'Dòng {i}: thời đoạn và tần suất bị trùng.')
        seen.add((d,p))
        base=None if not base else num(base,'Mưa cơ sở')
        mid=None if not mid else num(mid,'Tăng giữa thế kỷ');end=None if not end else num(end,'Tăng cuối thế kỷ')
        if base is not None and base<0 or any(v is not None and v<=-100 for v in (mid,end)):raise ValueError('Mưa cơ sở không âm; mức thay đổi phải > −100%.')
        output.append(dict(duration=d,p=p,full=x,base=base,xi=xi,mid=mid,end=end))
    for p in {r['p'] for r in output}:
        subset=sorted([r for r in output if r['p']==p],key=lambda r:r['duration'])
        if any(b['full']<a['full']-1e-8 for a,b in zip(subset[:-1],subset[1:])):raise ValueError('Mưa toàn chuỗi giảm khi thời đoạn tăng; cần kiểm tra số liệu, không tự sửa.')
    # Xi phải đồng nhất tại cùng thời đoạn; kiểm tra thứ tự mưa theo P.
    for d in {r['duration'] for r in output}:
        subset=sorted([r for r in output if r['duration']==d],key=lambda r:r['p'])
        if len({r['xi'] for r in subset})>1:raise ValueError('Xi tại cùng thời đoạn phải giống nhau giữa các tần suất.')
        for a,b in zip(subset[:-1],subset[1:]):
            if b['full']>a['full']+1e-8:raise ValueError('Mưa toàn chuỗi tăng khi P tăng; kiểm tra bảng.')
    return output

def parse_locations(text,reference_year=None):
    rows=table(text);output=[];seen=set()
    if len(rows[0])!=4:raise ValueError('Vị trí cần 4 cột theo mẫu.')
    for row in rows[1:]:
        label,h,year,ground=row
        year=num(year,'Năm điều tra')
        if year!=int(year) or not 1800<=year<=2300:raise ValueError('Năm điều tra phải nguyên và hợp lệ.')
        if not label or (label,year) in seen:raise ValueError('Tên vị trí trống hoặc vị trí/năm bị trùng.')
        seen.add((label,year))
        if reference_year is not None and year!=reference_year:raise ValueError(f'{label}: năm Hi={year:g} khác năm Xi={reference_year}; phải ghép mưa và mực nước cùng năm.')
        output.append(dict(location=label,Hi=num(h,'Cao độ điều tra'),year=int(year),ground=None if not ground else num(ground,'Cao độ đất')))
    return output

def calculate_inland(design,locations,b1,b2,c,ratio,scenario,minimum_days=1):
    if not 0<=c<=1 or not 0<=ratio<=1 or b1<0 or b2<0:raise ValueError('C và An/A phải trong [0;1]; β1, β2 không âm.')
    beta=b1+b2+c*ratio;details=[];summary=[];notes=[]
    for row in design:
        if row['duration']<minimum_days*1440:continue
        candidates=[('Hiện trạng',row['full'])]
        for name,key in [('Kịch bản 1','mid'),('Kịch bản 2','end')]:
            if row['base'] is not None and row[key] is not None:candidates.append((name,row['base']*(1+row[key]/100)))
        if scenario==SCENARIOS[3]:selected=candidates
        else:selected=[pair for pair in candidates if pair[0]==scenario]
        if not selected:
            notes.append(f"Bỏ thời đoạn {row['duration']/1440:g} ngày, P={row['p']:g}%: thiếu mưa cơ sở hoặc mức thay đổi cho {scenario}.");continue
        if row['xi'] is None:raise ValueError(f"Thiếu Xi tại thời đoạn {row['duration']/1440:g} ngày. Nhập mưa cùng năm với Hi trước khi tính úng.")
        for name,xp in selected:
            xp=num(xp,'Mưa thiết kế sau thay đổi');num((1+beta)*(xp-row['xi'])/1000,'Chênh lệch mực nước')
            details.append(dict(duration=row['duration'],p=row['p'],scenario=name,Xp=xp,X_full=row['full'],X_base=row['base'],change_percent=0 if name=='Hiện trạng' else row['mid'] if name=='Kịch bản 1' else row['end'],Xi=row['xi'],beta=beta,deltaH=(1+beta)*(xp-row['xi'])/1000))
    if not details:raise ValueError('Không có thời đoạn đủ dữ liệu để tính úng cho kịch bản đã chọn.')
    for p in sorted({r['p'] for r in details}):
        governing=max([r for r in details if r['p']==p],key=lambda r:r['deltaH'])
        for place in locations:
            hp=num(place['Hi']+governing['deltaH'],'Cao độ mực nước thiết kế')
            summary.append(dict(location=place['location'],p=p,Hi=place['Hi'],deltaH=governing['deltaH'],Hp=hp,days=governing['duration']/1440,scenario=governing['scenario'],ground=place['ground'],depth=None if place['ground'] is None else max(hp-place['ground'],0)))
    return details,summary,notes

OBS_HEADER='Nam;Phut;Xi_mm'

def parse_observed(text):
    if len(text.strip().splitlines())<2:return {}
    rows=table(text);mapping={}
    if len(rows[0])!=3:raise ValueError('Mưa điều tra cần 3 cột: Năm, thời đoạn phút, Xi mm.')
    for row in rows[1:]:
        year,duration,xi=[num(v,'Mưa năm điều tra') for v in row]
        if year!=int(year) or not 1800<=year<=2300 or duration<1440 or xi<0:raise ValueError('Năm phải nguyên; thời đoạn ≥1 ngày; Xi không âm.')
        key=(int(year),duration)
        if key in mapping:raise ValueError('Trùng năm và thời đoạn trong bảng Xi.')
        mapping[key]=xi
    return mapping

def calculate_multi(design,locations,observed,b1,b2,c,ratio,scenario,reference_year=None):
    details=[];per_year=[];notes=[]
    for place in locations:
        matched=[]
        for row in design:
            if row['duration']<1440:continue
            key=(place['year'],row['duration'])
            xi=observed.get(key)
            if xi is None and reference_year==place['year']:xi=row['xi']
            matched.append(dict(row,xi=xi))
        d,s,n=calculate_inland(matched,[place],b1,b2,c,ratio,scenario)
        details.extend(dict(row,location=place['location'],year=place['year'],Hi=place['Hi'],Hp=place['Hi']+row['deltaH']) for row in d)
        if {r['p'] for r in s}!={r['p'] for r in matched}:raise ValueError('Một tần suất đã chọn chưa có thời đoạn đủ dữ liệu cho kịch bản này.')
        per_year.extend(dict(row,year=place['year']) for row in s);notes.extend(n)
    if not per_year:raise ValueError('Không có kết quả mực nước điều tra.')
    summary=[]
    for key in sorted({(r['location'],r['p']) for r in per_year}):
        summary.append(max([r for r in per_year if (r['location'],r['p'])==key],key=lambda r:r['Hp']))
    if len({r['year'] for r in per_year})>1:notes.append('Kết quả khống chế lấy Htk lớn nhất trong các năm điều tra đã nhập; cần đánh giá chất lượng và tính đại diện của từng năm trước khi dùng thiết kế.')
    return details,per_year,summary,list(dict.fromkeys(notes))

class DataGrid(ttk.Frame):
    """Bảng biên tập: giữ chuỗi gốc, chỉ làm tròn lớp hiển thị."""
    def __init__(self,parent):
        super().__init__(parent);self.pack(fill='both',expand=True,pady=6)
        self.header=[];self.rows=[];self.editor=None
        bar=ttk.Frame(self);bar.grid(row=0,column=0,sticky='ew')
        for title,command in [('Thêm hàng',self.add_row),('Xóa hàng chọn',self.remove_rows),('Dán từ Excel',self.paste),('Sao chép bảng',self.copy),('Thêm thời đoạn',self.add_column)]:
            button=ttk.Button(bar,text=title,command=command);button.pack(side='left',padx=3,pady=3)
            if title=='Thêm thời đoạn':self.column_button=button
        ttk.Label(bar,text='Nhấp đúp để sửa ô · Enter: lưu · Esc: hủy').pack(side='right',padx=8)
        self.tree=ttk.Treeview(self,show='headings',selectmode='extended');self.tree.grid(row=1,column=0,sticky='nsew')
        self.tree.tag_configure('even',background='#f0f7fa');self.tree.tag_configure('odd',background='white')
        y=ttk.Scrollbar(self,command=self.tree.yview);y.grid(row=1,column=1,sticky='ns');x=ttk.Scrollbar(self,orient='horizontal',command=self.tree.xview);x.grid(row=2,column=0,sticky='ew')
        self.tree.configure(yscrollcommand=y.set,xscrollcommand=x.set);self.rowconfigure(1,weight=1);self.columnconfigure(0,weight=1)
        self.tree.bind('<Double-1>',self.edit);self.tree.bind('<Control-v>',lambda e:self.paste());self.tree.bind('<Control-c>',lambda e:self.copy())
    def title(self,key):
        names={'Nam':'Năm','Vi_tri':'Vị trí','Hi_m':'Cao độ điều tra Hi (m)','Nam_dieu_tra':'Năm điều tra','Z_dat_m':'Cao độ đất (m)','Phut':'Thời đoạn (phút)','P_pct':'Tần suất P (%)','X_toan_mm':'Mưa toàn chuỗi (mm)','X_coso_mm':'Mưa cơ sở (mm)','Xi_mm':'Mưa năm điều tra Xi (mm)','Tang_KB1_pct':'Thay đổi KB1 (%)','Tang_KB2_pct':'Thay đổi KB2 (%)','Tang_giua_pct':'Thay đổi KB1 (%)','Tang_cuoi_pct':'Thay đổi KB2 (%)'}
        if key in names:return names[key]
        try:
            minutes=float(key);return f'Mưa {minutes/1440:g} ngày (mm)' if minutes>=1440 and minutes%1440==0 else f'Mưa {minutes:g} phút (mm)'
        except ValueError:return key
    def display(self,value,column):
        if value=='':return ''
        if self.header[column] in ('Nam','Nam_dieu_tra','Vi_tri'):return value
        try:return f'{num(value):.3f}'
        except ValueError:return value
    def refresh(self):
        if self.header and self.header[0]=='Nam':self.column_button.pack(side='left',padx=3,pady=3)
        else:self.column_button.pack_forget()
        self.tree.delete(*self.tree.get_children());columns=[str(i) for i in range(len(self.header))];self.tree.configure(columns=columns)
        for i,key in enumerate(self.header):
            self.tree.heading(str(i),text=self.title(key));self.tree.column(str(i),width=max(125,min(190,len(self.title(key))*7)),anchor='center')
        for i,row in enumerate(self.rows):self.tree.insert('','end',iid=str(i),values=[self.display(v,j) for j,v in enumerate(row)],tags=('even' if i%2==0 else 'odd',))
    def commit(self):
        if self.editor:
            widget,row,column=self.editor;value=widget.get().strip();self.editor=None;widget.destroy();self.rows[row][column]=value;self.refresh()
    def edit(self,event):
        self.commit();item=self.tree.identify_row(event.y);column=self.tree.identify_column(event.x)
        if not item or not column:return
        j=int(column[1:])-1;box=self.tree.bbox(item,column)
        if not box:return
        widget=ttk.Entry(self.tree);widget.insert(0,self.rows[int(item)][j]);widget.place(x=box[0],y=box[1],width=box[2],height=box[3]);widget.select_range(0,'end');widget.focus_set();self.editor=(widget,int(item),j)
        widget.bind('<Return>',lambda e:self.commit());widget.bind('<FocusOut>',lambda e:self.commit());widget.bind('<Escape>',lambda e:self.cancel())
    def cancel(self):
        if self.editor:
            widget,_,_=self.editor;self.editor=None;widget.destroy()
    def snapshot(self):
        rows=[row.copy() for row in self.rows]
        if self.editor:
            widget,row,column=self.editor;rows[row][column]=widget.get().strip()
        out=io.StringIO();writer=csv.writer(out,delimiter=';',lineterminator='\n');writer.writerow(self.header);writer.writerows(rows);return out.getvalue()
    def get(self,*args):
        self.commit();return self.snapshot()
    def delete(self,*args):self.cancel();self.header=[];self.rows=[];self.refresh()
    def insert(self,index,text):
        self.cancel();lines=[line for line in text.strip().splitlines() if line.strip()]
        if not lines:return
        delimiter=';' if ';' in lines[0] else '\t' if '\t' in lines[0] else ','
        rows=list(csv.reader(lines,delimiter=delimiter));size=len(rows[0])
        if any(len(row)!=size for row in rows):raise ValueError('Các hàng không có cùng số cột.')
        self.header=rows[0];self.rows=rows[1:];self.refresh()
    def add_row(self):
        self.commit()
        if not self.header:return
        self.rows.append(['']*len(self.header));self.refresh();self.tree.see(str(len(self.rows)-1))
    def remove_rows(self):
        selected=[int(item) for item in self.tree.selection()];self.commit()
        for i in sorted(selected,reverse=True):del self.rows[i]
        self.refresh()
    def copy(self):
        self.commit();chosen=self.tree.selection();rows=[self.rows[int(i)] for i in chosen] if chosen else self.rows
        out=io.StringIO();writer=csv.writer(out,delimiter='\t',lineterminator='\n');writer.writerow(self.header);writer.writerows(rows)
        self.clipboard_clear();self.clipboard_append(out.getvalue())
    def paste(self):
        self.commit()
        try:
            text=self.clipboard_get();lines=text.strip('\r\n').splitlines()
            if not lines:return
            delimiter='\t' if '\t' in lines[0] else ';' if ';' in lines[0] else ',';rows=list(csv.reader(lines,delimiter=delimiter))
            if rows[0]==self.header:rows=rows[1:]
            if any(len(row)!=len(self.header) for row in rows):raise ValueError(f'Dán đủ {len(self.header)} cột; có thể kèm tiêu đề giống mẫu.')
            self.rows.extend(rows);self.refresh()
        except (ValueError,tk.TclError) as e:messagebox.showerror('Dán bảng',str(e))
    def add_column(self):
        self.commit()
        if not self.header or self.header[0]!='Nam':messagebox.showinfo('Thời đoạn','Chỉ thêm cột thời đoạn ở bảng mưa theo năm.');return
        from tkinter import simpledialog
        value=simpledialog.askstring('Thêm thời đoạn','Thời đoạn mới (phút):',parent=self)
        if value is None:return
        try:
            duration=num(value,'Thời đoạn')
            if duration<=0 or any(abs(num(v)-duration)<1e-9 for v in self.header[1:]):raise ValueError('Thời đoạn phải dương và chưa có trong bảng.')
            self.header.append(f'{duration:g}')
            for row in self.rows:row.append('')
            self.refresh()
        except ValueError as e:messagebox.showerror('Thời đoạn',str(e))

class InlandApp:
    def __init__(self,root):
        self.root=root;root.title('Thủy văn tuyến đường — mưa thiết kế và úng nội đồng');root.geometry('1280x920');root.minsize(1050,740)
        style=ttk.Style();style.theme_use('clam');style.configure('TNotebook.Tab',padding=(15,8));style.configure('Treeview',rowheight=25)
        self.result=None;self.stats=[];self._fit_source=None
        self.project=tk.StringVar(value='Dự án tuyến đường');self.station=tk.StringVar(value='Trạm mưa');self.p=tk.StringVar(value='1;2;4;10');self.method=tk.StringVar(value='Pearson III')
        self.start=tk.StringVar(value='');self.end=tk.StringVar(value='');self.year=tk.StringVar(value='');self.cs=tk.StringVar(value='')
        self.scenario=tk.StringVar(value=SCENARIOS[0]);self.beta1=tk.StringVar(value='');self.beta2=tk.StringVar(value='');self.c=tk.StringVar(value='');self.ratio=tk.StringVar(value='')
        self.status=tk.StringVar(value='Nhập số liệu trạm của dự án. Thời kỳ cơ sở và năm điều tra có thể để trống khi chỉ phân tích mưa.')
        header=ttk.Frame(root,padding=12);header.pack(fill='x');ttk.Label(header,text='MƯA THIẾT KẾ & ÚNG NỘI ĐỒNG',font=('Arial',16,'bold')).pack(side='left')
        for title,command in [('Mở dự án',self.open_project),('Lưu dự án',self.save_project),('Xuất Excel',self.export_excel),('Xuất Word',self.export_word)]:ttk.Button(header,text=title,command=command).pack(side='right',padx=3)
        meta=ttk.Frame(root,padding=(12,0));meta.pack(fill='x')
        for name,var,width in [('Dự án',self.project,35),('Trạm',self.station,24),('Năm Xi mặc định',self.year,8)]:ttk.Label(meta,text=name).pack(side='left',padx=4);ttk.Entry(meta,textvariable=var,width=width).pack(side='left')
        self.book=ttk.Notebook(root);self.book.pack(fill='both',expand=True,padx=12,pady=8)
        self.rain_tab=self.tab('1. Chọn P & phân tích mưa');self.design_tab=self.tab('2. Xp & kịch bản khí hậu');self.inland_tab=self.tab('3. Điều tra Hi, Xi & hệ số β');self.results_tab=self.tab('4. ΔH & Htk khống chế');self.plot_tab=self.tab('5. Đường tần suất mưa')
        self.build_rain();self.build_design();self.build_inland();self.build_results();self.build_plot()
        ttk.Label(root,textvariable=self.status,wraplength=1220,foreground='#155e75',padding=10).pack(fill='x')
        self.root.after(300,self.watch)
    def tab(self,title):
        frame=ttk.Frame(self.book,padding=12);self.book.add(frame,text=title);return frame
    def text(self,parent,height=15):
        return DataGrid(parent)
    def put(self,widget,value):widget.delete('1.0','end');widget.insert('1.0',value)
    def build_rain(self):
        ttk.Label(self.rain_tab,text='Mỗi hàng là một năm; cột thời đoạn dùng PHÚT; giá trị là lượng mưa cực đại năm (mm). Không tự biến mưa ngày thành mưa 5–60 phút.',wraplength=1150).pack(anchor='w')
        controls=ttk.Frame(self.rain_tab);controls.pack(fill='x',pady=8)
        ttk.Button(controls,text='Nhập CSV / TXT / Excel',command=lambda:self.import_table(self.annual)).pack(side='left')
        ttk.Button(controls,text='Nhập mưa ngày → cực đại 1,3,5,7 ngày',command=self.import_daily).pack(side='left',padx=6)
        self.annual=self.text(self.rain_tab);self.put(self.annual,'Nam;1440;4320;7200;10080\n')
        row=ttk.Frame(self.rain_tab);row.pack(fill='x')
        for title,var,width in [('P (%) phân cách ;',self.p,16),('Cơ sở từ',self.start,7),('đến',self.end,7),('Cs trực tiếp (trống: theo mẫu)',self.cs,8)]:ttk.Label(row,text=title).pack(side='left',padx=4);ttk.Entry(row,textvariable=var,width=width).pack(side='left')
        ttk.Combobox(row,textvariable=self.method,values=['Pearson III','Gumbel'],state='readonly',width=14).pack(side='left',padx=5)
        ttk.Button(self.rain_tab,text='Phân tích tần suất → tạo bảng mưa thiết kế',command=self.fit).pack(anchor='e',pady=8)
        ttk.Label(self.rain_tab,text='Phương pháp mômen: độ lệch chuẩn mẫu n−1, Cs mẫu hiệu chỉnh scipy. Cs nhập tay áp dụng cho các chuỗi; muốn dùng hệ số riêng từng thời đoạn, nhập bảng mưa đã kiểm chứng ở bước 2.').pack(anchor='w')
    def build_design(self):
        ttk.Label(self.design_tab,text='Nhập hoặc dán dữ liệu theo các cột bên dưới. Nhấp đúp vào ô để sửa.\nMưa khí hậu = mưa CƠ SỞ × (1 + mức thay đổi/100). Xi phải cùng năm và cùng thời đoạn với mực nước điều tra.',wraplength=1150).pack(anchor='w')
        row=ttk.Frame(self.design_tab);row.pack(fill='x',pady=8)
        ttk.Button(row,text='Nhập bảng mưa thiết kế',command=lambda:self.import_table(self.design)).pack(side='left',padx=5)
        climate=ttk.Frame(self.design_tab);climate.pack(fill='x',pady=6)
        self.duration_edit=tk.StringVar(value='1440');self.mid_edit=tk.StringVar(value='');self.end_edit=tk.StringVar(value='')
        for title,var in [('Thời đoạn (phút)',self.duration_edit),('Thay đổi KB1 (%)',self.mid_edit),('Thay đổi KB2 (%)',self.end_edit)]:
            ttk.Label(climate,text=title).pack(side='left',padx=5);ttk.Entry(climate,textvariable=var,width=12).pack(side='left')
        ttk.Button(climate,text='Áp dụng cho thời đoạn',command=self.apply_change).pack(side='left',padx=10)
        ttk.Button(climate,text='Tính % từ mưa cơ sở / tương lai',command=self.change_dialog).pack(side='left')
        self.design=self.text(self.design_tab);self.put(self.design,GROWTH_HEADER+'\n')
        ttk.Label(self.design_tab,text='Mức thay đổi do người dùng cung cấp theo trạm, thời đoạn và kịch bản áp dụng. Không có tỷ lệ mặc định cho địa phương nào. Để trống khi chưa có căn cứ.',wraplength=1150,foreground='#9a3412').pack(anchor='w')
    def change_dialog(self):
        window=tk.Toplevel(self.root);window.title('Mức thay đổi lượng mưa');window.resizable(False,False)
        entries=[]
        for i,title in enumerate(['Mưa thời kỳ cơ sở (mm)','Mưa kịch bản tương lai (mm)']):
            ttk.Label(window,text=title).grid(row=i,column=0,padx=12,pady=8);entry=ttk.Entry(window);entry.grid(row=i,column=1,padx=12);entries.append(entry)
        target=tk.StringVar(value='Kịch bản 2');ttk.Combobox(window,textvariable=target,state='readonly',values=SCENARIOS[1:3]).grid(row=2,column=0,columnspan=2,pady=8)
        ttk.Label(window,text='ΔX% = (X_tương_lai / X_cơ_sở − 1) × 100\nDùng số liệu cùng trạm, thời đoạn, tần suất. Không coi\nchênh lệch toàn chuỗi/cơ sở là dự báo khí hậu.',justify='left').grid(row=3,column=0,columnspan=2,padx=12)
        def apply():
            try:
                rate=change_rate(entries[0].get(),entries[1].get())
                (self.mid_edit if target.get()=='Kịch bản 1' else self.end_edit).set(f'{rate:.12g}')
                self.status.set(f'Mức thay đổi {rate:.3f}%. Kiểm tra thời đoạn rồi bấm Áp dụng cho thời đoạn.');window.destroy()
            except ValueError as e:messagebox.showerror('Mức thay đổi',str(e),parent=window)
        ttk.Button(window,text='Tính và điền mức thay đổi',command=apply).grid(row=4,column=0,columnspan=2,pady=12)

    def apply_change(self):
        try:
            rows=table(self.read_text(self.design));duration=num(self.duration_edit.get(),'Thời đoạn')
            for value in (self.mid_edit.get(),self.end_edit.get()):
                if value.strip() and num(value,'Mức thay đổi')<=-100:raise ValueError('Mức thay đổi phải > −100%.')
            count=0
            for row in rows[1:]:
                if num(row[0])==duration:
                    if not row[3]:raise ValueError('Thời đoạn chưa có mưa thời kỳ cơ sở; không tự lấy toàn chuỗi thay thế.')
                    row[5]=self.mid_edit.get();row[6]=self.end_edit.get();count+=1
            if not count:raise ValueError('Không tìm thấy thời đoạn này trong bảng.')
            self.put(self.design,'\n'.join(';'.join(row) for row in rows));self.status.set(f'Đã cập nhật mức thay đổi cho {count} tần suất ở thời đoạn {duration:g} phút.')
        except Exception as e:messagebox.showerror('Mức thay đổi',str(e))

    def fill_observed(self):
        try:
            years,durations,matrix=parse_annual(self.read_text(self.annual));locations=parse_locations(self.read_text(self.locations));rows=[OBS_HEADER]
            for year in sorted({r['year'] for r in locations}):
                if year not in years:raise ValueError(f'Chuỗi mưa không có năm {year}; cần nhập Xi từ nguồn phù hợp.')
                for j,d in enumerate(durations):
                    if d>=1440:rows.append(f'{year};{d:g};{matrix[years==year,j][0]:.12g}')
            self.put(self.observed,'\n'.join(rows));self.status.set('Đã ghép mưa Xi theo năm điều tra và thời đoạn; cần kiểm tra trạm đại diện trước khi tính.')
        except Exception as e:messagebox.showerror('Mưa năm điều tra',str(e))
    def build_inland(self):
        ttk.Label(self.inland_tab,text='Htk = Hi + (1 + β) × (Xp% − Xi) / 1000; β = β1 + β2 + C × An/A\nLấy ΔH lớn nhất giữa các thời đoạn ≥1 ngày. Giữ dấu âm của ΔH. Htk là cao độ nước, chưa phải cao độ nền đường.',font=('Arial',11),wraplength=1150).pack(anchor='w',pady=8)
        row=ttk.Frame(self.inland_tab);row.pack(fill='x')
        for title,var in [('β1',self.beta1),('β2',self.beta2),('C',self.c),('An/A (0–1)',self.ratio)]:ttk.Label(row,text=title).pack(side='left',padx=5);ttk.Entry(row,textvariable=var,width=9).pack(side='left')
        ttk.Combobox(row,textvariable=self.scenario,state='readonly',values=SCENARIOS,width=28).pack(side='left',padx=12)
        ttk.Label(self.inland_tab,text='Nhập β1, β2 theo điều kiện mặt phủ/canh tác của dự án; C là hệ số dòng chảy; An/A là tỷ lệ diện tích không ngập. Không dùng mặc định của địa phương khác.',wraplength=1150,foreground='#9a3412').pack(anchor='w',pady=6)
        ttk.Label(self.inland_tab,text='Các vị trí: Vi_tri;Hi_m;Nam_dieu_tra;Z_dat_m (cao độ đất có thể để trống). Có thể nhập nhiều năm tại cùng vị trí; Xi phải khớp năm Hi. Bảng Xi riêng được ưu tiên; Xi ở bước 2 chỉ dùng cho năm mặc định.').pack(anchor='w',pady=10)
        observed_tabs=ttk.Notebook(self.inland_tab);observed_tabs.pack(fill='both',expand=True)
        place_tab=ttk.Frame(observed_tabs);rain_tab=ttk.Frame(observed_tabs)
        observed_tabs.add(place_tab,text='Mực nước điều tra từng vị trí / năm');observed_tabs.add(rain_tab,text='Mưa Xi cùng năm, cùng thời đoạn')
        self.locations=self.text(place_tab);self.put(self.locations,POINT_HEADER+'\n')
        ttk.Button(rain_tab,text='Lấy Xi từ chuỗi mưa cho các năm điều tra',command=self.fill_observed).pack(anchor='w',pady=4)
        ttk.Button(rain_tab,text='Nhập bảng Xi',command=lambda:self.import_table(self.observed)).pack(anchor='w')
        self.observed=self.text(rain_tab);self.put(self.observed,OBS_HEADER+'\n')
        ttk.Button(self.inland_tab,text='Nhập vị trí tuyến',command=lambda:self.import_table(self.locations)).pack(side='left')
        ttk.Button(self.inland_tab,text='TÍNH ΔH → Htk KHỐNG CHẾ',command=self.calculate).pack(side='right',pady=10)
    def tree(self,parent,columns):
        frame=ttk.Frame(parent);frame.pack(fill='both',expand=True)
        widget=ttk.Treeview(frame,columns=list(columns),show='headings');widget.grid(row=0,column=0,sticky='nsew')
        frame.rowconfigure(0,weight=1);frame.columnconfigure(0,weight=1)
        for key,title in columns.items():widget.heading(key,text=title);widget.column(key,width=135,anchor='center')
        y=ttk.Scrollbar(frame,command=widget.yview);y.grid(row=0,column=1,sticky='ns');x=ttk.Scrollbar(frame,orient='horizontal',command=widget.xview);x.grid(row=1,column=0,sticky='ew');widget.configure(yscrollcommand=y.set,xscrollcommand=x.set);return widget
    def build_results(self):
        result_tabs=ttk.Notebook(self.results_tab);result_tabs.pack(fill='both',expand=True)
        final=ttk.Frame(result_tabs);by_year=ttk.Frame(result_tabs);detail=ttk.Frame(result_tabs)
        result_tabs.add(final,text='Htk khống chế theo vị trí / P');result_tabs.add(by_year,text='So sánh từng năm điều tra');result_tabs.add(detail,text='ΔH từng thời đoạn')
        self.tree_summary=self.tree(final,dict(location='Vị trí',p='P (%)',year='Năm khống chế',Hi='Hi (m)',deltaH='ΔH khống chế (m)',Hp='Htk (m)',days='Thời đoạn (ngày)',scenario='Kịch bản',depth='Chiều sâu ngập (m)'))
        self.tree_years=self.tree(by_year,dict(location='Vị trí',p='P (%)',year='Năm điều tra',Hi='Hi (m)',deltaH='ΔH (m)',Hp='Htk năm (m)',days='Thời đoạn (ngày)',scenario='Kịch bản'))
        self.tree_details=self.tree(detail,dict(location='Vị trí',year='Năm',p='P (%)',duration='Thời đoạn (phút)',Xp='Xp (mm)',Xi='Xi (mm)',beta='β',deltaH='ΔH (m)',Hp='Htk (m)'))
        ttk.Label(self.results_tab,text='Bảng tính chi tiết, dữ liệu gốc, giả thiết và các cảnh báo được lưu trong báo cáo Excel/Word.').pack(pady=8)
    def clear_tables(self):
        for tree in (self.tree_summary,self.tree_years,self.tree_details):tree.delete(*tree.get_children())
    def build_plot(self):
        row=ttk.Frame(self.plot_tab);row.pack(fill='x');self.plot_source=tk.StringVar(value='Hiện trạng');ttk.Combobox(row,textvariable=self.plot_source,values=SCENARIOS[:3],state='readonly',width=20).pack(side='left',padx=5);ttk.Button(row,text='Vẽ đường tần suất & điểm thực nghiệm',command=self.plot).pack(side='left');ttk.Button(row,text='Lưu biểu đồ PNG',command=self.save_plot).pack(side='left',padx=5)
        self.figure=Figure(figsize=(11,6),dpi=100);self.canvas=FigureCanvasTkAgg(self.figure,self.plot_tab);self.canvas.get_tk_widget().pack(fill='both',expand=True);NavigationToolbar2Tk(self.canvas,self.plot_tab)
    def read_text(self,widget):return widget.get('1.0','end').strip()
    def signature(self):
        variables=(self.project,self.station,self.p,self.start,self.end,self.year,self.cs,self.method,self.scenario,self.beta1,self.beta2,self.c,self.ratio)
        return tuple(v.get() for v in variables)+tuple(w.snapshot().strip() for w in (self.annual,self.design,self.locations,self.observed))
    def watch(self):
        if self.result is not None and self.signature()!=self._signature:
            self.result=None;self.clear_tables();self.figure.clear();self.canvas.draw();self.status.set('Dữ liệu đã thay đổi: cần tính lại trước khi xuất báo cáo.')
        self.root.after(300,self.watch)
    def import_table(self,widget):
        path=filedialog.askopenfilename(filetypes=[('Bảng số liệu','*.csv *.txt *.xlsx'),('Tất cả','*.*')])
        if not path:return
        try:
            if Path(path).suffix.lower()=='.xlsx' and widget is self.annual:
                try:text,meta=read_rain_workbook(path)
                except ValueError as e:
                    if not str(e).startswith('Có nhiều bảng mưa'):raise
                    from tkinter import simpledialog
                    sheet=simpledialog.askstring('Chọn trang tính',str(e)+'\nNhập tên trang cần sử dụng:',parent=self.root)
                    if not sheet:return
                    text,meta=read_rain_workbook(path,sheet)
                self.put(widget,text)
                if meta['station']:self.station.set(meta['station'])
                warning=''
                try:outside=num(self.end.get())<meta['year_start'] or num(self.start.get())>meta['year_end']
                except ValueError:outside=True
                if not self.start.get().strip() and not self.end.get().strip():
                    warning=' Thời kỳ cơ sở chưa khai báo: có thể phân tích toàn chuỗi trước; chỉ cần chọn cơ sở khi xét kịch bản khí hậu.'
                elif outside:
                    warning=f" Thời kỳ cơ sở {self.start.get()}–{self.end.get()} không nằm trong chuỗi; hãy chọn lại theo nguồn kịch bản, không tự dùng toàn chuỗi thay thế."
                self.status.set(f"Đã đọc trang {meta['sheet']}: {meta['count']} năm ({meta['year_start']}–{meta['year_end']}), các thời đoạn phút {meta['durations']}."+warning)
                return
            elif Path(path).suffix.lower()=='.xlsx':text=pd.read_excel(path,dtype=str).fillna('').to_csv(index=False,sep=';')
            else:text=Path(path).read_text(encoding='utf-8-sig')
            table(text);self.put(widget,text)
        except Exception as e:messagebox.showerror('Không nhập được',str(e))
    def import_daily(self):
        path=filedialog.askopenfilename(filetypes=[('Mưa ngày CSV / TXT','*.csv *.txt')])
        if not path:return
        try:
            text,omitted=daily_to_annual(Path(path).read_text(encoding='utf-8-sig'));self.put(self.annual,text)
            self.status.set('Tổng trượt trong từng năm; bỏ năm thiếu ngày: '+str(omitted)+'. Không lấy cửa sổ qua ranh giới năm.')
        except Exception as e:messagebox.showerror('Mưa ngày',str(e))
    def fit(self):
        try:
            years,durations,matrix=parse_annual(self.read_text(self.annual));probs=[num(v,'P') for v in self.p.get().split(';') if v.strip()]
            if not probs or len(set(probs))!=len(probs):raise ValueError('Danh sách P trống hoặc trùng.')
            year=num(self.year.get(),'Năm điều tra') if self.year.get().strip() else None
            if year is not None and (year!=int(year) or year not in years):raise ValueError('Năm điều tra phải nguyên và có trong chuỗi để lấy Xi; hoặc để trống khi chỉ phân tích mưa.')
            if bool(self.start.get().strip())!=bool(self.end.get().strip()):raise ValueError('Nhập cả hai năm thời kỳ cơ sở hoặc để trống cả hai.')
            baseline=None
            if self.start.get().strip():
                start=num(self.start.get());end=num(self.end.get())
                if start!=int(start) or end!=int(end) or start>end:raise ValueError('Thời kỳ cơ sở phải là năm nguyên, có thứ tự.')
                baseline=(years>=start)&(years<=end)
            stats=[];records=[GROWTH_HEADER]
            for j,duration in enumerate(durations):
                full,meta=frequency(matrix[:,j],probs,self.method.get(),self.cs.get().strip() or None)
                stats.append(dict(duration=duration,period='Toàn chuỗi',**meta))
                base=[None]*len(probs)
                if baseline is not None:
                    base,bmeta=frequency(matrix[baseline,j],probs,self.method.get(),self.cs.get().strip() or None)
                    stats.append(dict(duration=duration,period='Cơ sở',**bmeta))
                xi=float(matrix[years==year,j][0]) if year is not None else None
                for p,x,b in zip(probs,full,base):records.append(';'.join('' if value is None else f'{value:.12g}' for value in [duration,p,x,b,xi])+ ';;')
            self.put(self.design,'\n'.join(records));self.stats=stats;self._fit_source=(self.read_text(self.annual),self.method.get(),self.cs.get(),self.start.get(),self.end.get(),self.read_text(self.design))
            self.book.select(self.design_tab);self.status.set('Đã phân tích toàn chuỗi. '+('Đã tính thêm thời kỳ cơ sở. ' if baseline is not None else 'Chưa khai báo cơ sở: chưa tính mưa kịch bản khí hậu. ')+('Xi đã lấy theo năm điều tra. ' if year is not None else 'Chưa khai báo năm điều tra: chỉ tính mưa, chưa tính úng. ')+'Chuỗi <30 năm cần thận trọng khi ngoại suy P nhỏ.')
        except Exception as e:messagebox.showerror('Phân tích tần suất',str(e))
    def calculate(self):
        self.result=None;self.clear_tables()
        try:
            year=num(self.year.get(),'Năm Xi mặc định') if self.year.get().strip() else None
            if year is not None and year!=int(year):raise ValueError('Năm mặc định phải nguyên.')
            design=parse_design(self.read_text(self.design));locations=parse_locations(self.read_text(self.locations))
            probabilities=[num(value,'P thiết kế') for value in self.p.get().split(';') if value.strip()]
            if not probabilities or any(not 0<p<100 for p in probabilities):raise ValueError('Chọn P thiết kế trong (0;100).')
            if any(not any(r['p']==p for r in design) for p in probabilities):raise ValueError('Bảng Xp chưa có đủ các tần suất thiết kế đã chọn.')
            design=[r for r in design if r['p'] in probabilities]
            details,per_year,summary,notes=calculate_multi(design,locations,parse_observed(self.read_text(self.observed)),num(self.beta1.get()),num(self.beta2.get()),num(self.c.get()),num(self.ratio.get()),self.scenario.get(),year)
            if self._fit_source and self._fit_source!=(self.read_text(self.annual),self.method.get(),self.cs.get(),self.start.get(),self.end.get(),self.read_text(self.design)):notes.append('Bảng mưa hoặc đầu vào phân tích đã được sửa; không gắn thống kê cũ vào kết quả hiện tại.')
            for tree,records in [(self.tree_summary,summary),(self.tree_years,per_year),(self.tree_details,details)]:
                tree.delete(*tree.get_children())
                for row in records:tree.insert('','end',values=[row[k] if isinstance(row[k],str) else '—' if row[k] is None else f'{row[k]:.3f}' for k in tree['columns']])
            self.result=dict(design=design,details=details,summary=summary,notes=notes,project=self.project.get(),station=self.station.get(),scenario=self.scenario.get(),year=year,per_year=per_year,beta=num(self.beta1.get())+num(self.beta2.get())+num(self.c.get())*num(self.ratio.get()),locations=locations,coefficients={'beta1':num(self.beta1.get()),'beta2':num(self.beta2.get()),'C':num(self.c.get()),'An_A':num(self.ratio.get())})
            self._signature=self.signature();self.book.select(self.results_tab);self.status.set(f'Đã tính {len(summary)} kết quả Htk khống chế. '+(f'Có {len(notes)} ghi chú về dữ liệu thiếu/đã sửa; xem báo cáo xuất. Không tự gán mức tăng cho thời đoạn thiếu.' if notes else 'Đủ dữ liệu cho kịch bản đang dùng.'))
        except Exception as e:messagebox.showerror('Không tính được',str(e))
    def current(self):
        if self.result is None or self.signature()!=self._signature:
            messagebox.showwarning('Cần tính lại','Chưa có kết quả hợp lệ ứng với dữ liệu hiện tại.');return False
        return True
    def plot(self):
        self.figure.clear()
        try:
            years,durations,matrix=parse_annual(self.read_text(self.annual))
            selected=self.plot_source.get();rows=parse_design(self.read_text(self.design))
            mask=np.ones(len(years),dtype=bool)
            if selected!='Hiện trạng':
                start=num(self.start.get(),'Năm đầu thời kỳ cơ sở');end=num(self.end.get(),'Năm cuối thời kỳ cơ sở')
                if start!=int(start) or end!=int(end) or start>end:raise ValueError('Thời kỳ cơ sở không hợp lệ.')
                mask=(years>=start)&(years<=end)
            ax=self.figure.add_subplot(111)
            # Giấy xác suất chuẩn: tọa độ x = Φ⁻¹(P vượt/100), nhãn vẫn là P%.
            ticks=np.array([.1,.2,.5,1,2,5,10,20,30,50,70,80,90,95,98,99,99.5,99.8,99.9])
            probabilities=norm.cdf(np.linspace(norm.ppf(.001),norm.ppf(.999),600))*100
            plotted=0;omitted=[]
            for j,duration in enumerate(durations):
                if duration<1440:continue
                values=matrix[mask,j];factor=1.
                if selected!='Hiện trạng':
                    key='mid' if selected=='Kịch bản 1' else 'end'
                    rates={r[key] for r in rows if r['duration']==duration and r[key] is not None}
                    if not rates:omitted.append(f'{duration/1440:g} ngày');continue
                    if len(rates)!=1:raise ValueError('Mức thay đổi theo P khác nhau: chưa thể xác định một đường phân phối khí hậu duy nhất cho thời đoạn này.')
                    factor=1+next(iter(rates))/100
                _,meta=frequency(values,[50],self.method.get(),self.cs.get().strip() or None)
                mean=meta['mean'];sd=meta['cv']*mean
                if sd==0:quantiles=np.full(len(probabilities),mean)
                elif self.method.get()=='Gumbel':
                    scale=sd*np.sqrt(6)/np.pi
                    quantiles=gumbel_r.isf(probabilities/100,loc=mean-np.euler_gamma*scale,scale=scale)
                else:quantiles=mean+sd*pearson3.isf(probabilities/100,skew=meta['cs_used'])
                valid=np.isfinite(quantiles)&(quantiles>=0)
                line,=ax.plot(norm.ppf(probabilities[valid]/100),quantiles[valid]*factor,label=f'{duration/1440:g} ngày — {meta["method"]}, n={meta["n"]}')
                empirical=np.sort(values)[::-1];p_emp=np.arange(1,len(values)+1)/(len(values)+1)
                ax.scatter(norm.ppf(p_emp),empirical*factor,color=line.get_color(),s=22,facecolors='none')
                plotted+=1
            if not plotted:raise ValueError('Không có chuỗi mưa và mức thay đổi đủ dữ liệu để vẽ.')
            ax.set_xticks(norm.ppf(ticks/100));ax.set_xticklabels([f'{v:g}' for v in ticks],rotation=45,fontsize=8)
            ax.set_xlim(norm.ppf(.001),norm.ppf(.999));ax.set_ylim(bottom=0)
            ax.set(xlabel='Tần suất vượt P (%) — thang xác suất chuẩn',ylabel='Lượng mưa X (mm)',title='ĐƯỜNG TẦN SUẤT MƯA — '+self.station.get()+' — '+selected)
            ax.grid(alpha=.3);ax.legend(fontsize=8)
            subtitle='Đường liền: phân phối lý luận (ước lượng mômen). Điểm: P = m/(n+1), xếp mưa giảm dần (Weibull).'
            if selected!='Hiện trạng':subtitle+=' Điểm kịch bản được nhân hệ số thay đổi; không phải quan trắc tương lai.'
            self.figure.text(.5,.015,subtitle,ha='center',fontsize=8,wrap=True)
            self.figure.tight_layout(rect=(0,.07,1,1));self.canvas.draw();self.book.select(self.plot_tab)
            self.status.set('Đã vẽ trên giấy xác suất chuẩn; không nối các điểm P thiết kế. Phần phân vị mưa âm không hiển thị.'+(' Thiếu kịch bản: '+', '.join(omitted) if omitted else ''))
        except Exception as e:self.figure.clear();self.canvas.draw();messagebox.showerror('Biểu đồ',str(e))
    def save_plot(self):
        self.plot()
        if not self.figure.axes:return
        path=filedialog.asksaveasfilename(defaultextension='.png',filetypes=[('Ảnh PNG','*.png')])
        if path:self.figure.savefig(path,dpi=200)
    def project_data(self):
        return dict(version=1,vars={name:getattr(self,name).get() for name in ['project','station','p','method','start','end','year','cs','scenario','beta1','beta2','c','ratio']},annual=self.read_text(self.annual),design=self.read_text(self.design),locations=self.read_text(self.locations),observed=self.read_text(self.observed))
    def save_project(self):
        path=filedialog.asksaveasfilename(defaultextension='.ung',filetypes=[('Dự án úng nội đồng','*.ung')])
        if path:
            try:Path(path).write_text(json.dumps(self.project_data(),ensure_ascii=False,indent=2),encoding='utf-8')
            except OSError as e:messagebox.showerror('Lưu dự án',str(e))
    def open_project(self):
        path=filedialog.askopenfilename(filetypes=[('Dự án úng nội đồng','*.ung')])
        if not path:return
        try:
            data=json.loads(Path(path).read_text(encoding='utf-8'));vars=data['vars']
            for name in ['project','station','p','method','start','end','year','cs','scenario','beta1','beta2','c','ratio']:
                if name in vars:getattr(self,name).set(vars[name])
            mapping={'Giữa thế kỷ':SCENARIOS[1],'Cuối thế kỷ':SCENARIOS[2],'Bao lớn nhất 3 kịch bản':SCENARIOS[3]}
            if self.scenario.get() in mapping:self.scenario.set(mapping[self.scenario.get()])
            for name in ['annual','design','locations']:self.put(getattr(self,name),data[name])
            self.put(self.observed,data.get('observed',OBS_HEADER+'\n'))
            self.result=None;self.stats=[];self._fit_source=None;self.clear_tables();self.figure.clear();self.canvas.draw();self.status.set('Đã mở dự án. Bấm tính để tạo kết quả.')
        except Exception as e:messagebox.showerror('Mở dự án',str(e))
    def assumptions(self):
        return ['Htk là cao độ mặt nước (m); X là lượng mưa (mm). Không tự chọn cao độ nền đường.',f"Các năm điều tra: {sorted({r['year'] for r in self.result['locations']})}; β={self.result['beta']:.6g}; thành phần: {self.result['coefficients']}.",'Htk = Hi + (1+β)(Xp−Xi)/1000; lấy ΔH lớn nhất theo thời đoạn, không làm tròn trước khi tính.',f"Kịch bản: {self.result['scenario']}. Mức thay đổi áp dụng lên mưa thời kỳ cơ sở; không nhân toàn chuỗi nếu chưa có căn cứ.",'Hi và Xi phải cùng năm; Xi và Xp phải cùng thời đoạn; Hi và cao độ đất phải cùng hệ cao độ; lấy ΔH lớn nhất theo thời đoạn từng năm rồi Htk lớn nhất trong các năm đã nhập. Không làm tròn số trung gian.','Pearson III/Gumbel ước lượng mômen; người dùng cần kiểm tra độ phù hợp phân phối, độ dài chuỗi và tính đại diện của trạm.']+self.result['notes']
    def export_excel(self):
        if not self.current():return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',filetypes=[('Excel','*.xlsx')])
        if not path:return
        try:
            with pd.ExcelWriter(path,engine='openpyxl') as writer:
                for sheet,key in [('Mua_thiet_ke','design'),('DeltaH_chi_tiet','details'),('Hp_tuyen','summary'),('Vi_tri','locations'),('Htk_tung_nam','per_year')]:pd.DataFrame(self.result[key]).rename(columns=COLUMN_TITLES).to_excel(writer,sheet_name=sheet,index=False)
                pd.DataFrame({'Ghi_chu':[self.project.get(),self.station.get()]+self.assumptions()}).to_excel(writer,sheet_name='Can_cu_gia_thiet',index=False)
                if self.stats and self._fit_source and self._fit_source[:5]==(self.read_text(self.annual),self.method.get(),self.cs.get(),self.start.get(),self.end.get()):
                    pd.DataFrame(self.stats).to_excel(writer,sheet_name='Thong_ke_tan_suat',index=False)
                pd.DataFrame([{'Ký hiệu':key,'Diễn giải':title} for key,title in COLUMN_TITLES.items()]).to_excel(writer,sheet_name='Dien_giai_ky_hieu',index=False)
                for name in ('annual','design','locations','observed'):
                    rows=table(self.read_text(getattr(self,name))) if len(self.read_text(getattr(self,name)).splitlines())>1 else []
                    if rows:pd.DataFrame(rows[1:],columns=rows[0]).to_excel(writer,sheet_name='Goc_'+name,index=False)
                from openpyxl.styles import Font,PatternFill
                for ws in writer.book:
                    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
                    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='155E75')
                    for row in ws.iter_rows(min_row=2):
                        for cell in row:
                            if isinstance(cell.value,(float,int)):cell.number_format='0.000'
                    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(65,max(16,len(str(column[0].value))+3))
            self.status.set('Đã xuất Excel: số liệu gốc, bảng mưa, ΔH, Htk và giả thiết.')
        except Exception as e:messagebox.showerror('Xuất Excel',str(e))
    def export_word(self):
        if not self.current():return
        path=filedialog.asksaveasfilename(defaultextension='.docx',filetypes=[('Word','*.docx')])
        if not path:return
        try:
            from docx import Document
            from docx.shared import Cm,Pt
            doc=Document();section=doc.sections[0];section.page_width=Cm(29.7);section.page_height=Cm(21);section.left_margin=section.right_margin=Cm(1.5)
            doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('TÍNH MỰC NƯỚC ÚNG NỘI ĐỒNG',0);doc.add_paragraph(self.result['project']+' — '+self.result['station'])
            for note in self.assumptions():doc.add_paragraph(note)
            for title,key,cols in [('Mưa và chênh lệch mực nước','details',['location','year','duration','p','Xp','Xi','deltaH']),('Mực nước thiết kế từng năm điều tra','per_year',['location','p','year','Hi','deltaH','Hp','days','scenario']),('Mực nước thiết kế khống chế theo vị trí','summary',['location','p','year','Hi','deltaH','Hp','days','scenario'])]:
                doc.add_heading(title,1);t=doc.add_table(rows=1,cols=len(cols));t.style='Table Grid'
                for cell,keycol in zip(t.rows[0].cells,cols):cell.text=COLUMN_TITLES.get(keycol,keycol)
                for row in self.result[key]:
                    for cell,keycol in zip(t.add_row().cells,cols):
                        val=row[keycol];cell.text='—' if val is None else f'{val:.3f}' if isinstance(val,(int,float)) else str(val)
            doc.save(path);self.status.set('Đã xuất Word với công thức, giả thiết và bảng kết quả.')
        except Exception as e:messagebox.showerror('Xuất Word',str(e))

if __name__=='__main__':
    root=tk.Tk();app=InlandApp(root);root.mainloop()
