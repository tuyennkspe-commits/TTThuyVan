#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPHinhthaiMC v2 - Quan hệ H-Q, H-V theo phương pháp hình thái mặt cắt (Manning, chia dòng chủ/bãi)
+ diện tích thoát nước cầu (Bê-lê-li-út-xki) + nước dềnh trước cầu.

Cấu trúc:
  1. Lõi tính toán thuần (không phụ thuộc Tkinter -> kiểm thử được)
       - hình học ướt CHÍNH XÁC (có giao điểm mực nước với đáy sông)
       - chế độ "tương thích bản cũ" tái hiện đúng hình học & cách nội suy Htk của bản cũ
  2. Xuất Excel (công thức sống cho Manning)
  3. Biểu đồ (Figure, không dùng pyplot)
  4. Giao diện Tkinter

Yêu cầu: Python >= 3.9, numpy, scipy, pandas (chỉ để đọc Excel), openpyxl, matplotlib.
"""
from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass, field

import numpy as np
from scipy.optimize import brentq

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    HAS_TK = True
except ImportError:  # pragma: no cover
    HAS_TK = False
try:
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
    HAS_MPL = True
except ImportError:  # pragma: no cover
    HAS_MPL = False

APP_TITLE = "THỦY LỰC MẶT CẮT SÔNG - PHƯƠNG PHÁP HÌNH THÁI (H-Q, H-V) - THOÁT NƯỚC CẦU & NƯỚC DỀNH"
G = 9.81

DEFAULT_MC = [
    (0.0, 21.28), (41.79, 20.84), (56.07, 19.92), (79.27, 19.05), (89.04, 18.35),
    (92.66, 16.30), (97.03, 15.68), (100.00, 15.80), (105.10, 16.50), (106.58, 18.50),
    (118.41, 19.22), (140.14, 19.06), (163.84, 19.59), (171.40, 21.98), (182.84, 22.25),
]


# =============================================================================
# 1. LÕI TÍNH TOÁN
# =============================================================================
def parse_float(text, name="giá trị") -> float:
    try:
        return float(str(text).strip().replace(",", "."))
    except ValueError:
        raise ValueError(f"Ô «{name}» không phải số hợp lệ: “{text}”") from None


def _num(tok):
    return float(tok.replace(",", ".")) if re.fullmatch(r"[-+]?\d+([.,]\d+)?", tok) else None


def parse_xz_text(text: str):
    """Đọc cặp (X, Z) từ văn bản: tab/khoảng trắng/;/, đều được, chấp nhận dấu phẩy thập phân. Bỏ qua dòng tiêu đề."""
    pts = []
    for line in text.splitlines():
        toks = [t for t in re.split(r"[\s;]+", line.strip()) if t]
        if len(toks) < 2:
            toks = [t for t in line.strip().split(",") if t.strip()]
        vals = [v for v in (_num(t.strip()) for t in toks) if v is not None]
        if len(vals) >= 2:
            pts.append((vals[0], vals[1]))
    return pts


def detect_banks(pts):
    """Như bản cũ: mép bờ = đầu trên của đoạn dốc nhất ở mỗi phía so với điểm thấp nhất."""
    xs, zs = [p[0] for p in pts], [p[1] for p in pts]
    n = len(pts)
    i_min = int(np.argmin(zs))
    sl = [abs(zs[i + 1] - zs[i]) / max(abs(xs[i + 1] - xs[i]), 1e-4) for i in range(n - 1)]
    left = sl[:i_min]
    right = sl[i_min:]
    i_l = int(np.argmax(left)) if left else 0
    i_r = i_min + int(np.argmax(right)) + 1 if right else n - 1
    return xs[i_l], xs[i_r]


def default_levels(pts):
    zs = [p[1] for p in pts]
    z_min, z_max = min(zs), max(zs)
    dh = round((z_max - z_min) / 8.0, 2) or 0.3
    return z_max, round(z_min + dh, 2), dh


def make_levels(h_max, h_min, dh):
    if dh <= 0 or h_max <= h_min:
        raise ValueError("Cần Hmax > Hmin và bước dH > 0.")
    n = int(round((h_max - h_min) / dh)) + 1
    return [round(h_max - i * dh, 2) for i in range(n)]


# ---- Hình học ướt chính xác --------------------------------------------------
def clip_zone(pts, x0, x1):
    xs, zs = [p[0] for p in pts], [p[1] for p in pts]
    return [(x0, float(np.interp(x0, xs, zs)))] + [p for p in pts if x0 < p[0] < x1] + [(x1, float(np.interp(x1, xs, zs)))]


def wet_piece(a, b, H):
    """Phần ướt của đoạn a-b dưới mực nước H (cắt tại giao điểm). None nếu khô hoàn toàn."""
    (x1, z1), (x2, z2) = a, b
    h1, h2 = H - z1, H - z2
    if h1 <= 0 and h2 <= 0:
        return None
    if h1 >= 0 and h2 >= 0:
        xa, ha, xb, hb = x1, h1, x2, h2
    else:
        xc = x1 + (x2 - x1) * h1 / (h1 - h2)
        xa, ha, xb, hb = (xc, 0.0, x2, h2) if h1 < 0 else (x1, h1, xc, 0.0)
    bw = abs(xb - xa)
    return dict(xa=xa, xb=xb, ha=ha, hb=hb, b=bw, w=0.5 * (ha + hb) * bw, dh=hb - ha, c=math.hypot(bw, hb - ha))


ZONE_LABEL = {"l": "Tổng cộng bãi trái", "m": "Tổng cộng dòng chủ", "r": "Tổng cộng bãi phải"}
ZONE_START = {"l": "Bãi trái", "m": "Dòng chủ", "r": "Bãi phải"}


def _total_row(z, b, w, dh, c):
    return dict(bophan=ZONE_LABEL[z], x=None, z=None, h=None, b=b, w=w, dh=dh, c=c, is_tot=True, raw_b=b, raw_w=w, raw_c=c)


def rows_exact(zone, zpts, H):
    rows, last, tb, tw, tdh, tc = [], None, 0.0, 0.0, 0.0, 0.0
    for a, b in zip(zpts[:-1], zpts[1:]):
        p = wet_piece(a, b, H)
        if not p:
            continue
        if last is None or abs(last - p["xa"]) > 1e-9:
            rows.append(dict(bophan=ZONE_START[zone] if not rows else "(vũng riêng)", x=p["xa"], z=H - p["ha"], h=p["ha"],
                             b=None, w=None, dh=None, c=None, is_tot=False))
        rows.append(dict(bophan="", x=p["xb"], z=H - p["hb"], h=p["hb"], b=p["b"], w=p["w"], dh=p["dh"], c=p["c"], is_tot=False))
        last = p["xb"]
        tb += p["b"]; tw += p["w"]; tdh += p["dh"]; tc += p["c"]
    rows.append(_total_row(zone, tb, tw, tdh, tc))
    return rows


# ---- Hình học của bản cũ (port nguyên văn để đối chiếu) ---------------------
def _legacy_prepare(mc, x_left, x_right):
    pts = sorted(list(mc), key=lambda p: p[0])
    xs = [p[0] for p in pts]
    x_min, x_max = xs[0], xs[-1]
    x_left = max(x_min, min(x_left, x_max))
    x_right = max(x_left, min(x_right, x_max))
    new_pts, ins_l, ins_r = [], False, False
    for i in range(len(pts)):
        p = pts[i]
        if not ins_l and p[0] >= x_left:
            if p[0] > x_left and i > 0:
                q = pts[i - 1]
                new_pts.append((x_left, q[1] + (x_left - q[0]) * (p[1] - q[1]) / (p[0] - q[0])))
            ins_l = True
        if not ins_r and p[0] >= x_right:
            if p[0] > x_right and i > 0:
                q = new_pts[-1] if new_pts else pts[i - 1]
                new_pts.append((x_right, q[1] + (x_right - q[0]) * (p[1] - q[1]) / (p[0] - q[0])))
            ins_r = True
        if not new_pts or abs(new_pts[-1][0] - p[0]) > 1e-5:
            new_pts.append(p)
    return new_pts, x_left, x_right


def rows_legacy(pts, H, xl_main, xr_main):
    new_pts, xl, xr = _legacy_prepare(pts, xl_main, xr_main)
    zones = {"l": [p for p in new_pts if p[0] <= xl + 1e-5],
             "m": [p for p in new_pts if xl - 1e-5 <= p[0] <= xr + 1e-5],
             "r": [p for p in new_pts if p[0] >= xr - 1e-5]}
    out = {}
    for z, zp in zones.items():
        if not zp:
            out[z] = [_total_row(z, 0.0, 0.0, 0.0, 0.0)]
            continue
        act = [dict(x=p[0], z=p[1], h=max(0.0, H - p[1])) for p in zp]
        rows, tb, tw, tdh, tc = [], 0.0, 0.0, 0.0, 0.0
        for i, pt in enumerate(act):
            if i == 0:
                hack = z == "l" and pt["h"] > 0
                rows.append(dict(bophan=ZONE_START[z] if z != "r" else "", x=pt["x"], z=H if hack else pt["z"], h=0.0 if hack else pt["h"],
                                 b=None, w=None, dh=None, c=None, is_tot=False))
            else:
                prev = act[i - 1]
                h_prev = 0.0 if (z == "l" and i == 1 and act[0]["h"] > 0) else prev["h"]
                b = pt["x"] - prev["x"]
                w = 0.5 * (h_prev + pt["h"]) * b
                dh = pt["h"] - h_prev
                c = math.sqrt(b * b + dh * dh)
                tb += b; tw += w; tdh += dh; tc += c
                rows.append(dict(bophan="", x=pt["x"], z=pt["z"], h=pt["h"], b=b, w=w, dh=dh, c=c, is_tot=False))
        rows.append(_total_row(z, tb, tw, tdh, tc))
        out[z] = rows
    return out


# ---- Thủy lực ---------------------------------------------------------------
@dataclass
class Params:
    pts: list
    x_left: float
    x_right: float
    inv_n: tuple            # (1/n bãi trái, 1/n dòng chủ, 1/n bãi phải)
    slope: float
    legacy: bool = False

    def sorted_pts(self):
        return sorted(self.pts, key=lambda p: p[0])

    def validate(self):
        errs, warns = [], []
        if len(self.pts) < 3:
            errs.append("Mặt cắt cần ít nhất 3 điểm")
        else:
            xs = [p[0] for p in self.sorted_pts()]
            if len(set(xs)) < len(xs):
                warns.append("Có điểm trùng cao độ X (thành đứng) - kiểm tra lại số liệu mặt cắt")
            if not (xs[0] <= self.x_left < self.x_right <= xs[-1]):
                errs.append(f"Cần X_min ≤ mép trái < mép phải ≤ X_max ({xs[0]:g} … {xs[-1]:g})")
        if any(n <= 0 for n in self.inv_n):
            errs.append("1/n phải > 0")
        if self.slope <= 0:
            errs.append("Độ dốc i phải > 0")
        return errs, warns


def evaluate(P: Params, H: float) -> dict:
    """Thủy lực toàn mặt cắt tại mực nước H. Trả về tổng từng vùng + bảng chi tiết."""
    pts = P.sorted_pts()
    if P.legacy:
        zrows = rows_legacy(pts, H, P.x_left, P.x_right)
    else:
        xs = [p[0] for p in pts]
        bounds = {"l": (xs[0], P.x_left), "m": (P.x_left, P.x_right), "r": (P.x_right, xs[-1])}
        zrows = {z: rows_exact(z, clip_zone(pts, *bounds[z]), H) for z in "lmr"}
    sq = math.sqrt(P.slope)
    zones, rows = {}, []
    for z, inv_n in zip("lmr", P.inv_n):
        t = zrows[z][-1]
        w, c, b = t["raw_w"], t["raw_c"], t["raw_b"]
        R = w / c if c > 0 else 0.0
        V = inv_n * R ** (2.0 / 3.0) * sq if R > 0 else 0.0
        zones[z] = dict(w=w, c=c, b=b, R=R, V=V, Q=w * V, inv_n=inv_n)
        rows += zrows[z]
    w_tot = sum(zones[z]["w"] for z in "lmr")
    Q_tot = sum(zones[z]["Q"] for z in "lmr")
    return dict(H=H, zones=zones, w_tot=w_tot, Q_tot=Q_tot, V_tb=Q_tot / w_tot if w_tot > 0 else 0.0,
                B=sum(zones[z]["b"] for z in "lmr"), rows=rows)


def _extrap_interp(x, xp, fp):
    """Nội suy tuyến tính có ngoại suy tuyến tính hai đầu (giống interp1d(..., fill_value='extrapolate'))."""
    if x < xp[0]:
        return fp[0] + (x - xp[0]) * (fp[1] - fp[0]) / (xp[1] - xp[0])
    if x > xp[-1]:
        return fp[-1] + (x - xp[-1]) * (fp[-1] - fp[-2]) / (xp[-1] - xp[-2])
    return float(np.interp(x, xp, fp))


def solve_H(P: Params, Q: float, table=None):
    """Tìm Htk ứng với Qtk. Trả về (H, [cảnh báo]).
    - Chế độ chính xác: giải phương trình Q(H) = Qtk bằng brentq trên hàm liên tục.
    - Chế độ tương thích bản cũ: nội suy tuyến tính trên bảng các cấp H (có ngoại suy) như bản cũ."""
    warns = []
    if P.legacy:
        if table is None or len(table) < 2:
            raise ValueError("Cần ≥ 2 cấp mực nước để nội suy Htk (chế độ bản cũ).")
        pairs = sorted((t["Q_tot"], t["H"]) for t in table)
        qs, hs = [p[0] for p in pairs], [p[1] for p in pairs]
        if Q > qs[-1] or Q < qs[0]:
            warns.append("Qtk nằm ngoài dải Q của các cấp H → Htk được NGOẠI SUY tuyến tính (kém tin cậy)")
        return _extrap_interp(Q, qs, hs), warns
    zs = [p[1] for p in P.pts]
    lo, hi = min(zs), max(zs)
    f = lambda H: evaluate(P, H)["Q_tot"] - Q
    if f(hi) < 0:
        top = hi
        while f(top) < 0 and top < hi + 30:
            top += 0.5
        if f(top) < 0:
            raise ValueError("Qtk quá lớn so với khả năng thoát của mặt cắt.")
        warns.append(f"Htk = {brentq(f, hi, top):.2f} m cao hơn điểm cao nhất của mặt cắt đo ({hi:.2f} m): "
                     "nước tràn ra ngoài phạm vi khảo sát, cần đo kéo dài mặt cắt")
        return brentq(f, hi, top), warns
    return brentq(f, lo, hi, xtol=1e-9), warns


def hq_curve(P: Params, n=70):
    zs = [p[1] for p in P.pts]
    hs = np.linspace(min(zs) + 0.02, max(zs), n)
    ev = [evaluate(P, float(h)) for h in hs]
    return hs, ev


def design_state(P: Params, Q: float, H: float) -> dict:
    ev = evaluate(P, H)
    z = ev["zones"]
    w1, wc, w2 = z["l"]["w"], z["m"]["w"], z["r"]["w"]
    b1, bc, b2 = z["l"]["b"], z["m"]["b"], z["r"]["b"]
    wb, bb = w1 + w2, b1 + b2
    Vc = z["m"]["Q"] / wc if wc > 0 else 0.0
    hc = wc / bc if bc > 0 else 0.0
    return dict(ev=ev, Q=Q, H=H, w1=w1, wc=wc, w2=w2, b1=b1, bc=bc, b2=b2, wb=wb, bb=bb, B=b1 + bc + b2,
                wtot=w1 + wc + w2, R1=z["l"]["R"], Rc=z["m"]["R"], R2=z["r"]["R"],
                Q1=z["l"]["Q"], Qc=z["m"]["Q"], Q2=z["r"]["Q"], Qb=z["l"]["Q"] + z["r"]["Q"],
                Vc=Vc, Vtb=Q / (w1 + wc + w2) if (w1 + wc + w2) > 0 else 0.0,
                hc=hc, h1=w1 / b1 if b1 > 0 else 0.0, h2=w2 / b2 if b2 > 0 else 0.0, hb=wb / bb if bb > 0 else 0.0,
                Fr=Vc / math.sqrt(G * hc) if hc > 0 else 0.0)


# ---- Bảng kết quả tab 4 / tab 5 ---------------------------------------------
@dataclass
class TRow:
    stt: str
    item: str
    unit: str = ""
    value: object = None
    note: str = ""
    kind: str = "row"      # sec | row | hl
    nd: int = 2


def rows_khau_do(S: dict, inv_n, mu, p_xoi, alpha_deg, htb):
    cos_a = math.cos(math.radians(alpha_deg))
    Q = S["Q"]
    k1 = S["Q1"] / Q if Q > 0 else 0.0
    k2 = S["Q2"] / Q if Q > 0 else 0.0
    den = mu * p_xoi * S["Vc"]
    omega = Q / den if den > 0 else 0.0
    L = omega / (htb * cos_a) if htb * cos_a > 0 else 0.0
    R = [
        TRow("I", "CÁC THÔNG SỐ LINK TỰ ĐỘNG TỪ MẶT CẮT TẠI Htk:", kind="sec"),
        TRow("1", "Lưu lượng thiết kế: QTK", "m³/s", Q, "", nd=2),
        TRow("2", "Diện tích ướt bãi trái: w1", "m²", S["w1"]),
        TRow("3", "Diện tích ướt bãi phải: w2", "m²", S["w2"], f"wb = {S['wb']:.2f} m²"),
        TRow("4", "Diện tích ướt lòng chủ: wc", "m²", S["wc"]),
        TRow("5", "Chiều rộng mặt nước bãi trái: b1", "m", S["b1"]),
        TRow("6", "Chiều rộng mặt nước bãi phải: b2", "m", S["b2"], f"B = {S['B']:.2f} m"),
        TRow("7", "Chiều rộng mặt nước lòng chủ: bc", "m", S["bc"]),
        TRow("8", "Bán kính thủy lực bãi trái: R1", "m", S["R1"], nd=3),
        TRow("9", "Bán kính thủy lực bãi phải: R2", "m", S["R2"], nd=3),
        TRow("10", "Bán kính thủy lực lòng chủ: Rc", "m", S["Rc"], nd=3),
        TRow("11", "Chiều sâu trung bình dòng bãi: hb = wb / (b1+b2)", "m", S["hb"], nd=3),
        TRow("12", "Hệ số nhám bãi trái: n1", "-", 1 / inv_n[0], "Quy đổi từ 1/n", nd=3),
        TRow("13", "Hệ số nhám bãi phải: n2", "-", 1 / inv_n[2], "Quy đổi từ 1/n", nd=3),
        TRow("14", "Hệ số nhám lòng chủ: nc", "-", 1 / inv_n[1], "Quy đổi từ 1/n", nd=3),
        TRow("15", "Hệ số phân phối lưu lượng bãi trái: k1 = Q1/Qtk", "-", k1, nd=4),
        TRow("16", "Hệ số phân phối lưu lượng bãi phải: k2 = Q2/Qtk", "-", k2, nd=4),
        TRow("17", "Lưu lượng dòng chủ: Qc", "m³/s", S["Qc"]),
        TRow("18", "Lưu tốc dòng chủ: Vc = Qc / wc", "m/s", S["Vc"]),
        TRow("II", "KẾT QUẢ TÍNH KHẨU ĐỘ THEO THAM SỐ CÔNG TRÌNH:", kind="sec"),
        TRow("1", "Hệ số thắt hẹp dòng chảy: m", "-", mu, "Người dùng nhập"),
        TRow("2", "Hệ số xói cho phép: P < Pmax", "-", p_xoi, "Người dùng nhập"),
        TRow("3", "Góc giữa dòng chảy và pháp tuyến tim cầu: α", "độ", alpha_deg, f"cosα = {cos_a:.4f}", nd=1),
        TRow("4", "Chiều sâu trung bình tính khẩu độ: htb", "m", htb, f"hc tại Htk = {S['hc']:.3f} m", nd=3),
        TRow("5", "DIỆN TÍCH THOÁT NƯỚC CẦN THIẾT: ωct = Qtk / (m·P·Vc)", "m²", omega, kind="hl"),
        TRow("6", "KHẨU ĐỘ THOÁT NƯỚC CẦN THIẾT TƯƠNG ĐƯƠNG: Lct = ωct / (htb·cosα)", "m", L, kind="hl"),
    ]
    return R, dict(omega=omega, L=L)


def rows_nuoc_denh(S: dict, slope, L_hl, n_piers, b_pier, eta, dz_manual=None):
    Q = S["Q"]
    H_cau = S["H"] + L_hl * slope
    w_dc = S["wb"] + S["wc"]
    w_tru = n_piers * b_pier * S["hc"]
    Vm = S["Vc"]
    Vo = Q / w_dc if w_dc > 0 else 0.0
    kb = S["Qb"] / Q * 100.0 if Q > 0 else 0.0
    dz_formula = eta * max(0.0, Vm ** 2 - Vo ** 2)
    dz = dz_manual if dz_manual is not None else dz_formula
    R = [
        TRow("I", "CÁC THÔNG SỐ LINK TỰ ĐỘNG TỪ MẶT CẮT TẠI Htk:", kind="sec"),
        TRow("1", "Lưu lượng thiết kế: QTK", "m³/s", Q),
        TRow("2", "Tổng lưu lượng qua hai bãi: ΣQb = Q1+Q2", "m³/s", S["Qb"]),
        TRow("3", "Mực nước thiết kế chuyển về tim cầu: HTK = Htk + Lhl·i", "m", H_cau, f"Hạ lưu cách tim {L_hl:g} m", nd=3),
        TRow("4", "Lưu tốc dòng chủ: Vc", "m/s", S["Vc"]),
        TRow("5", "Tổng diện tích ướt hai bãi: Swb = w1+w2", "m²", S["wb"]),
        TRow("6", "Diện tích ướt lòng chủ: wc", "m²", S["wc"]),
        TRow("7", "Tổng diện tích ướt thoát nước dưới cầu: Swdc", "m²", w_dc),
        TRow("8", "Diện tích chắn nước do trụ: wtrụ = số trụ·bt·hc", "m²", w_tru, f"{n_piers} trụ × {b_pier:g} m (chỉ hiển thị, không đưa vào ΔZ)"),
        TRow("9", "Tổng chiều rộng mặt nước: B", "m", S["B"]),
        TRow("10", "Chiều sâu trung bình bãi trái: h1 = w1 / b1", "m", S["h1"], nd=3),
        TRow("11", "Chiều sâu trung bình bãi phải: h2 = w2 / b2", "m", S["h2"], nd=3),
        TRow("12", "Chiều sâu trung bình lòng chủ: hc = wc / bc", "m", S["hc"], nd=3),
        TRow("13", "Gia tốc trọng trường: g", "m/s²", G),
        TRow("II", "KẾT QUẢ TÍNH ĐỘ DỀNH:", kind="sec"),
        TRow("1", "Lưu tốc trung bình dưới cầu trước xói: Vm", "m/s", Vm, "Vm = Vc"),
        TRow("2", "Lưu tốc trung bình mặt cắt tự nhiên: Vo", "m/s", Vo, "Vo = Qtk / Swdc"),
        TRow("3", "Tỉ số lưu lượng bãi / lưu lượng thiết kế: Kb", "%", kb, "Kb = ΣQb/Qtk·100"),
        TRow("4", "Hệ số η", "-", eta, "Người dùng nhập"),
        TRow("5", "ĐỘ DỀNH TRƯỚC CẦU: ΔZ = η·(Vm² − Vo²)", "m", dz,
             "NHẬP TAY (bỏ qua công thức)" if dz_manual is not None else f"theo công thức", kind="hl", nd=3),
        TRow("6", "MỰC NƯỚC DỀNH TRƯỚC CẦU: HTKdềnh = HTK + ΔZ", "m", H_cau + dz, kind="hl", nd=3),
    ]
    return R, dict(H_cau=H_cau, dz=dz, dz_formula=dz_formula, H_denh=H_cau + dz, Vo=Vo, Kb=kb)


def fmt_val(r: TRow) -> str:
    return "" if r.value is None or r.value == "" else (f"{r.value:.{r.nd}f}" if isinstance(r.value, (int, float)) else str(r.value))


# =============================================================================
# 2. XUẤT EXCEL (Manning là công thức sống; các bảng khác là số thật, không phải chuỗi)
# =============================================================================
def _safe_name(s):
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_") or "DuAn"


_FN = "Times New Roman"
_F_T, _F_B, _F_R = Font(name=_FN, size=13, bold=True), Font(name=_FN, size=11, bold=True), Font(name=_FN, size=11)
_F_H = Font(name=_FN, size=11, bold=True, color="FFFFFF")
_FILL_H = PatternFill("solid", start_color="203764", end_color="203764")
_FILL_IN = PatternFill("solid", start_color="FFFDE7", end_color="FFFDE7")
_FILL_TOT = PatternFill("solid", start_color="FFF3E0", end_color="FFF3E0")
_FILL_SEC = PatternFill("solid", start_color="E1F5FE", end_color="E1F5FE")
_S = Side(border_style="thin", color="000000")
_BOX = Border(left=_S, right=_S, top=_S, bottom=_S)
_CTR, _LFT, _RGT = (Alignment(horizontal=h, vertical="center", wrap_text=(h == "left")) for h in ("center", "left", "right"))


def _hdr(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row, j, h)
        c.font, c.fill, c.alignment, c.border = _F_H, _FILL_H, Alignment(horizontal="center", vertical="center", wrap_text=True), _BOX


def _tbl_sheet(ws, title, sub, rows):
    ws["A1"], ws["A1"].font = title, _F_T
    ws["A2"], ws["A2"].font = sub, _F_B
    _hdr(ws, 4, ["STT", "Hạng mục tính toán", "Đơn vị", "Trị số", "Ghi chú"])
    for i, r in enumerate(rows, start=5):
        vals = [r.stt, r.item, r.unit, r.value, r.note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(i, j, v)
            c.border, c.font = _BOX, (_F_B if r.kind != "row" else _F_R)
            c.alignment = _CTR if j in (1, 3) else (_RGT if j == 4 else _LFT)
            if r.kind == "sec":
                c.fill = _FILL_SEC
            elif r.kind == "hl":
                c.fill = _FILL_TOT
        if isinstance(r.value, (int, float)):
            ws.cell(i, 4).number_format = "0." + "0" * r.nd if r.nd else "0"
    for col, w in zip("ABCDE", (7, 70, 10, 14, 44)):
        ws.column_dimensions[col].width = w


def export_workbook(path, res):
    P, wb = res["params"], openpyxl.Workbook()
    name, L_hl = res["bridge"], res["L_hl"]

    # -- Đầu vào
    ws = wb.active
    ws.title = "Dau vao"
    ws["A1"], ws["A1"].font = f"THÔNG SỐ ĐẦU VÀO - {name}", _F_T
    items = [("Mép bờ trái dòng chủ X (m)", P.x_left), ("Mép bờ phải dòng chủ X (m)", P.x_right),
             ("1/n bãi trái", P.inv_n[0]), ("1/n dòng chủ", P.inv_n[1]), ("1/n bãi phải", P.inv_n[2]),
             ("Độ dốc mặt nước i", P.slope), ("Lưu lượng thiết kế Qtk (m³/s)", res["Q"]),
             ("Chế độ hình học", "Tương thích bản cũ" if P.legacy else "Chính xác (có giao điểm mực nước)")]
    for i, (k, v) in enumerate(items, start=3):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.font, b.font, a.border, b.border = _F_R, _F_B, _BOX, _BOX
        if not isinstance(v, str):
            b.fill = _FILL_IN
    r0 = 3 + len(items) + 1
    _hdr(ws, r0, ["STT", "X (m)", "Z (m)"])
    for k, (x, z) in enumerate(P.sorted_pts(), start=1):
        for j, v in enumerate((k, x, z), start=1):
            c = ws.cell(r0 + k, j, v)
            c.border, c.font = _BOX, _F_R
            if j > 1:
                c.number_format, c.fill = "0.00", _FILL_IN
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 36, 22, 12

    # -- Tổng hợp H-Q, H-V (R, V, Q là công thức)
    ws = wb.create_sheet("T.hop")
    ws["A1"], ws["A1"].font = f"BẢNG TỔNG HỢP QUAN HỆ H-Q VÀ H-V - {name}", _F_T
    for j, (lab, v) in enumerate((("1/n bãi trái", P.inv_n[0]), ("1/n dòng chủ", P.inv_n[1]),
                                  ("1/n bãi phải", P.inv_n[2]), ("Độ dốc i", P.slope))):
        ws.cell(2, 1 + j, lab).font = _F_B
        c = ws.cell(3, 1 + j, v)
        c.fill, c.font, c.border, c.alignment = _FILL_IN, _F_B, _BOX, _CTR
    heads = ["H (m)", "ω bãi trái", "χ bãi trái", "R bãi trái", "Q bãi trái", "ω dòng chủ", "χ dòng chủ", "R dòng chủ",
             "Vc (m/s)", "Q dòng chủ", "ω bãi phải", "χ bãi phải", "R bãi phải", "Q bãi phải", "Σω (m²)", "ΣQ (m³/s)", "Vtb (m/s)"]
    _hdr(ws, 5, heads)
    for i, e in enumerate(res["table"], start=6):
        z = e["zones"]
        f = {1: e["H"], 2: z["l"]["w"], 3: z["l"]["c"], 4: f"=IF(C{i}>0,B{i}/C{i},0)", 5: f"=B{i}*$A$3*D{i}^(2/3)*SQRT($D$3)",
             6: z["m"]["w"], 7: z["m"]["c"], 8: f"=IF(G{i}>0,F{i}/G{i},0)", 9: f"=IF(H{i}>0,$B$3*H{i}^(2/3)*SQRT($D$3),0)", 10: f"=F{i}*I{i}",
             11: z["r"]["w"], 12: z["r"]["c"], 13: f"=IF(L{i}>0,K{i}/L{i},0)", 14: f"=K{i}*$C$3*M{i}^(2/3)*SQRT($D$3)",
             15: f"=B{i}+F{i}+K{i}", 16: f"=E{i}+J{i}+N{i}", 17: f"=IF(O{i}>0,P{i}/O{i},0)"}
        for j, v in f.items():
            c = ws.cell(i, j, v)
            c.font, c.border, c.alignment, c.number_format = _F_R, _BOX, (_CTR if j == 1 else _RGT), "0.00" if j != 4 and j != 8 and j != 13 else "0.000"
    for j in range(1, 18):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.freeze_panes = "B6"
    ws.cell(7 + len(res["table"]), 1, "ω, χ: từ hình học mặt cắt (số liệu); R, V, Q: công thức Manning Q = ω·(1/n)·R^(2/3)·√i").font = Font(name=_FN, size=10, italic=True)

    # -- Chi tiết từng cấp
    for k, e in enumerate(res["table"], start=1):
        w = wb.create_sheet(f"H{k}")
        w["A1"], w["A1"].font = f"MẶT CẮT LƯU LƯỢNG TẠI HẠ LƯU CÁCH TIM {L_hl:g}M", _F_T
        w["A2"], w["A2"].font = f"BẢNG TÍNH DIỆN TÍCH THOÁT NƯỚC & CHU VI ƯỚT {name} - CẤP H{k} = {e['H']:.2f} m", _F_B
        _hdr(w, 4, ["Bộ phận", "X (m)", "Cao độ TN (m)", "Độ sâu h (m)", "K/cách lẻ b (m)", "Diện tích ωi (m²)", "Δh (m)", "Chu vi ướt χi (m)"])
        for i, r in enumerate(e["rows"], start=5):
            for j, key in enumerate(("bophan", "x", "z", "h", "b", "w", "dh", "c"), start=1):
                c = w.cell(i, j, r[key])
                c.border, c.font = _BOX, (_F_B if r["is_tot"] else _F_R)
                c.alignment = _LFT if j == 1 else _RGT
                if j > 1:
                    c.number_format = "0.00"
                if r["is_tot"]:
                    c.fill = _FILL_TOT
        w.column_dimensions["A"].width = 24
        for j in range(2, 9):
            w.column_dimensions[get_column_letter(j)].width = 15

    _tbl_sheet(wb.create_sheet("Dien tich thoat nuoc"), "BẢNG TÍNH DIỆN TÍCH THOÁT NƯỚC CẦN THIẾT CẦU",
               f"{name} - Htk = {res['H_tk']:.3f} m", res["rows4"])
    _tbl_sheet(wb.create_sheet("H1%Denh"), "BẢNG TÍNH MỰC NƯỚC DỀNH TRƯỚC CẦU", f"{name} - (ỨNG VỚI LŨ THIẾT KẾ)", res["rows5"])
    if res["warnings"]:
        w = wb.create_sheet("Canh bao")
        w["A1"], w["A1"].font = "CẢNH BÁO KHI TÍNH", _F_T
        for i, t in enumerate(res["warnings"], start=3):
            w.cell(i, 1, t).font = Font(name=_FN, size=11, color="B71C1C")
        w.column_dimensions["A"].width = 130
    wb.save(path)


# =============================================================================
# 3. BIỂU ĐỒ (Figure thuần - không dùng pyplot)
# =============================================================================
def draw_figure(fig, res):
    P, S = res["params"], res["state"]
    pts = P.sorted_pts()
    xs, zs = [p[0] for p in pts], [p[1] for p in pts]
    H_tk, Q = res["H_tk"], res["Q"]
    fig.clear()
    ax1 = fig.add_subplot(2, 1, 1)
    ax2 = fig.add_subplot(2, 2, 3)
    ax3 = fig.add_subplot(2, 2, 4)

    # (1) mặt cắt
    ax1.fill_between(xs, zs, min(zs) - 0.6, color="#e0d0b0", alpha=0.6)
    ax1.plot(xs, zs, "k-o", lw=1.8, ms=3, label="Địa hình mặt cắt")
    grid = np.linspace(xs[0], xs[-1], 600)
    zg = np.interp(grid, xs, zs)
    ax1.fill_between(grid, zg, H_tk, where=zg < H_tk, color="#42a5f5", alpha=0.45, label=f"Nước tại Htk = {H_tk:.2f} m")
    ax1.axhline(H_tk, color="#c62828", lw=1.6)
    for e in res["table"]:
        ax1.axhline(e["H"], color="#90caf9", lw=0.7, ls="--")
    for x, lab in ((P.x_left, "Mép trái dòng chủ"), (P.x_right, "Mép phải dòng chủ")):
        ax1.axvline(x, color="#b71c1c", ls=":", lw=1.3)
        ax1.text(x, min(zs) - 0.55, lab, rotation=90, va="bottom", ha="right", fontsize=8, color="#b71c1c")
    ax1.set_title(f"MẶT CẮT SÔNG - {res['bridge']}", fontweight="bold", fontsize=11)
    ax1.set_xlabel("Khoảng cách X (m)")
    ax1.set_ylabel("Cao trình Z (m)")
    ax1.grid(True, ls=":", alpha=0.6)
    ax1.legend(loc="lower right", fontsize=8)

    hs, ev = res["curve"]
    q = [e["Q_tot"] for e in ev]
    ax2.plot(q, hs, color="#1565c0", lw=2, label="H-Q (liên tục)")
    ax2.scatter([e["Q_tot"] for e in res["table"]], [e["H"] for e in res["table"]], color="#c62828", s=26, zorder=5, label="Các cấp H tính")
    ax2.scatter([Q], [H_tk], color="#2e7d32", s=70, zorder=6, label=f"Qtk={Q:.1f} → Htk={H_tk:.2f}")
    ax2.plot([0, Q, Q], [H_tk, H_tk, min(hs)], color="#2e7d32", ls="--", lw=1)
    ax2.set_xlabel("Q (m³/s)"); ax2.set_ylabel("H (m)"); ax2.set_title("QUAN HỆ H - Q", fontweight="bold", fontsize=10)
    ax2.grid(True, ls=":", alpha=0.6); ax2.legend(fontsize=7.5, loc="lower right")

    ax3.plot([e["zones"]["m"]["V"] for e in ev], hs, color="#d32f2f", lw=2, label="Vc dòng chủ")
    ax3.plot([e["V_tb"] for e in ev], hs, color="#2e7d32", lw=2, ls="--", label="Vtb toàn mặt cắt")
    ax3.scatter([e["zones"]["m"]["V"] for e in res["table"]], [e["H"] for e in res["table"]], color="#d32f2f", s=22, zorder=5)
    ax3.axhline(H_tk, color="#1565c0", ls=":", lw=1.4, label=f"Htk = {H_tk:.2f} m (Vc={S['Vc']:.2f}, Vtb={S['Vtb']:.2f})")
    ax3.set_xlabel("V (m/s)"); ax3.set_ylabel("H (m)"); ax3.set_title("QUAN HỆ H - V", fontweight="bold", fontsize=10)
    ax3.grid(True, ls=":", alpha=0.6); ax3.legend(fontsize=7.5, loc="lower right")
    fig.tight_layout()


# =============================================================================
# 4. GIAO DIỆN TKINTER
# =============================================================================
MANUAL_BG, AUTO_BG = "#fffde7", "#e1f5fe"
DEFAULTS = dict(bridge_name="CẦU HỮU NGHỊ", x_left="", x_right="", n_left="5.0", n_main="20.0", n_right="5.0",
                slope_i="0.00107", h_max="", h_min="", dh="", q_tk="159.39",
                dt_mu="0.98", dt_p="1.04", dt_alpha="33.0", dt_htb="2.885",
                denh_lhl="25.0", denh_piers="0", denh_bpier="1.4", denh_eta="0.07", dz_manual="",
                legacy=False, auto_calc=True)


class CrossSectionApp:
    def __init__(self, root):
        self.root = root
        root.title(APP_TITLE)
        root.geometry("1180x780")
        root.minsize(980, 640)
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.pts = list(DEFAULT_MC)
        self.res = None
        self.level_vals = []
        self._after = None
        self._internal = False
        self.v = {}
        for k, val in DEFAULTS.items():
            var = tk.BooleanVar(value=val) if isinstance(val, bool) else tk.StringVar(value=val)
            var.trace_add("write", self._changed)
            self.v[k] = var
        self.v["level_sel"] = tk.StringVar(value="")
        self._build_ui()
        self.auto_detect(silent=True)

    # ------------------------------------------------------------ helpers
    def _changed(self, *_):
        if self._internal or not self.v["auto_calc"].get():
            return
        if self._after:
            self.root.after_cancel(self._after)
        self._after = self.root.after(250, self.calculate)

    def _set(self, key, val):
        self._internal = True
        try:
            self.v[key].set(val)
        finally:
            self._internal = False

    def num(self, key, name=None):
        return parse_float(self.v[key].get(), name or key)

    def show_status(self, msg, level="ok"):
        bg, fg = {"ok": ("#e8f5e9", "#1b5e20"), "warn": ("#fff8e1", "#e65100"), "error": ("#ffebee", "#b71c1c")}[level]
        self.status.config(text=msg, bg=bg, fg=fg)

    def _entry(self, parent, key, width=9, auto=False):
        return tk.Entry(parent, textvariable=self.v[key], width=width, justify="center", bg=AUTO_BG if auto else MANUAL_BG,
                        fg="#0277bd" if auto else "#b71c1c", font=("Arial", 10, "bold"))

    def _tree(self, parent, cols, heads, widths, height=12, anchors=None):
        tv = ttk.Treeview(parent, columns=cols, show="headings", height=height)
        for i, (c, h, w) in enumerate(zip(cols, heads, widths)):
            tv.heading(c, text=h)
            tv.column(c, width=w, anchor=(anchors[i] if anchors else "e"))
        tv.tag_configure("tot", background="#fff3e0", font=("Arial", 9, "bold"))
        tv.tag_configure("sec", background="#e1f5fe", foreground="#0277bd", font=("Arial", 9, "bold"))
        tv.tag_configure("hl", background="#fff3e0", foreground="#b71c1c", font=("Arial", 9, "bold"))
        return tv

    # --------------------------------------------------------------- UI
    def _build_ui(self):
        top = ttk.Frame(self.root, padding=(10, 6))
        top.pack(fill="x")
        ttk.Label(top, text="Tên cầu / công trình:", font=("Arial", 10, "bold")).pack(side="left")
        tk.Entry(top, textvariable=self.v["bridge_name"], width=30, font=("Arial", 10, "bold"), fg="#b71c1c").pack(side="left", padx=8)
        ttk.Button(top, text="📂 Mở dự án…", command=self.open_project).pack(side="right", padx=3)
        ttk.Button(top, text="💾 Lưu dự án…", command=self.save_project).pack(side="right", padx=3)

        self.status = tk.Label(self.root, anchor="w", justify="left", wraplength=1140, padx=8, pady=3, relief="sunken")
        self.status.pack(fill="x", side="bottom")
        bar = ttk.Frame(self.root, padding=(10, 6))
        bar.pack(fill="x", side="bottom")
        tk.Button(bar, text="🚀 TÍNH TOÁN TOÀN BỘ", font=("Arial", 10, "bold"), bg="#154360", fg="white", padx=16, pady=4,
                  command=self.run_clicked).pack(side="left")
        ttk.Checkbutton(bar, text="Tự động tính lại khi sửa số liệu", variable=self.v["auto_calc"]).pack(side="left", padx=12)
        ttk.Button(bar, text="📊 Xuất Excel", command=self.export_excel).pack(side="right", padx=3)
        ttk.Button(bar, text="🖼 Xuất ảnh biểu đồ", command=self.export_png).pack(side="right", padx=3)

        nb = self.nb = ttk.Notebook(self.root)
        nb.pack(fill="both", expand=True, padx=8, pady=2)
        self.tabs = [ttk.Frame(nb) for _ in range(7)]
        for t, name in zip(self.tabs, ["1. Mặt cắt & thông số", "2. Bảng H-Q, H-V", "3. Mực nước thiết kế Htk",
                                       "4. Diện tích thoát nước cầu", "5. Nước dềnh trước cầu", "6. Chi tiết từng cấp H", "7. Biểu đồ"]):
            nb.add(t, text=f"  {name}  ")
        self._tab_input(self.tabs[0])
        self._tab_summary(self.tabs[1])
        self._tab_htk(self.tabs[2])
        self._tab_bridge(self.tabs[3], 4)
        self._tab_bridge(self.tabs[4], 5)
        self._tab_detail(self.tabs[5])
        self._tab_plot(self.tabs[6])

    def _tab_input(self, t):
        left = ttk.LabelFrame(t, text=" Tọa độ mặt cắt (X: khoảng cách cộng dồn, Z: cao độ đáy) ", padding=6)
        left.pack(side="left", fill="both", expand=True, padx=(8, 4), pady=6)
        self.tree_mc = self._tree(left, ("stt", "x", "z"), ("STT", "X (m)", "Z (m)"), (50, 110, 110), 14, ("center", "e", "e"))
        self.tree_mc.pack(fill="both", expand=True)
        self.tree_mc.bind("<<TreeviewSelect>>", self.on_pick_point)
        r = ttk.Frame(left); r.pack(fill="x", pady=4)
        self.ent_x, self.ent_z = tk.Entry(r, width=8), tk.Entry(r, width=8)
        ttk.Label(r, text="X:").pack(side="left"); self.ent_x.pack(side="left", padx=2)
        ttk.Label(r, text="Z:").pack(side="left"); self.ent_z.pack(side="left", padx=2)
        ttk.Button(r, text="➕ Thêm", command=self.add_point).pack(side="left", padx=2)
        ttk.Button(r, text="✏ Cập nhật", command=self.update_point).pack(side="left", padx=2)
        ttk.Button(r, text="❌ Xóa", command=self.delete_point).pack(side="left", padx=2)
        r2 = ttk.Frame(left); r2.pack(fill="x")
        ttk.Button(r2, text="📋 Dán từ Excel (clipboard)", command=self.paste_points).pack(side="left", padx=2)
        ttk.Button(r2, text="📂 Nạp Excel/TXT/CSV", command=self.load_points_file).pack(side="left", padx=2)

        right = ttk.LabelFrame(t, text=" Thông số thủy lực (vàng: nhập tay, xanh: tự nhận diện) ", padding=10)
        right.pack(side="right", fill="both", expand=True, padx=(4, 8), pady=6)
        g = ttk.Frame(right); g.pack(fill="x")
        rows = [("Mép bờ trái dòng chủ X (m)", "x_left", True), ("Mép bờ phải dòng chủ X (m)", "x_right", True),
                ("1/n bãi trái", "n_left", False), ("1/n dòng chủ", "n_main", False), ("1/n bãi phải", "n_right", False),
                ("Độ dốc mặt nước lũ i", "slope_i", False), ("Lưu lượng thiết kế Qtk (m³/s)", "q_tk", False),
                ("H max của dải cấp H (m)", "h_max", True), ("H min (m)", "h_min", True), ("Bước dH (m)", "dh", True)]
        for i, (lab, key, auto) in enumerate(rows):
            ttk.Label(g, text=lab).grid(row=i, column=0, sticky="w", pady=3)
            self._entry(g, key, 10, auto).grid(row=i, column=1, padx=8)
        ttk.Button(right, text="🔄 Tự động nhận diện bờ & dải H", command=lambda: self.auto_detect(False)).pack(anchor="w", pady=(8, 2))
        ttk.Separator(right).pack(fill="x", pady=8)
        ttk.Checkbutton(right, text="Tương thích bản cũ (hình học & nội suy Htk như bản cũ - chỉ để đối chiếu)",
                        variable=self.v["legacy"]).pack(anchor="w")
        ttk.Label(right, text="Mặc định: hình học chính xác (cắt đúng giao điểm mực nước với đáy sông) và giải Htk\n"
                              "trực tiếp từ Q(H) = Qtk thay vì nội suy trên vài cấp H.", foreground="#555").pack(anchor="w", pady=4)

    def _tab_summary(self, t):
        cols = ("h", "wl", "cl", "ql", "wc", "cc", "vc", "qc", "wr", "cr", "qr", "wt", "vtb", "qt")
        heads = ("H (m)", "ω bãi trái", "χ bãi trái", "Q bãi trái", "ω dòng chủ", "χ dòng chủ", "Vc (m/s)", "Q dòng chủ",
                 "ω bãi phải", "χ bãi phải", "Q bãi phải", "Σω (m²)", "Vtb (m/s)", "ΣQ (m³/s)")
        self.tree2 = self._tree(t, cols, heads, [70] + [80] * 13, 16)
        self.tree2.pack(fill="both", expand=True, padx=6, pady=6)

    def _tab_htk(self, t):
        f = ttk.LabelFrame(t, text=" Lưu lượng thiết kế → mực nước thiết kế ", padding=8)
        f.pack(fill="x", padx=8, pady=6)
        r = ttk.Frame(f); r.pack(fill="x")
        ttk.Label(r, text="Qtk (m³/s):", font=("Arial", 10, "bold")).pack(side="left")
        self._entry(r, "q_tk", 10).pack(side="left", padx=8)
        self.lbl_card = tk.Label(f, text="", font=("Arial", 10, "bold"), fg="#0d47a1", bg="#e3f2fd", justify="left", padx=10, pady=6, relief="groove")
        self.lbl_card.pack(fill="x", pady=6)
        self.tree3 = self._detail_tree(t)
        self.tree3.pack(fill="both", expand=True, padx=8, pady=4)

    def _detail_tree(self, parent):
        return self._tree(parent, ("bp", "x", "z", "h", "b", "w", "dh", "c"),
                          ("Bộ phận", "X (m)", "Cao độ TN (m)", "Độ sâu h (m)", "K/cách lẻ b (m)", "Diện tích ωi (m²)", "Δh (m)", "Chu vi χi (m)"),
                          (150, 80, 100, 100, 110, 120, 80, 110), 12, ("w",) + ("e",) * 7)

    def _tab_bridge(self, t, which):
        f = ttk.LabelFrame(t, text=" Tham số công trình cầu (người dùng nhập) ", padding=8)
        f.pack(fill="x", padx=8, pady=6)
        r = ttk.Frame(f); r.pack(fill="x")
        spec = ([("Hệ số thắt hẹp m", "dt_mu"), ("Hệ số xói P", "dt_p"), ("Góc lệch dòng α (°)", "dt_alpha"), ("htb (m)", "dt_htb")] if which == 4 else
                [("Cự ly tim cầu Lhl (m)", "denh_lhl"), ("Số trụ", "denh_piers"), ("Bề rộng trụ bt (m)", "denh_bpier"),
                 ("η", "denh_eta"), ("ΔZ nhập tay (để trống = dùng công thức)", "dz_manual")])
        for lab, key in spec:
            ttk.Label(r, text=lab).pack(side="left", padx=(0, 2))
            self._entry(r, key, 8).pack(side="left", padx=(0, 10))
        if which == 4:
            ttk.Button(r, text="Lấy htb = hc tại Htk", command=self.htb_from_hc).pack(side="left")
        tv = self._tree(t, ("stt", "item", "unit", "val", "note"), ("STT", "Hạng mục tính toán", "Đơn vị", "Trị số", "Ghi chú"),
                        (45, 520, 70, 110, 300), 16, ("center", "w", "center", "e", "w"))
        tv.pack(fill="both", expand=True, padx=8, pady=4)
        setattr(self, f"tree{which}", tv)

    def _tab_detail(self, t):
        r = ttk.Frame(t); r.pack(fill="x", pady=6, padx=8)
        ttk.Label(r, text="Chọn cấp mực nước:").pack(side="left")
        self.cbo_level = ttk.Combobox(r, textvariable=self.v["level_sel"], state="readonly", width=24)
        self.cbo_level.pack(side="left", padx=6)
        self.cbo_level.bind("<<ComboboxSelected>>", lambda e: self.show_level())
        self.lbl_level = ttk.Label(r, text="", font=("Arial", 10, "bold"), foreground="#c62828")
        self.lbl_level.pack(side="left", padx=10)
        self.tree6 = self._detail_tree(t)
        self.tree6.pack(fill="both", expand=True, padx=8, pady=4)

    def _tab_plot(self, t):
        if not HAS_MPL:
            ttk.Label(t, text="Cần cài matplotlib để xem biểu đồ.").pack(pady=20)
            self.fig = self.canvas = None
            return
        self.fig = Figure(figsize=(11, 7), dpi=100)
        self.canvas = FigureCanvasTkAgg(self.fig, master=t)
        NavigationToolbar2Tk(self.canvas, t).update()
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

    # ------------------------------------------------------- điểm mặt cắt
    def refresh_points(self):
        self.pts.sort(key=lambda p: p[0])
        self.tree_mc.delete(*self.tree_mc.get_children())
        for i, (x, z) in enumerate(self.pts, start=1):
            self.tree_mc.insert("", "end", values=(i, f"{x:.2f}", f"{z:.2f}"))
        self._changed()

    def on_pick_point(self, _=None):
        sel = self.tree_mc.selection()
        if sel:
            i = int(self.tree_mc.item(sel[0])["values"][0]) - 1
            self.ent_x.delete(0, "end"); self.ent_x.insert(0, f"{self.pts[i][0]:g}")
            self.ent_z.delete(0, "end"); self.ent_z.insert(0, f"{self.pts[i][1]:g}")

    def _xz(self):
        return parse_float(self.ent_x.get(), "X"), parse_float(self.ent_z.get(), "Z")

    def add_point(self):
        try:
            self.pts.append(self._xz())
        except ValueError as e:
            return self.show_status(f"✖ {e}", "error")
        self.refresh_points()

    def update_point(self):
        sel = self.tree_mc.selection()
        if not sel:
            return self.show_status("Chọn một điểm trong bảng trước khi cập nhật.", "warn")
        try:
            self.pts[int(self.tree_mc.item(sel[0])["values"][0]) - 1] = self._xz()
        except ValueError as e:
            return self.show_status(f"✖ {e}", "error")
        self.refresh_points()

    def delete_point(self):
        sel = self.tree_mc.selection()
        if sel:
            del self.pts[int(self.tree_mc.item(sel[0])["values"][0]) - 1]
            self.refresh_points()

    def _adopt_points(self, pts, src):
        if len(pts) < 3:
            return messagebox.showwarning("Cảnh báo", "Cần ít nhất 3 điểm (X Z).")
        self.pts = pts
        self.refresh_points()
        self.auto_detect(silent=True)
        self.show_status(f"✔ Đã nạp {len(pts)} điểm từ {src}; đã nhận diện lại bờ và dải H.", "ok")

    def paste_points(self):
        try:
            self._adopt_points(parse_xz_text(self.root.clipboard_get()), "clipboard")
        except Exception as e:
            self.show_status(f"✖ Không đọc được clipboard: {e}", "error")

    def load_points_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel/Text/CSV", "*.xlsx *.xls *.txt *.csv"), ("Tất cả", "*.*")])
        if not path:
            return
        try:
            if path.lower().endswith((".xls", ".xlsx")):
                import pandas as pd
                df = pd.read_excel(path, sheet_name=0)
                cols = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])]
                low = {c: str(c).lower() for c in cols}
                cx = next((c for c in cols if re.search(r"\bx\b|khoảng|kc|cộng dồn|distance", low[c])), None)
                cz = next((c for c in cols if re.search(r"\bz\b|cao|đáy|elev", low[c]) and c != cx), None)
                if cx is None or cz is None:
                    cx, cz = (cols[0], cols[1]) if len(cols) >= 2 else (None, None)
                if cx is None:
                    return messagebox.showwarning("Cảnh báo", "Không tìm thấy 2 cột số (X, Z).")
                sub = df[[cx, cz]].dropna()
                pts = list(zip(sub.iloc[:, 0].astype(float), sub.iloc[:, 1].astype(float)))
                src = f"{path} (cột X = «{cx}», Z = «{cz}»)"
            else:
                with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                    pts = parse_xz_text(fh.read())
                src = path
            self._adopt_points(pts, src)
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không đọc được tệp mặt cắt: {e}")

    def auto_detect(self, silent=True):
        if len(self.pts) < 3:
            return
        xl, xr = detect_banks(self.pts)
        hmax, hmin, dh = default_levels(self.pts)
        self._internal = True
        try:
            for k, val in (("x_left", f"{xl:.2f}"), ("x_right", f"{xr:.2f}"), ("h_max", f"{hmax:.2f}"),
                           ("h_min", f"{hmin:.2f}"), ("dh", f"{dh:.2f}")):
                self.v[k].set(val)
        finally:
            self._internal = False
        self.calculate()
        if not silent:
            self.show_status(f"✔ Bờ trái X = {xl:.2f}, bờ phải X = {xr:.2f}; dải H: {hmax:.2f} → {hmin:.2f}, bước {dh:.2f} m.", "ok")

    # --------------------------------------------------------------- tính
    def run_clicked(self):
        self.calculate()
        if self.res:
            self.nb.select(self.tabs[1])

    def htb_from_hc(self):
        if self.res:
            self._set("dt_htb", f"{self.res['state']['hc']:.3f}")
            self.calculate()

    def read_params(self):
        return Params(pts=list(self.pts), x_left=self.num("x_left", "Mép bờ trái"), x_right=self.num("x_right", "Mép bờ phải"),
                      inv_n=(self.num("n_left", "1/n bãi trái"), self.num("n_main", "1/n dòng chủ"), self.num("n_right", "1/n bãi phải")),
                      slope=self.num("slope_i", "i"), legacy=self.v["legacy"].get())

    def calculate(self):
        if self._after:
            self.root.after_cancel(self._after)
            self._after = None
        self.res = None
        try:
            P = self.read_params()
            Q = self.num("q_tk", "Qtk")
            levels = make_levels(self.num("h_max", "Hmax"), self.num("h_min", "Hmin"), self.num("dh", "dH"))
            errs, warns = P.validate()
            if Q <= 0:
                errs.append("Qtk phải > 0")
            if errs:
                return self.show_status("✖ " + "; ".join(errs), "error")
            table = [evaluate(P, H) for H in levels]
            H_tk, w2 = solve_H(P, Q, table)
            warns += w2
            S = design_state(P, Q, H_tk)
            rows4, k4 = rows_khau_do(S, P.inv_n, self.num("dt_mu", "m"), self.num("dt_p", "P"), self.num("dt_alpha", "α"), self.num("dt_htb", "htb"))
            dz = self.v["dz_manual"].get().strip()
            rows5, k5 = rows_nuoc_denh(S, P.slope, self.num("denh_lhl", "Lhl"), int(self.num("denh_piers", "số trụ")),
                                       self.num("denh_bpier", "bt"), self.num("denh_eta", "η"), parse_float(dz, "ΔZ") if dz else None)
        except (ValueError, ZeroDivisionError) as e:
            return self.show_status(f"✖ {e}", "error")
        if S["Fr"] >= 1:
            warns.append(f"Fr = {S['Fr']:.2f} ≥ 1 ở dòng chủ: dòng siết, giả thiết Manning chảy đều cần cân nhắc")
        if S["wc"] <= 0:
            warns.append("Tại Htk dòng chủ không có nước - kiểm tra vị trí mép bờ / cao độ")
        if S["Qb"] <= 0:
            warns.append("Tại Htk hai bãi chưa ngập (ΣQb = 0)")
        if P.legacy:
            warns.append("Đang ở chế độ TƯƠNG THÍCH BẢN CŨ: hình học không cắt giao điểm mực nước; Htk nội suy tuyến tính trên các cấp H")
        if dz:
            warns.append("ΔZ đang được NHẬP TAY, không dùng công thức η(Vm²−Vo²)")
        self.res = dict(params=P, Q=Q, table=table, levels=levels, H_tk=H_tk, state=S, rows4=rows4, rows5=rows5, k4=k4, k5=k5,
                        curve=hq_curve(P), bridge=self.v["bridge_name"].get().strip().upper() or "CẦU",
                        L_hl=self.num("denh_lhl", "Lhl"), warnings=warns)
        self._render()
        self.show_status("⚠ " + " | ".join(warns) if warns else "✔ Đã tính xong, không có cảnh báo.", "warn" if warns else "ok")

    def _render(self):
        r, S = self.res, self.res["state"]
        self.tree2.delete(*self.tree2.get_children())
        for e in r["table"]:
            z = e["zones"]
            self.tree2.insert("", "end", values=(f"{e['H']:.2f}", f"{z['l']['w']:.2f}", f"{z['l']['c']:.2f}", f"{z['l']['Q']:.2f}",
                                                 f"{z['m']['w']:.2f}", f"{z['m']['c']:.2f}", f"{z['m']['V']:.3f}", f"{z['m']['Q']:.2f}",
                                                 f"{z['r']['w']:.2f}", f"{z['r']['c']:.2f}", f"{z['r']['Q']:.2f}",
                                                 f"{e['w_tot']:.2f}", f"{e['V_tb']:.3f}", f"{e['Q_tot']:.2f}"))
        self.lbl_card.config(text=(f"🎯 {r['bridge']}:  Qtk = {r['Q']:.2f} m³/s  →  Mực nước thiết kế Htk = {r['H_tk']:.3f} m\n"
                                   f"Tổng diện tích ướt ωtk = {S['wtot']:.2f} m²   |   Vc = {S['Vc']:.3f} m/s   |   Vtb = {S['Vtb']:.3f} m/s   |   "
                                   f"Fr(dòng chủ) = {S['Fr']:.2f}"))
        self._fill_detail(self.tree3, S["ev"]["rows"])
        for tv, rows in ((self.tree4, r["rows4"]), (self.tree5, r["rows5"])):
            tv.delete(*tv.get_children())
            for x in rows:
                tv.insert("", "end", tags=(x.kind,) if x.kind != "row" else (),
                          values=(x.stt, x.item, x.unit, fmt_val(x), "" if x.kind != "row" and not x.note else x.note))
        self.level_vals = [f"Cấp H{i + 1}: {e['H']:.2f} m" for i, e in enumerate(r["table"])]
        self.cbo_level["values"] = self.level_vals
        if self.v["level_sel"].get() not in self.level_vals:
            self._set("level_sel", self.level_vals[0])
        self.show_level()
        if self.fig is not None:
            draw_figure(self.fig, r)
            self.canvas.draw_idle()

    @staticmethod
    def _fill_detail(tv, rows):
        tv.delete(*tv.get_children())
        f = lambda v: "" if v is None else f"{v:.2f}"
        for x in rows:
            tv.insert("", "end", tags=("tot",) if x["is_tot"] else (),
                      values=(x["bophan"], f(x["x"]), f(x["z"]), f(x["h"]), f(x["b"]), f(x["w"]), f(x["dh"]), f(x["c"])))

    def show_level(self):
        if not self.res or self.v["level_sel"].get() not in self.level_vals:
            return
        e = self.res["table"][self.level_vals.index(self.v["level_sel"].get())]
        self.lbl_level.config(text=f"H = {e['H']:.2f} m   |   ωtổng = {e['w_tot']:.2f} m²   |   Q = {e['Q_tot']:.2f} m³/s")
        self._fill_detail(self.tree6, e["rows"])

    # ------------------------------------------------------ xuất / lưu / mở
    def _ensure(self):
        if self._after:
            self.calculate()
        if not self.res:
            messagebox.showerror("Lỗi", "Số liệu đầu vào chưa hợp lệ - xem thanh trạng thái phía dưới.")
        return self.res

    def export_excel(self):
        res = self._ensure()
        if not res:
            return
        path = filedialog.asksaveasfilename(initialfile=f"HinhThaiMC_{_safe_name(res['bridge'])}.xlsx", defaultextension=".xlsx",
                                            filetypes=[("Excel", "*.xlsx")])
        if path:
            try:
                export_workbook(path, res)
                messagebox.showinfo("Thành công", f"Đã xuất Excel:\n{path}")
            except Exception as e:
                messagebox.showerror("Lỗi", f"Không xuất được Excel: {e}")

    def export_png(self):
        if not self._ensure() or self.fig is None:
            return
        path = filedialog.asksaveasfilename(defaultextension=".png", filetypes=[("PNG", "*.png"), ("PDF", "*.pdf")])
        if path:
            self.fig.savefig(path, dpi=200)
            self.show_status(f"✔ Đã lưu biểu đồ: {path}", "ok")

    KEYS = [k for k in DEFAULTS]

    def save_project(self):
        path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("Dự án JSON", "*.json")])
        if not path:
            return
        try:
            data = {"version": 2, "pts": [list(p) for p in self.pts], **{k: self.v[k].get() for k in self.KEYS}}
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(data, fh, ensure_ascii=False, indent=2)
            self.show_status(f"✔ Đã lưu dự án: {path}", "ok")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không lưu được: {e}")

    def open_project(self):
        path = filedialog.askopenfilename(filetypes=[("Dự án JSON", "*.json"), ("Tất cả", "*.*")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as fh:
                d = json.load(fh)
            pts = [tuple(p) for p in d.get("pts", d.get("mc_data", DEFAULT_MC))]     # đọc được cả file của bản cũ
            self._internal = True
            try:
                for k in self.KEYS:
                    if k in d:
                        self.v[k].set(d[k])
                if "q_tk" in d and "version" not in d:
                    self.v["legacy"].set(False)
            finally:
                self._internal = False
            self.pts = pts
            self.refresh_points()
            self.calculate()
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không mở được dự án: {e}")

    def on_close(self):
        if self._after:
            self.root.after_cancel(self._after)
        self.root.destroy()


def main():
    if not HAS_TK:
        raise SystemExit("Python này chưa có tkinter (Linux: sudo apt install python3-tk).")
    root = tk.Tk()
    CrossSectionApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
