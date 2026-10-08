"""Integration checks for the real Tk calculation and Word/Excel report writers.

Run with a working DISPLAY and the application's Python dependencies:
    python -m unittest discover -s tests -p test_scour_reports.py -v
"""
import importlib.util
import math
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

SOURCE = Path(__file__).resolve().parents[1] / 'TinhXoiCau' / 'TTXoiCau_v4.py'
spec = importlib.util.spec_from_file_location('scour_app', SOURCE)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class ScourReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.app = module.MainScourApplication()
        cls.app.withdraw()
        for pier in cls.app.project['piers_detail']:
            pier['loai_tru'] = ('Trụ đơn đặc' if pier['name'] == 'T25' else
                                'Lộ bệ' if pier['name'] in ('T26', 'T27') else 'Lộ bệ & cọc')
        # Deliberately differ from saved project metadata to detect stale report headers.
        cls.app.t1_entries['htk'].delete(0, 'end')
        cls.app.t1_entries['htk'].insert(0, '1.56')
        def fail(*args, **kwargs):
            raise AssertionError(args)
        with patch.object(module.messagebox, 'showinfo'), patch.object(module.messagebox, 'showerror', fail):
            cls.app.run_full_system()
        cls.xlsx = Path(cls.temp.name) / 'scour.xlsx'
        cls.word = Path(cls.temp.name) / 'scour.docx'
        cls.app.export_excel_report(cls.xlsx)
        cls.app.export_word_report(cls.word)
        cls.wb = load_workbook(cls.xlsx)
        cls.doc = Document(cls.word)

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        cls.app.destroy()
        cls.temp.cleanup()

    def test_live_input_header_and_complete_exports(self):
        self.assertEqual(len(self.wb.worksheets), 9)
        for ws in self.wb:
            self.assertIn('Htt = 1.56 m', ws['A3'].value)
            expected = 'portrait' if ws.title in ('XCB-lo coc', 'XCB-lo be') else 'landscape'
            self.assertEqual(ws.page_setup.orientation, expected)
            self.assertEqual(ws.page_setup.fitToWidth, 1)
            self.assertTrue(ws.print_area)
        self.assertEqual(len(self.doc.sections), 9)
        self.assertTrue(any('Htt = 1.56 m' in p.text for p in self.doc.paragraphs))

    def test_pile_group_has_all_four_tables_in_both_formats(self):
        ws = self.wb['XCB-lo coc']
        text = [c.value for row in ws for c in row if isinstance(c.value, str)]
        for title in ('1. Xói cục bộ do thân trụ gây ra', '2. Xói cục bộ do bệ trụ',
                      '3. Xói cục bộ do nhóm cọc', '4. Kết quả phân tích xói cục bộ tại trụ'):
            self.assertIn(title, text)
            self.assertTrue(any(p.text == title for p in self.doc.paragraphs))
        for name in ('T28', 'T29', 'T30'):
            self.assertEqual(sum(c.value == name for row in ws for c in row), 4)
        self.assertNotIn('T26', text)

    def test_classification_names_and_empty_case_are_preserved(self):
        names = [c.value for row in self.wb['Xoi cuc tru'] for c in row]
        self.assertIn('T25', names)
        self.assertNotIn('T28', names)
        lb = [c.value for row in self.wb['XCB-lo be'] for c in row]
        self.assertIn('T26', lb)
        self.assertIn('T27', lb)
        original = self.app._report_piers
        try:
            self.app._report_piers = []
            path = Path(self.temp.name) / 'empty.xlsx'
            self.app.export_excel_report(path)
            wb = load_workbook(path)
            self.assertTrue(any('Không có mố/trụ' in str(c.value)
                                for row in wb['XCB-lo coc'] for c in row))
            wb.close()
        finally:
            self.app._report_piers = original

    def test_summary_keeps_degradation_and_actual_totals(self):
        ws = self.wb['Tong hop']
        self.assertTrue(any(c.value == 'Hạ thấp dài hạn' for row in ws for c in row))
        source = {r['name']: r for r in self.app.scour_results}
        rows = [row for row in ws.iter_rows(values_only=True) if row[0] in source]
        self.assertEqual(len(rows), len(source))
        for row in rows:
            r = source[row[0]]
            self.assertAlmostEqual(row[-2], r['y_tot'], places=10)
            self.assertAlmostEqual(row[-1], r['z_scour'], places=10)
            self.assertAlmostEqual(row[-2], r['y_deg'] + r['ysc'] + r['ys_local'], places=10)
            self.assertAlmostEqual(row[-1], row[1] - row[-2], places=10)

    def test_small_grains_scientific_values_and_equations(self):
        ws = self.wb['Xói chung']
        data = next(row for row in ws.iter_rows() if row[0].value == 'Lòng sông')
        self.assertAlmostEqual(data[3].value, 0.025)
        self.assertAlmostEqual(data[8].value, 0.000005)
        self.assertEqual(data[3].number_format, '0.000')
        self.assertEqual(data[8].number_format, '0.000000')
        self.assertEqual(len(ws._images), 3)
        self.assertGreaterEqual(len(self.doc.inline_shapes), 15)

    def test_merged_headers_and_fixed_word_table_widths(self):
        self.assertGreater(len(self.wb['XCB-lo coc'].merged_cells.ranges), 20)
        grid_tables = [t for t in self.doc.tables if t.style.name == 'Table Grid']
        self.assertEqual(len(grid_tables), 14)
        for table in grid_tables:
            self.assertFalse(table.autofit)
            for row in table.rows[:3]:
                self.assertIsNotNone(row._tr.trPr.find(qn('w:tblHeader')))
            for row in table.rows:
                self.assertIsNotNone(row._tr.trPr.find(qn('w:cantSplit')))

    def test_formula_failure_reports_error_and_cleans_temporary_images(self):
        created = []
        original = self.app._create_equation_img
        def create(equation, **kwargs):
            if created:
                return None, 0, 0
            result = original(equation, **kwargs)
            created.append(result[0])
            return result
        with patch.object(self.app, '_create_equation_img', create):
            with self.assertRaisesRegex(ValueError, 'Không tạo được ảnh công thức'):
                self.app.export_excel_report(Path(self.temp.name) / 'failure.xlsx')
        self.assertTrue(created)
        self.assertFalse(Path(created[0]).exists())


if __name__ == '__main__':
    unittest.main()
