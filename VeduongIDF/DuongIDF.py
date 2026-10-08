import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import numpy as np
import pandas as pd
from scipy.stats import pearson3, norm

# Matplotlib nhúng vào Tkinter
import matplotlib
matplotlib.use("TkAgg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
import matplotlib.pyplot as plt

plt.rcParams['font.family'] = 'serif'
plt.rcParams['font.serif'] = ['Times New Roman', 'DejaVu Serif']

# OpenPyXL xử lý file Excel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text, RichText
from openpyxl.drawing.text import (
    Paragraph, ParagraphProperties, CharacterProperties,
    Font as DrawingFont, RegularTextRun
)
from openpyxl.utils import get_column_letter


class HydrologyCore:
    """Module tính toán chuyên sâu thủy văn và Pearson III."""
    
    # 9 Tần suất thiết kế chuẩn công trình thoát nước & thủy lợi
    STANDARD_P = [1.0, 2.0, 3.0, 4.0, 5.0, 10.0, 20.0, 25.0, 50.0]
    T_REPEAT = [100, 50, 33.3, 25, 20, 10, 5, 4, 2]

    # Bảng dải tần suất chi tiết cho bảng lý luận
    FULL_P = [0.1, 0.5, 1.0, 2.0, 3.0, 4.0, 5.0, 10.0, 20.0, 25.0, 33.3, 50.0, 75.0, 90.0, 95.0]

    def __init__(self):
        self.station_name = "CHƯA XÁC ĐỊNH"
        self.durations = []
        self.df_data = pd.DataFrame()
        self.stats = {}           # dur -> {mean, cv, cs, n, std}
        self.df_empirical = {}    # dur -> DataFrame(rank m, year, val, P_empirical)
        self.df_theoretical = {}  # dur -> DataFrame(P, T, H_p, I_min, I_h, q)
        self.df_Hp_summary = None
        self.cs_cv_ratio = 2.0

    def load_excel(self, filepath):
        excel_file = pd.ExcelFile(filepath)
        sheet_name = None
        for s in excel_file.sheet_names:
            if "rtd" in s.lower() or "mua" in s.lower():
                sheet_name = s
                break
        if not sheet_name:
            sheet_name = excel_file.sheet_names[0]

        df_raw = pd.read_excel(filepath, sheet_name=sheet_name, header=None)

        # 1. Nhận diện tên trạm
        cell_tram = str(df_raw.iloc[0, 0]) if pd.notna(df_raw.iloc[0, 0]) else ""
        if ":" in cell_tram:
            self.station_name = cell_tram.split(":")[-1].strip().upper()
        elif cell_tram.strip():
            self.station_name = cell_tram.strip().upper()
        else:
            self.station_name = "TRẠM ĐO MƯA"

        # 2. Quét dòng thời đoạn mưa ngắn
        dur_row_idx = None
        col_map = {}
        for r in range(min(12, len(df_raw))):
            row_vals = df_raw.iloc[r].values
            temp_map = {}
            for c, val in enumerate(row_vals):
                if pd.notna(val) and "phút" in str(val).lower():
                    try:
                        dur = int(str(val).lower().replace("phút", "").strip())
                        temp_map[dur] = c
                    except:
                        pass
            if len(temp_map) >= 2:
                dur_row_idx = r
                col_map = temp_map
                break

        if not col_map:
            raise ValueError("Không tìm thấy các cột thời đoạn ('... phút') trong dữ liệu.")

        self.durations = sorted(list(col_map.keys()))

        # 3. Trích xuất chuỗi số liệu quan trắc qua các năm
        records = []
        for r in range(dur_row_idx + 2, len(df_raw)):
            yr_val = df_raw.iloc[r, 0]
            if pd.isna(yr_val) or "max" in str(yr_val).lower() or "tb" in str(yr_val).lower():
                continue
            try:
                yr = int(yr_val)
                row_dict = {"Year": yr}
                for dur in self.durations:
                    v = df_raw.iloc[r, col_map[dur]]
                    row_dict[dur] = float(v) if pd.notna(v) else np.nan
                records.append(row_dict)
            except:
                pass

        if not records:
            raise ValueError("Không trích xuất được chuỗi số liệu quan trắc các năm.")

        self.df_data = pd.DataFrame(records).set_index("Year")
        self.compute_all()

    def compute_all(self):
        self.stats.clear()
        self.df_empirical.clear()
        self.df_theoretical.clear()

        self.df_Hp_summary = pd.DataFrame(index=self.durations, columns=self.STANDARD_P)

        for dur in self.durations:
            series = self.df_data[dur].dropna()
            n = len(series)
            if n < 3:
                continue

            mean = series.mean()
            std = series.std(ddof=1)
            cv = std / mean
            cs = self.cs_cv_ratio * cv

            self.stats[dur] = {
                "n": n, "mean": mean, "std": std, "cv": cv, "cs": cs
            }

            # Bảng tần suất kinh nghiệm (Weibull: P = m / (n + 1) * 100%)
            df_emp = pd.DataFrame({
                "Year": series.index,
                "H": series.values
            }).sort_values(by="H", ascending=False).reset_index(drop=True)
            df_emp["Rank"] = df_emp.index + 1
            df_emp["P_emp"] = (df_emp["Rank"] / (n + 1.0)) * 100.0
            self.df_empirical[dur] = df_emp

            # Bảng tần suất lý luận Pearson III
            theo_records = []
            for p in self.FULL_P:
                prob = 1.0 - p / 100.0
                hp = pearson3.ppf(prob, skew=cs, loc=mean, scale=std)
                hp = max(0.0, hp)
                i_min = hp / dur
                i_h = i_min * 60.0
                q_ha = 166.7 * i_min
                t_repeat = round(100.0 / p, 1) if p > 0 else np.nan
                theo_records.append({
                    "P(%)": p,
                    "T(năm)": t_repeat,
                    "H_p(mm)": hp,
                    "I(mm/phút)": i_min,
                    "I(mm/h)": i_h,
                    "q(l/s.ha)": q_ha
                })
            self.df_theoretical[dur] = pd.DataFrame(theo_records)

            for p in self.STANDARD_P:
                prob = 1.0 - p / 100.0
                self.df_Hp_summary.loc[dur, p] = pearson3.ppf(prob, skew=cs, loc=mean, scale=std)


class IDFApplication(tk.Tk):
    """Giao diện phần mềm đồ họa hoàn chỉnh."""

    def __init__(self):
        super().__init__()
        self.title("Phần Mềm Thủy Văn: Phân Tích Tần Suất & Xây Dựng Đường Cong IDF")
        self.geometry("1260x840")
        self.minsize(1050, 720)

        self.core = HydrologyCore()
        self.current_filepath = None

        self._apply_style()
        self._build_header_ui()
        self._build_body_tabs()
        self._build_statusbar()

    def _apply_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except:
            pass
        style.configure("Treeview", font=("Times New Roman", 10), rowheight=24)
        style.configure("Treeview.Heading", font=("Times New Roman", 10, "bold"))
        style.configure("TNotebook.Tab", font=("Times New Roman", 10, "bold"), padding=[10, 5])
        style.configure("TButton", font=("Times New Roman", 10))
        style.configure("TLabel", font=("Times New Roman", 10))

    def _build_header_ui(self):
        top_frame = ttk.LabelFrame(self, text=" Bảng Điều Khiển & Thiết Lập ", padding=10)
        top_frame.pack(fill="x", padx=10, pady=5)

        btn_open = ttk.Button(top_frame, text="📁 Chọn file Excel số liệu...", command=self.on_open_file)
        btn_open.grid(row=0, column=0, padx=5, pady=2, sticky="w")

        self.lbl_file = ttk.Label(top_frame, text="Chưa mở file dữ liệu nào", foreground="#555")
        self.lbl_file.grid(row=0, column=1, padx=10, pady=2, sticky="w")

        lbl_cs = ttk.Label(top_frame, text="Tỷ số Cs/Cv:")
        lbl_cs.grid(row=0, column=2, padx=(20, 5), pady=2, sticky="e")
        self.spin_cs = ttk.Spinbox(top_frame, from_=1.0, to=5.0, increment=0.5, width=6)
        self.spin_cs.set(2.0)
        self.spin_cs.grid(row=0, column=3, padx=5, pady=2, sticky="w")

        btn_recalc = ttk.Button(top_frame, text="🔄 Tính Lại", command=self.on_recalculate)
        btn_recalc.grid(row=0, column=4, padx=5, pady=2, sticky="w")

        btn_export = ttk.Button(top_frame, text="💾 Xuất Báo Cáo Excel (9 Biểu Đồ)...", command=self.on_export_excel)
        btn_export.grid(row=0, column=5, padx=(30, 5), pady=2, sticky="e")

    def _build_body_tabs(self):
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=5)

        # Tab 1: Tham số thống kê
        self.tab_stats = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_stats, text=" 📊 Tham Số Thống Kê & Chuỗi Đo ")
        self._init_tab_stats()

        # Tab 2: Bảng tần suất kinh nghiệm & lý luận
        self.tab_freq_tables = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_freq_tables, text=" 📋 Bảng Tần Suất (Kinh Nghiệm - Lý Luận) ")
        self._init_tab_freq_tables()

        # Tab 3: Biểu đồ tần suất Pearson III (Giấy xác suất thủy văn)
        self.tab_freq_chart = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_freq_chart, text=" 📈 Biểu Đồ Tần Suất Chuẩn (Pearson III) ")
        self._init_tab_freq_chart()

        # Tab 4: 4 Bảng tra chỉ tiêu IDF
        self.tab_idf_tables = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_idf_tables, text=" 🌧️ Bảng Tra Cường Độ IDF ")
        self._init_tab_idf_tables()

        # Tab 5: Biểu đồ đường cong IDF riêng từng tần suất
        self.tab_idf_chart = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_idf_chart, text=" 📉 Đường Cong IDF Từng Tần Suất (I - t) ")
        self._init_tab_idf_chart()

    def _build_statusbar(self):
        self.statusbar = ttk.Label(self, text="Sẵn sàng.", relief="sunken", anchor="w", padding=4)
        self.statusbar.pack(fill="x", side="bottom")

    # ------------------ CẤU HÌNH GIAO DIỆN CÁC TAB ------------------ #
    def _init_tab_stats(self):
        paned = ttk.PanedWindow(self.tab_stats, orient="vertical")
        paned.pack(fill="both", expand=True, padx=5, pady=5)

        f_top = ttk.LabelFrame(paned, text=" Tham số thống kê mưa thời đoạn ngắn ", padding=5)
        paned.add(f_top, weight=1)

        cols_stat = ("Thời đoạn (phút)", "Số năm (n)", "X_tb (mm)", "Độ lệch chuẩn (S)", "Cv", "Cs")
        self.tree_stats = ttk.Treeview(f_top, columns=cols_stat, show="headings", height=5)
        for col in cols_stat:
            self.tree_stats.heading(col, text=col)
            self.tree_stats.column(col, anchor="center", width=140)
        self.tree_stats.pack(fill="both", expand=True)

        f_bot = ttk.LabelFrame(paned, text=" Chuỗi số liệu quan trắc qua các năm (mm) ", padding=5)
        paned.add(f_bot, weight=2)

        self.tree_raw = ttk.Treeview(f_bot, show="headings", height=10)
        sc_y = ttk.Scrollbar(f_bot, orient="vertical", command=self.tree_raw.yview)
        sc_x = ttk.Scrollbar(f_bot, orient="horizontal", command=self.tree_raw.xview)
        self.tree_raw.configure(yscrollcommand=sc_y.set, xscrollcommand=sc_x.set)
        sc_y.pack(side="right", fill="y")
        sc_x.pack(side="bottom", fill="x")
        self.tree_raw.pack(fill="both", expand=True)

    def _init_tab_freq_tables(self):
        ctrl = ttk.Frame(self.tab_freq_tables, padding=5)
        ctrl.pack(fill="x")
        ttk.Label(ctrl, text="Chọn thời đoạn xem bảng tần suất: ", font=("Times New Roman", 10, "bold")).pack(side="left")
        self.combo_dur_table = ttk.Combobox(ctrl, state="readonly", width=15)
        self.combo_dur_table.pack(side="left", padx=5)
        self.combo_dur_table.bind("<<ComboboxSelected>>", self._on_select_dur_table)

        body = ttk.Frame(self.tab_freq_tables)
        body.pack(fill="both", expand=True, padx=5, pady=5)

        # Trái: Bảng tần suất kinh nghiệm
        f_left = ttk.LabelFrame(body, text=" Tần suất kinh nghiệm (Weibull: P = m / (n + 1)) ", padding=5)
        f_left.pack(side="left", fill="both", expand=True, padx=3)

        cols_emp = ("Thứ tự (m)", "Năm", "Lượng mưa H (mm)", "Tần suất P_emp (%)")
        self.tree_emp = ttk.Treeview(f_left, columns=cols_emp, show="headings")
        for c in cols_emp:
            self.tree_emp.heading(c, text=c)
            self.tree_emp.column(c, anchor="center", width=100)
        sc_emp = ttk.Scrollbar(f_left, orient="vertical", command=self.tree_emp.yview)
        self.tree_emp.configure(yscrollcommand=sc_emp.set)
        sc_emp.pack(side="right", fill="y")
        self.tree_emp.pack(fill="both", expand=True)

        # Phải: Bảng tần suất lý luận
        f_right = ttk.LabelFrame(body, text=" Tần suất lý luận (Phân phối Pearson III) ", padding=5)
        f_right.pack(side="right", fill="both", expand=True, padx=3)

        cols_theo = ("Tần suất P (%)", "Chu kỳ T (năm)", "H_p (mm)", "I (mm/phút)", "I (mm/h)", "q (l/s.ha)")
        self.tree_theo = ttk.Treeview(f_right, columns=cols_theo, show="headings")
        for c in cols_theo:
            self.tree_theo.heading(c, text=c)
            self.tree_theo.column(c, anchor="center", width=105)
        sc_theo = ttk.Scrollbar(f_right, orient="vertical", command=self.tree_theo.yview)
        self.tree_theo.configure(yscrollcommand=sc_theo.set)
        sc_theo.pack(side="right", fill="y")
        self.tree_theo.pack(fill="both", expand=True)

    def _init_tab_freq_chart(self):
        ctrl = ttk.Frame(self.tab_freq_chart, padding=5)
        ctrl.pack(fill="x")
        ttk.Label(ctrl, text="Thời đoạn mưa ngắn: ", font=("Times New Roman", 10, "bold")).pack(side="left")
        self.combo_dur_chart = ttk.Combobox(ctrl, state="readonly", width=15)
        self.combo_dur_chart.pack(side="left", padx=5)
        self.combo_dur_chart.bind("<<ComboboxSelected>>", self._on_select_dur_chart)

        self.fig_freq = Figure(figsize=(9.5, 6), dpi=100)
        self.ax_freq = self.fig_freq.add_subplot(111)

        self.canvas_freq = FigureCanvasTkAgg(self.fig_freq, master=self.tab_freq_chart)
        self.canvas_freq.get_tk_widget().pack(fill="both", expand=True)

        toolbar = NavigationToolbar2Tk(self.canvas_freq, self.tab_freq_chart)
        toolbar.update()
        toolbar.pack(side="bottom", fill="x")

    def _init_tab_idf_tables(self):
        paned = ttk.PanedWindow(self.tab_idf_tables, orient="vertical")
        paned.pack(fill="both", expand=True, padx=5, pady=5)

        self.tree_idf_h = self._create_idf_tree(paned, "1. Lượng mưa thời đoạn tính toán H_p (mm)")
        self.tree_idf_imin = self._create_idf_tree(paned, "2. Cường độ mưa I (mm/phút)")
        self.tree_idf_ih = self._create_idf_tree(paned, "3. Cường độ mưa I (mm/h)")
        self.tree_idf_q = self._create_idf_tree(paned, "4. Cường độ mưa q (l/s/ha)")

    def _create_idf_tree(self, parent, title):
        frame = ttk.LabelFrame(parent, text=f" {title} ", padding=5)
        parent.add(frame, weight=1)
        tree = ttk.Treeview(frame, show="headings", height=4)
        tree.pack(fill="both", expand=True)
        return tree

    def _init_tab_idf_chart(self):
        ctrl = ttk.Frame(self.tab_idf_chart, padding=5)
        ctrl.pack(fill="x")

        ttk.Label(ctrl, text="Chọn tần suất cần vẽ biểu đồ IDF: ", font=("Times New Roman", 10, "bold")).pack(side="left")
        self.combo_idf_freq = ttk.Combobox(ctrl, state="readonly", width=36)
        self.combo_idf_freq.pack(side="left", padx=5)
        self.combo_idf_freq.bind("<<ComboboxSelected>>", self._on_select_idf_freq)

        self.fig_idf = Figure(figsize=(9.5, 6), dpi=100)
        self.ax_idf = self.fig_idf.add_subplot(111)

        self.canvas_idf = FigureCanvasTkAgg(self.fig_idf, master=self.tab_idf_chart)
        self.canvas_idf.get_tk_widget().pack(fill="both", expand=True)

        toolbar = NavigationToolbar2Tk(self.canvas_idf, self.tab_idf_chart)
        toolbar.update()
        toolbar.pack(side="bottom", fill="x")

    # ------------------ XỬ LÝ SỰ KIỆN & CẬP NHẬT GIAO DIỆN ------------------ #
    def on_open_file(self):
        path = filedialog.askopenfilename(
            title="Chọn file dữ liệu mưa thời đoạn ngắn",
            filetypes=[("Excel Files", "*.xlsx *.xls")]
        )
        if not path:
            return

        try:
            self.core.load_excel(path)
            self.current_filepath = path
            self.lbl_file.config(text=f"{os.path.basename(path)} | Trạm: {self.core.station_name}", foreground="darkgreen")
            self.update_all_views()
            self.statusbar.config(text=f"Đã nạp trạm {self.core.station_name} ({len(self.core.durations)} thời đoạn quan trắc).")
        except Exception as e:
            messagebox.showerror("Lỗi dữ liệu", f"Không thể xử lý file Excel đã chọn:\n{e}")

    def on_recalculate(self):
        if not self.core.durations:
            return
        try:
            ratio = float(self.spin_cs.get())
            self.core.cs_cv_ratio = ratio
            self.core.compute_all()
            self.update_all_views()
            self.statusbar.config(text=f"Đã cập nhật tính toán theo tỷ số Cs/Cv = {ratio:.2f}")
        except Exception as e:
            messagebox.showerror("Lỗi tham số", f"Tỷ số Cs/Cv không hợp lệ:\n{e}")

    def update_all_views(self):
        dur_list = [f"{d} phút" for d in self.core.durations]
        self.combo_dur_table["values"] = dur_list
        self.combo_dur_chart["values"] = dur_list
        if dur_list:
            self.combo_dur_table.current(0)
            self.combo_dur_chart.current(0)

        # Cấu hình danh sách chọn tần suất cho biểu đồ IDF
        freq_options = []
        for p, t in zip(self.core.STANDARD_P, self.core.T_REPEAT):
            freq_options.append(f"Tần suất P = {p:.1f}%  (Chu kỳ T = {t} năm)")
        freq_options.append("--- Hiển thị tất cả tần suất trên cùng 1 biểu đồ ---")
        self.combo_idf_freq["values"] = freq_options
        self.combo_idf_freq.current(0)

        self._render_tab_stats()
        self._on_select_dur_table()
        self._on_select_dur_chart()
        self._render_tab_idf_tables()
        self._on_select_idf_freq()

    def _render_tab_stats(self):
        for r in self.tree_stats.get_children():
            self.tree_stats.delete(r)
        for dur in self.core.durations:
            st = self.core.stats[dur]
            self.tree_stats.insert("", "end", values=(
                f"{dur} phút",
                st["n"],
                f"{st['mean']:.2f}",
                f"{st['std']:.2f}",
                f"{st['cv']:.3f}",
                f"{st['cs']:.3f}"
            ))

        for r in self.tree_raw.get_children():
            self.tree_raw.delete(r)
        cols = ["Năm"] + [f"{d} phút" for d in self.core.durations]
        self.tree_raw["columns"] = cols
        for c in cols:
            self.tree_raw.heading(c, text=c)
            self.tree_raw.column(c, anchor="center", width=85)

        for yr, row in self.core.df_data.iterrows():
            vals = [yr] + [f"{row[d]:.2f}" if pd.notna(row[d]) else "-" for d in self.core.durations]
            self.tree_raw.insert("", "end", values=vals)

    def _on_select_dur_table(self, event=None):
        idx = self.combo_dur_table.current()
        if idx < 0 or not self.core.durations:
            return
        dur = self.core.durations[idx]

        for r in self.tree_emp.get_children():
            self.tree_emp.delete(r)
        df_emp = self.core.df_empirical.get(dur, pd.DataFrame())
        for _, row in df_emp.iterrows():
            self.tree_emp.insert("", "end", values=(
                int(row["Rank"]),
                int(row["Year"]),
                f"{row['H']:.2f}",
                f"{row['P_emp']:.2f}%"
            ))

        for r in self.tree_theo.get_children():
            self.tree_theo.delete(r)
        df_theo = self.core.df_theoretical.get(dur, pd.DataFrame())
        for _, row in df_theo.iterrows():
            self.tree_theo.insert("", "end", values=(
                f"{row['P(%)']:.2f}%",
                row['T(năm)'],
                f"{row['H_p(mm)']:.2f}",
                f"{row['I(mm/phút)']:.3f}",
                f"{row['I(mm/h)']:.2f}",
                f"{row['q(l/s.ha)']:.2f}"
            ))

    # ------------------ HIỆU CHỈNH CHUẨN BIỂU ĐỒ TẦN SUẤT ------------------ #
    def _on_select_dur_chart(self, event=None):
        idx = self.combo_dur_chart.current()
        if idx < 0 or not self.core.durations:
            return
        dur = self.core.durations[idx]

        st = self.core.stats[dur]
        df_emp = self.core.df_empirical[dur]

        self.ax_freq.clear()

        # Dải tần suất thiết kế: 0.05% -> 99.5%
        p_smooth = np.linspace(0.05, 99.5, 400)
        prob_smooth = 1.0 - p_smooth / 100.0
        h_smooth = pearson3.ppf(prob_smooth, skew=st["cs"], loc=st["mean"], scale=st["std"])

        # Trục hoành chuẩn hóa Gauss: u = norm.ppf(P / 100)
        # Giúp P nhỏ (lũ/mưa cực trị) nằm bên TRÁI, P lớn nằm bên PHẢI; đường cong dốc tự nhiên từ TRÁI trên xuống PHẢI dưới.
        x_smooth = norm.ppf(p_smooth / 100.0)
        x_emp = norm.ppf(df_emp["P_emp"].values / 100.0)

        # 1. Vẽ đường tần suất lý luận Pearson III
        self.ax_freq.plot(x_smooth, h_smooth, color="#D00000", lw=2.2,
                          label="Đường tần suất lý luận Pearson III", zorder=3)

        # 2. Vẽ điểm kinh nghiệm Weibull (hình tròn xanh viền đậm)
        self.ax_freq.scatter(x_emp, df_emp["H"].values, color="#1F4E79", edgecolors="black",
                             linewidths=0.8, s=42, zorder=5, label="Tần suất kinh nghiệm (Weibull)")

        # 3. Thiết lập hệ trục & đường lưới xác suất thủy văn (Probability Paper Grid)
        major_p = [0.1, 0.5, 1, 2, 5, 10, 20, 30, 50, 70, 80, 90, 95, 98, 99]
        minor_p = [0.05, 0.2, 0.3, 0.4, 0.6, 0.7, 0.8, 0.9, 3, 4, 6, 7, 8, 9, 15, 25, 35, 40, 60, 75, 85, 97, 99.5]

        major_x = norm.ppf(np.array(major_p) / 100.0)
        minor_x = norm.ppf(np.array(minor_p) / 100.0)

        self.ax_freq.set_xticks(major_x)
        self.ax_freq.set_xticklabels([str(p) for p in major_p], fontsize=9)
        self.ax_freq.set_xticks(minor_x, minor=True)

        self.ax_freq.set_xlim(norm.ppf(0.0005), norm.ppf(0.9995))
        self.ax_freq.set_ylim(bottom=0)

        # Lưới tọa độ giấy xác suất
        self.ax_freq.grid(True, which="major", color="#888888", linestyle="-", linewidth=0.7, alpha=0.6)
        self.ax_freq.grid(True, which="minor", color="#BBBBBB", linestyle=":", linewidth=0.5, alpha=0.5)

        # Tiêu đề và nhãn
        self.ax_freq.set_title(f"ĐƯỜNG TẦN SUẤT LƯỢNG MƯA THỜI ĐOẠN {dur} PHÚT - TRẠM {self.core.station_name}",
                               fontsize=12, fontweight="bold", pad=12)
        self.ax_freq.set_xlabel("Tần suất thiết kế P (%)  [Giấy xác suất Gauss/P-III]", fontsize=10, fontweight="bold")
        self.ax_freq.set_ylabel(f"Lượng mưa H_{dur} (mm)", fontsize=10, fontweight="bold")

        # Hộp thông số thống kê
        info_box = (f"THÔNG SỐ THỐNG KÊ:\n"
                    f"• Số năm (n) = {st['n']}\n"
                    f"• X_tb = {st['mean']:.2f} mm\n"
                    f"• Cv = {st['cv']:.3f}\n"
                    f"• Cs = {st['cs']:.3f} (Cs=2Cv)")
        self.ax_freq.text(0.03, 0.05, info_box, transform=self.ax_freq.transAxes,
                          fontsize=9, verticalalignment="bottom",
                          bbox=dict(boxstyle="square,pad=0.6", facecolor="#FFFFE0", edgecolor="#888888", alpha=0.9))

        self.ax_freq.legend(loc="upper right", frameon=True, facecolor="white", edgecolor="#888888", fontsize=9)
        self.fig_freq.tight_layout()
        self.canvas_freq.draw()

    def _render_tab_idf_tables(self):
        cols = ["Thời đoạn t (phút)"] + [f"P={p}%\n(T={t}n)" for p, t in zip(self.core.STANDARD_P, self.core.T_REPEAT)]

        for tree in (self.tree_idf_h, self.tree_idf_imin, self.tree_idf_ih, self.tree_idf_q):
            for r in tree.get_children():
                tree.delete(r)
            tree["columns"] = cols
            for c in cols:
                tree.heading(c, text=c)
                tree.column(c, anchor="center", width=105)

        for dur in self.core.durations:
            row_h = [f"{dur} phút"]
            row_imin = [f"{dur} phút"]
            row_ih = [f"{dur} phút"]
            row_q = [f"{dur} phút"]

            for p in self.core.STANDARD_P:
                hp = self.core.df_Hp_summary.loc[dur, p]
                imin = hp / dur
                ih = imin * 60.0
                q = 166.7 * imin

                row_h.append(f"{hp:.2f}")
                row_imin.append(f"{imin:.3f}")
                row_ih.append(f"{ih:.2f}")
                row_q.append(f"{q:.2f}")

            self.tree_idf_h.insert("", "end", values=row_h)
            self.tree_idf_imin.insert("", "end", values=row_imin)
            self.tree_idf_ih.insert("", "end", values=row_ih)
            self.tree_idf_q.insert("", "end", values=row_q)

    # ------------------ HIỆU CHỈNH VẼ RIÊNG BIỂU ĐỒ IDF TỪNG TẦN SUẤT ------------------ #
    def _on_select_idf_freq(self, event=None):
        if not self.core.durations:
            return

        idx = self.combo_idf_freq.current()
        self.ax_idf.clear()

        dur_fit = [d for d in self.core.durations if d <= 180]
        if not dur_fit:
            dur_fit = self.core.durations
        t_arr = np.array(dur_fit)

        # Chế độ 1: Vẽ riêng 1 biểu đồ độc lập cho tần suất được chọn
        if idx < len(self.core.STANDARD_P):
            p = self.core.STANDARD_P[idx]
            t_repeat = self.core.T_REPEAT[idx]
            color = "#C00000"

            i_vals = np.array([self.core.df_Hp_summary.loc[d, p] / d for d in dur_fit])

            # Hồi quy hàm lũy thừa I = A * t^b
            log_t = np.log(t_arr)
            log_i = np.log(i_vals)
            poly = np.polyfit(log_t, log_i, 1)
            b = poly[0]
            A = np.exp(poly[1])

            # Tính R^2
            i_pred = A * (t_arr ** b)
            ss_res = np.sum((i_vals - i_pred) ** 2)
            ss_tot = np.sum((i_vals - np.mean(i_vals)) ** 2)
            r2 = 1.0 - (ss_res / ss_tot) if ss_tot > 0 else 1.0

            # Đường cong trơn
            t_curve = np.linspace(min(t_arr), max(t_arr), 150)
            i_curve = A * (t_curve ** b)

            # Vẽ điểm thực nghiệm
            self.ax_idf.scatter(t_arr, i_vals, color=color, s=55, edgecolors="black",
                                label="Số liệu tính toán từ P-III", zorder=5)

            # Vẽ đường hồi quy
            self.ax_idf.plot(t_curve, i_curve, color=color, lw=2.2,
                             label=f"Đường tương quan lũy thừa")

            # Hộp công thức hồi quy
            eq_text = (f"CÔNG THỨC ĐƯỜNG CONG IDF:\n"
                       f"• Phương trình: I = {A:.3f} · t^({b:.4f})\n"
                       f"• Hệ số tương quan: R² = {r2:.4f}\n"
                       f"• Chu kỳ lặp lại: N = {t_repeat} năm (P = {p}%)")
            self.ax_idf.text(0.55, 0.85, eq_text, transform=self.ax_idf.transAxes,
                             fontsize=10, verticalalignment="top",
                             bbox=dict(boxstyle="round,pad=0.7", facecolor="#F2F2F2", edgecolor=color, lw=1.5))

            self.ax_idf.set_title(f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = {p}% (T = {t_repeat} NĂM)\nTRẠM {self.core.station_name}",
                                  fontsize=12, fontweight="bold", pad=12)

        # Chế độ 2: Vẽ họ đường cong gộp tất cả tần suất
        else:
            colors = ["#C00000", "#ED7D31", "#A5A5A5", "#2E75B6", "#70AD47", "#385723", "#4472C4", "#7030A0", "#002060"]
            for i, p in enumerate(self.core.STANDARD_P):
                i_vals = np.array([self.core.df_Hp_summary.loc[d, p] / d for d in dur_fit])
                log_t = np.log(t_arr)
                log_i = np.log(i_vals)
                poly = np.polyfit(log_t, log_i, 1)
                b = poly[0]
                A = np.exp(poly[1])

                t_curve = np.linspace(min(t_arr), max(t_arr), 100)
                i_curve = A * (t_curve ** b)

                c = colors[i % len(colors)]
                self.ax_idf.scatter(t_arr, i_vals, color=c, s=35)
                self.ax_idf.plot(t_curve, i_curve, color=c, lw=1.8,
                                 label=f"P={p}% (T={self.core.T_REPEAT[i]}n): I={A:.2f}·t^({b:.3f})")

            self.ax_idf.set_title(f"TỔNG HỢP CÁC ĐƯỜNG CONG CƯỜNG ĐỘ MƯA (IDF) - TRẠM {self.core.station_name}",
                                  fontsize=12, fontweight="bold", pad=12)

        self.ax_idf.set_xlabel("Thời gian mưa t (phút)", fontsize=10, fontweight="bold")
        self.ax_idf.set_ylabel("Cường độ mưa I (mm/phút)", fontsize=10, fontweight="bold")
        self.ax_idf.grid(True, which="both", linestyle="--", alpha=0.6)
        self.ax_idf.legend(loc="lower left" if idx < len(self.core.STANDARD_P) else "upper right", frameon=True, fontsize=9)

        self.fig_idf.tight_layout()
        self.canvas_idf.draw()

    # ------------------ XUẤT BÁO CÁO EXCEL CHUYÊN NGHIỆP ------------------ #
    def on_export_excel(self):
        if not self.core.durations:
            messagebox.showwarning("Chưa có dữ liệu", "Vui lòng mở file và tính toán trước khi xuất kết quả.")
            return

        default_name = f"BaoCao_IDF_9_BieuDo_{self.core.station_name}.xlsx"
        save_path = filedialog.asksaveasfilename(
            title="Lưu báo cáo Excel (9 biểu đồ IDF)",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel Workbook", "*.xlsx")]
        )
        if not save_path:
            return

        try:
            self._export_full_excel(save_path)
            messagebox.showinfo("Hoàn thành", f"Báo cáo đầy đủ 9 biểu đồ IDF đã lưu tại:\n{save_path}")
            self.statusbar.config(text=f"Đã lưu thành công: {save_path}")
        except Exception as e:
            messagebox.showerror("Lỗi xuất file", f"Không thể lưu file Excel:\n{e}")

    def _export_full_excel(self, filepath):
        wb = openpyxl.Workbook()
        ws_idf = wb.active
        ws_idf.title = "Duong cong IDF"

        font_title = Font(name="Times New Roman", size=13, bold=True)
        font_tbl_header = Font(name="Times New Roman", size=10, bold=True)
        font_cell = Font(name="Times New Roman", size=10)
        fill_header = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        fill_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        thin = Side(border_style="thin", color="000000")
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

        p_headers = self.core.STANDARD_P
        t_repeat = self.core.T_REPEAT

        ws_idf.column_dimensions["A"].width = 12
        for c in ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
            ws_idf.column_dimensions[c].width = 13.5
        for c_idx in range(13, 35):
            ws_idf.column_dimensions[get_column_letter(c_idx)].width = 12

        ws_idf["A2"] = "XÂY DỰNG ĐƯỜNG CONG CƯỜNG ĐỘ MƯA (IDF)"
        ws_idf["A2"].font = font_title
        ws_idf["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws_idf.merge_cells("A2:K2")

        ws_idf["A3"] = f"TRẠM {self.core.station_name}"
        ws_idf["A3"].font = font_title
        ws_idf["A3"].alignment = Alignment(horizontal="center", vertical="center")
        ws_idf.merge_cells("A3:K3")

        def build_idf_table(ws, start_r, title_text):
            ws.cell(row=start_r, column=1, value="Thời đoạn")
            ws.cell(row=start_r, column=2, value=title_text)
            ws.merge_cells(start_row=start_r, start_column=2, end_row=start_r, end_column=11)

            ws.cell(row=start_r + 1, column=1, value="t (phút)")
            for idx, p in enumerate(p_headers, start=2):
                ws.cell(row=start_r + 1, column=idx, value=f"{p:.2f}")
            ws.cell(row=start_r + 1, column=11, value="Tần suất%")

            for idx, t in enumerate(t_repeat, start=2):
                ws.cell(row=start_r + 2, column=idx, value=t)
            ws.cell(row=start_r + 2, column=11, value="Chu kỳ lặp")
            ws.merge_cells(start_row=start_r + 1, start_column=1, end_row=start_r + 2, end_column=1)

            for r in range(start_r, start_r + 3):
                for c in range(1, 12):
                    cell = ws.cell(row=r, column=c)
                    cell.font = font_tbl_header
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.fill = fill_header
                    cell.border = border_all

        # 1. Bảng Lượng mưa H (mm)
        build_idf_table(ws_idf, 6, "Lượng mưa thời đoạn I (mm) tương ứng với chu kỳ lặp lại N")
        for r_idx, dur in enumerate(self.core.durations, start=9):
            ws_idf.cell(row=r_idx, column=1, value=dur)
            st = self.core.stats[dur]
            for c_idx, p in enumerate(p_headers, start=2):
                prob = 1.0 - p / 100.0
                hp = pearson3.ppf(prob, skew=st["cs"], loc=st["mean"], scale=st["std"])
                c = ws_idf.cell(row=r_idx, column=c_idx, value=round(float(hp), 2))
                c.number_format = "0.00"
            for c in range(1, 12):
                cell = ws_idf.cell(row=r_idx, column=c)
                cell.font = font_cell
                cell.border = border_all
                cell.alignment = Alignment(horizontal="center" if c == 1 else "right", vertical="center")

        # 2. Bảng Cường độ I (mm/phút)
        build_idf_table(ws_idf, 23, "Cường độ mưa I (mm/phút) tương ứng với chu kỳ lặp lại N")
        for r_idx, dur in enumerate(self.core.durations, start=26):
            ws_idf.cell(row=r_idx, column=1, value=dur)
            hp_r = r_idx - 17
            for c_idx in range(2, 11):
                col_let = get_column_letter(c_idx)
                c = ws_idf.cell(row=r_idx, column=c_idx, value=f"={col_let}{hp_r}/$A{r_idx}")
                c.number_format = "0.000"
            for c in range(1, 12):
                cell = ws_idf.cell(row=r_idx, column=c)
                cell.font = font_cell
                cell.border = border_all
                cell.alignment = Alignment(horizontal="center" if c == 1 else "right", vertical="center")

        # 3. Bảng Cường độ I (mm/h)
        build_idf_table(ws_idf, 40, "Cường độ mưa I (mm/h) tương ứng với chu kỳ lặp lại N")
        for r_idx, dur in enumerate(self.core.durations, start=43):
            ws_idf.cell(row=r_idx, column=1, value=dur)
            ip_r = r_idx - 17
            for c_idx in range(2, 11):
                col_let = get_column_letter(c_idx)
                c = ws_idf.cell(row=r_idx, column=c_idx, value=f"={col_let}{ip_r}*60")
                c.number_format = "0.00"
            for c in range(1, 12):
                cell = ws_idf.cell(row=r_idx, column=c)
                cell.font = font_cell
                cell.border = border_all
                cell.alignment = Alignment(horizontal="center" if c == 1 else "right", vertical="center")

        # 4. Bảng Cường độ q (l/s/ha)
        build_idf_table(ws_idf, 57, "Cường độ mưa q (l/s/ha) tương ứng với chu kỳ lặp lại N")
        for r_idx, dur in enumerate(self.core.durations, start=60):
            ws_idf.cell(row=r_idx, column=1, value=dur)
            ip_r = r_idx - 34
            for c_idx in range(2, 11):
                col_let = get_column_letter(c_idx)
                c = ws_idf.cell(row=r_idx, column=c_idx, value=f"=166.7*{col_let}{ip_r}")
                c.number_format = "0.00"
            for c in range(1, 12):
                cell = ws_idf.cell(row=r_idx, column=c)
                cell.font = font_cell
                cell.border = border_all
                cell.alignment = Alignment(horizontal="center" if c == 1 else "right", vertical="center")

        # NHÚNG TOÀN BỘ 9 BIỂU ĐỒ CHO TỪNG TẦN SUẤT RIÊNG BIỆT
        max_fit_row = 26 + len([d for d in self.core.durations if d <= 180]) - 1
        xvalues = Reference(ws_idf, min_col=1, min_row=26, max_row=max_fit_row)

        chart_configs = [
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 1% (T = 100 NĂM)\nTRẠM {self.core.station_name}", 2, "M6", "C00000"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 2% (T = 50 NĂM)\nTRẠM {self.core.station_name}", 3, "W6", "ED7D31"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 3% (T = 33.3 NĂM)\nTRẠM {self.core.station_name}", 4, "M26", "A5A5A5"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 4% (T = 25 NĂM)\nTRẠM {self.core.station_name}", 5, "W26", "2E75B6"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 5% (T = 20 NĂM)\nTRẠM {self.core.station_name}", 6, "M46", "70AD47"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 10% (T = 10 NĂM)\nTRẠM {self.core.station_name}", 7, "W46", "385723"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 20% (T = 5 NĂM)\nTRẠM {self.core.station_name}", 8, "M66", "4472C4"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 25% (T = 4 NĂM)\nTRẠM {self.core.station_name}", 9, "W66", "7030A0"),
            (f"ĐƯỜNG CONG CƯỜNG ĐỘ MƯA TẦN SUẤT P = 50% (T = 2 NĂM)\nTRẠM {self.core.station_name}", 10, "M86", "002060")
        ]

        for title, col_num, cell_loc, marker_color in chart_configs:
            chart = ScatterChart()
            chart.style = 13
            chart.width = 15.5
            chart.height = 9.8
            self._set_chart_title_tnr(chart, title, size=1050, bold=True)
            self._set_axis_title_tnr(chart.x_axis, "Thời gian t (phút)", size=950, bold=False)
            self._set_axis_title_tnr(chart.y_axis, "Cường độ mưa I (mm/phút)", size=950, bold=False)

            chart.x_axis.majorGridlines = ChartLines()
            chart.y_axis.majorGridlines = ChartLines()

            yvalues = Reference(ws_idf, min_col=col_num, min_row=26, max_row=max_fit_row)
            series = Series(yvalues, xvalues, title_from_data=False)
            series.marker.symbol = "circle"
            series.marker.size = 6
            series.marker.graphicalProperties.solidFill = marker_color
            series.marker.graphicalProperties.line.solidFill = marker_color
            series.graphicalProperties.line.noFill = True

            trendline = Trendline()
            trendline.trendlineType = "power"
            trendline.dispEq = True
            trendline.dispRSqr = True
            series.trendline = trendline

            chart.series.append(series)
            chart.legend = None
            ws_idf.add_chart(chart, cell_loc)

        # SHEET 2: BẢNG TẦN SUẤT CHI TIẾT TỪNG THỜI ĐOẠN
        ws_ts = wb.create_sheet(title="TanSuat_TungThoiDoan")
        ws_ts.views.sheetView[0].showGridLines = True

        r_curr = 2
        for dur in self.core.durations:
            st = self.core.stats[dur]
            ws_ts.cell(row=r_curr, column=1, value=f"PHÂN TÍCH TẦN SUẤT THỜI ĐOẠN {dur} PHÚT (X_tb={st['mean']:.2f}mm, Cv={st['cv']:.3f}, Cs={st['cs']:.3f})").font = font_title
            r_curr += 1

            ws_ts.cell(row=r_curr, column=1, value="Bảng Tần Suất Kinh Nghiệm (Weibull)").font = font_tbl_header
            ws_ts.cell(row=r_curr, column=6, value="Bảng Tần Suất Lý Luận (Pearson III)").font = font_tbl_header
            r_curr += 1

            emp_headers = ["Hạng (m)", "Năm", "H (mm)", "P_kinh_nghiệm (%)"]
            for ci, h in enumerate(emp_headers, start=1):
                c = ws_ts.cell(row=r_curr, column=ci, value=h)
                c.font = font_tbl_header
                c.fill = fill_sub
                c.border = border_all
                c.alignment = Alignment(horizontal="center")

            theo_headers = ["Tần suất P (%)", "Chu kỳ T (năm)", "H_p (mm)", "I (mm/phút)", "I (mm/h)", "q (l/s.ha)"]
            for ci, h in enumerate(theo_headers, start=6):
                c = ws_ts.cell(row=r_curr, column=ci, value=h)
                c.font = font_tbl_header
                c.fill = fill_sub
                c.border = border_all
                c.alignment = Alignment(horizontal="center")
            r_curr += 1

            df_e = self.core.df_empirical[dur]
            df_t = self.core.df_theoretical[dur]
            max_rows = max(len(df_e), len(df_t))

            for i in range(max_rows):
                if i < len(df_e):
                    row_e = df_e.iloc[i]
                    ws_ts.cell(row=r_curr, column=1, value=int(row_e["Rank"])).alignment = Alignment(horizontal="center")
                    ws_ts.cell(row=r_curr, column=2, value=int(row_e["Year"])).alignment = Alignment(horizontal="center")
                    ws_ts.cell(row=r_curr, column=3, value=round(row_e["H"], 2)).number_format = "0.00"
                    ws_ts.cell(row=r_curr, column=4, value=round(row_e["P_emp"], 2)).number_format = "0.00"
                    for ci in range(1, 5):
                        ws_ts.cell(row=r_curr, column=ci).border = border_all
                        ws_ts.cell(row=r_curr, column=ci).font = font_cell

                if i < len(df_t):
                    row_t = df_t.iloc[i]
                    ws_ts.cell(row=r_curr, column=6, value=round(row_t["P(%)"], 2)).number_format = "0.00"
                    ws_ts.cell(row=r_curr, column=7, value=row_t["T(năm)"]).alignment = Alignment(horizontal="center")
                    ws_ts.cell(row=r_curr, column=8, value=round(row_t["H_p(mm)"], 2)).number_format = "0.00"
                    ws_ts.cell(row=r_curr, column=9, value=round(row_t["I(mm/phút)"], 3)).number_format = "0.000"
                    ws_ts.cell(row=r_curr, column=10, value=round(row_t["I(mm/h)"], 2)).number_format = "0.00"
                    ws_ts.cell(row=r_curr, column=11, value=round(row_t["q(l/s.ha)"], 2)).number_format = "0.00"
                    for ci in range(6, 12):
                        ws_ts.cell(row=r_curr, column=ci).border = border_all
                        ws_ts.cell(row=r_curr, column=ci).font = font_cell

                r_curr += 1
            r_curr += 2

        for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
            ws_ts.column_dimensions[col].width = 16

        wb.save(filepath)

    @staticmethod
    def _set_chart_title_tnr(chart, title_text, size=1050, bold=True):
        font_tnr = DrawingFont(typeface="Times New Roman")
        cp = CharacterProperties(latin=font_tnr, sz=size, b=bold)
        lines = title_text.split("\n")
        paras = [Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(t=line, rPr=cp)]) for line in lines]
        chart.title = Title(tx=Text(rich=RichText(p=paras)))

    @staticmethod
    def _set_axis_title_tnr(axis, title_text, size=950, bold=False):
        font_tnr = DrawingFont(typeface="Times New Roman")
        cp = CharacterProperties(latin=font_tnr, sz=size, b=bold)
        run = RegularTextRun(t=title_text, rPr=cp)
        axis.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[run])])))
        axis.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(latin=font_tnr, sz=850)))])


if __name__ == "__main__":
    app = IDFApplication()
    app.mainloop()