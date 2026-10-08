import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.ticker import MultipleLocator
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Cấu hình font chữ chuẩn kỹ thuật
plt.rcParams['font.family'] = 'serif'
plt.rcParams['font.serif'] = ['Times New Roman', 'Arial', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False


class MultiYearWaterLevelApp:
    def __init__(self, root):
        self.root = root
        self.root.title("TẦN SUẤT LŨY TÍCH MỰC NƯỚC GIỜ NHIỀU NĂM")
        self.root.geometry("750x610")
        self.root.resizable(False, False)

        # Thoát dứt điểm tiến trình chạy ngầm khi tắt
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.plot_win = None
        self.tbl_win = None

        # Dữ liệu
        self.years_data = {}      # {2010: np.array([...]), 2011: ...}
        self.all_data = np.array([])
        self.actual_max = 0.0
        self.actual_min = 0.0
        self.df_multi = None
        self.lower_bounds = None
        self.cum_freq_pct = None

        self.setup_ui()

    def on_close(self):
        try:
            plt.close('all')
            if self.plot_win and self.plot_win.winfo_exists(): self.plot_win.destroy()
            if self.tbl_win and self.tbl_win.winfo_exists(): self.tbl_win.destroy()
            self.root.quit()
            self.root.destroy()
        except Exception:
            pass
        finally:
            os._exit(0)

    def setup_ui(self):
        # 1. TIÊU ĐỀ
        hdr = tk.Frame(self.root, pady=8)
        hdr.pack(fill=tk.X)
        tk.Label(hdr, text="PHÂN TÍCH TẦN SUẤT LŨY TÍCH MỰC NƯỚC GIỜ", 
                 font=("Times New Roman", 15, "bold"), fg="#154360").pack()
        tk.Label(hdr, text="[ Xử lý chuỗi một năm hoặc nhiều năm - Tự động nhận diện Min/Max ]", 
                 font=("Times New Roman", 9, "italic"), fg="#7f8c8d").pack()

        # 2. CHỌN FILE
        file_fr = tk.LabelFrame(self.root, text=" 1. Nạp dữ liệu Excel (1 file nhiều sheet hoặc chọn nhiều file) ", 
                                font=("Arial", 9, "bold"), padx=10, pady=8)
        file_fr.pack(fill=tk.X, padx=15, pady=4)

        self.lbl_file = tk.Label(file_fr, text="Chưa nạp file dữ liệu...", fg="gray", anchor=tk.W)
        self.lbl_file.pack(side=tk.LEFT, fill=tk.X, expand=True)

        tk.Button(file_fr, text="📁 Chọn file Excel...", font=("Arial", 9, "bold"),
                  bg="#e0e0e0", command=self.load_excel_files).pack(side=tk.RIGHT, padx=5)

        # 3. THÔNG TIN TRẠM & KHUNG TÊN
        info_fr = tk.LabelFrame(self.root, text=" 2. Thông tin trạm quan trắc ", font=("Arial", 9, "bold"), padx=10, pady=8)
        info_fr.pack(fill=tk.X, padx=15, pady=4)

        r1 = tk.Frame(info_fr)
        r1.pack(fill=tk.X, pady=2)
        tk.Label(r1, text="Tên trạm:", width=10, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_station = tk.Entry(r1, width=22)
        self.txt_station.insert(0, "TÂN AN")
        self.txt_station.pack(side=tk.LEFT, padx=(0, 15))

        tk.Label(r1, text="Sông:").pack(side=tk.LEFT)
        self.txt_river = tk.Entry(r1, width=22)
        self.txt_river.insert(0, "VÀM CỎ TÂY")
        self.txt_river.pack(side=tk.LEFT, padx=(0, 15))

        tk.Label(r1, text="Đơn vị:").pack(side=tk.LEFT)
        self.txt_unit = tk.Entry(r1, width=6)
        self.txt_unit.insert(0, "cm")
        self.txt_unit.pack(side=tk.LEFT)

        self.lbl_years_info = tk.Label(info_fr, text="Các năm có dữ liệu: Chưa có số liệu", 
                                       font=("Arial", 9, "italic"), fg="blue", anchor=tk.W)
        self.lbl_years_info.pack(fill=tk.X, pady=(4, 0))

        # 4. THIẾT LẬP CẤP MỰC NƯỚC (CÓ TỰ ĐỘNG ĐIỀN MIN/MAX)
        step_fr = tk.LabelFrame(self.root, text=" 3. Phân chia cấp mực nước tính toán ", font=("Arial", 9, "bold"), padx=10, pady=8)
        step_fr.pack(fill=tk.X, padx=15, pady=4)

        # Dòng hiển thị giá trị Min/Max thực tế đo được
        self.lbl_actual_stat = tk.Label(step_fr, text="Số liệu đo thực tế: Max = ---  |  Min = ---", 
                                        font=("Arial", 9, "bold"), fg="#d32f2f", anchor=tk.W)
        self.lbl_actual_stat.pack(fill=tk.X, pady=(0, 6))

        r3 = tk.Frame(step_fr)
        r3.pack(fill=tk.X, pady=2)

        tk.Label(r3, text="Bước cấp (Step):").pack(side=tk.LEFT)
        self.txt_step = tk.Entry(r3, width=5)
        self.txt_step.insert(0, "5")
        self.txt_step.pack(side=tk.LEFT, padx=3)
        tk.Label(r3, text="(cm)").pack(side=tk.LEFT, padx=(0, 15))

        tk.Label(r3, text="Cấp Max:").pack(side=tk.LEFT)
        self.txt_hmax = tk.Entry(r3, width=7, font=("Arial", 9, "bold"), fg="#154360")
        self.txt_hmax.pack(side=tk.LEFT, padx=3)

        tk.Label(r3, text="Cấp Min:").pack(side=tk.LEFT)
        self.txt_hmin = tk.Entry(r3, width=7, font=("Arial", 9, "bold"), fg="#154360")
        self.txt_hmin.pack(side=tk.LEFT, padx=3)

        tk.Button(r3, text="🔄 Gợi ý lại theo Step", font=("Arial", 8, "bold"), 
                  bg="#e8f5e9", fg="#2e7d32", command=self.auto_fill_min_max).pack(side=tk.LEFT, padx=10)

        tk.Label(step_fr, text="* Sau khi nạp file, Cấp Max/Min sẽ tự động điền giá trị làm tròn. Bạn có thể tự do gõ sửa lại số theo ý muốn.", 
                 font=("Arial", 8, "italic"), fg="#7f8c8d").pack(anchor=tk.W, pady=(4, 0))

        # 5. CÔNG CỤ TRA CỨU MỰC NƯỚC THEO TẦN SUẤT BẤT KỲ
        lookup_fr = tk.LabelFrame(self.root, text=" 4. Tra cứu mực nước theo tần suất bất kỳ ", 
                                  font=("Arial", 9, "bold"), fg="#b71c1c", padx=10, pady=8)
        lookup_fr.pack(fill=tk.X, padx=15, pady=4)

        r4 = tk.Frame(lookup_fr)
        r4.pack(fill=tk.X, pady=2)
        tk.Label(r4, text="Nhập tần suất P (%):", font=("Arial", 9, "bold")).pack(side=tk.LEFT)
        self.txt_lookup_p = tk.Entry(r4, width=8, font=("Arial", 9, "bold"))
        self.txt_lookup_p.insert(0, "50.0")
        self.txt_lookup_p.pack(side=tk.LEFT, padx=8)

        tk.Button(r4, text="🔍 TÍNH MỰC NƯỚC", font=("Arial", 9, "bold"), bg="#b71c1c", fg="white",
                  command=self.lookup_water_level).pack(side=tk.LEFT, padx=8)

        self.lbl_lookup_result = tk.Label(r4, text="Mực nước H = --- cm", font=("Arial", 10, "bold"), fg="#b71c1c")
        self.lbl_lookup_result.pack(side=tk.LEFT, padx=12)

        # 6. HÀNG NÚT BẤM HÀNH ĐỘNG
        btn_fr = tk.Frame(self.root, pady=10)
        btn_fr.pack(fill=tk.X, padx=15)

        tk.Button(btn_fr, text="📋 XEM BẢNG TỔNG HỢP NHIỀU NĂM", font=("Arial", 9, "bold"),
                  bg="#27ae60", fg="white", padx=12, pady=6, command=self.show_multi_year_table).pack(side=tk.LEFT)

        tk.Button(btn_fr, text="📈 VẼ BIỂU ĐỒ TẦN SUẤT LŨY TÍCH >>", font=("Arial", 10, "bold"),
                  bg="#154360", fg="white", padx=15, pady=6, command=self.plot_graph).pack(side=tk.RIGHT)

    def auto_fill_min_max(self):
        """Tự động tính toán Cấp Max và Cấp Min làm tròn bao trọn chuỗi số liệu thực tế"""
        if len(self.all_data) == 0:
            return
        try:
            step = float(self.txt_step.get())
            if step <= 0: step = 5.0
        except ValueError:
            step = 5.0

        # Làm tròn Cấp Max lên bội số của step, Cấp Min xuống bội số của step
        suggest_max = int(np.ceil(self.actual_max / step) * step)
        suggest_min = int(np.floor(self.actual_min / step) * step)

        self.txt_hmax.delete(0, tk.END)
        self.txt_hmax.insert(0, str(suggest_max))

        self.txt_hmin.delete(0, tk.END)
        self.txt_hmin.insert(0, str(suggest_min))

    def extract_hours_from_df(self, df):
        """Hàm đọc số liệu chuẩn: lọc hàng ngày 1-31 và 24 cột giờ 0-23"""
        hour_cols = list(range(1, 25))
        for r in range(min(15, len(df))):
            cols_found = []
            for c in range(len(df.columns)):
                try:
                    v = int(float(df.iloc[r, c]))
                    if 0 <= v <= 23: cols_found.append(c)
                except Exception:
                    pass
            if len(cols_found) >= 24:
                hour_cols = cols_found[:24]
                break

        vals = []
        for r in range(len(df)):
            val0 = str(df.iloc[r, 0]).strip()
            try:
                d = float(val0)
                if d.is_integer() and 1 <= int(d) <= 31:
                    for c in hour_cols:
                        v = float(df.iloc[r, c])
                        if not np.isnan(v):
                            vals.append(v)
            except Exception:
                continue
        return np.array(vals)

    def load_excel_files(self):
        """Hỗ trợ chọn 1 file có nhiều sheet hoặc chọn nhiều file Excel cùng lúc"""
        files = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")])
        if not files: return

        self.years_data = {}
        st_name, rv_name = "", ""

        try:
            for f in files:
                xls = pd.ExcelFile(f)
                for sheet in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet, header=None)
                    if df.empty or len(df) < 5: continue

                    yr = None
                    for r in range(min(8, len(df))):
                        for c in range(min(15, len(df.columns))):
                            val = str(df.iloc[r, c]).strip()
                            if "Trạm:" in val and not st_name:
                                st_name = val.replace("Trạm:", "").strip()
                            elif "Sông:" in val and not rv_name:
                                rv_name = val.replace("Sông:", "").strip()
                            elif ("NĂM:" in val or "Năm:" in val) and yr is None:
                                try:
                                    yr = int(float(str(df.iloc[r, c+1]).strip()))
                                except Exception:
                                    pass

                    if yr is None:
                        try:
                            s_int = int(sheet)
                            yr = 2000 + s_int if s_int < 50 else 1900 + s_int
                        except Exception:
                            continue

                    h_vals = self.extract_hours_from_df(df)
                    if len(h_vals) > 0:
                        self.years_data[yr] = h_vals

            if not self.years_data:
                messagebox.showerror("Lỗi", "Không tìm thấy chuỗi số liệu giờ (0h-23h) hợp lệ trong tệp!")
                return

            if st_name: 
                self.txt_station.delete(0, tk.END); self.txt_station.insert(0, st_name)
            if rv_name: 
                self.txt_river.delete(0, tk.END); self.txt_river.insert(0, rv_name)

            sorted_years = sorted(self.years_data.keys())
            all_list = [self.years_data[y] for y in sorted_years]
            self.all_data = np.concatenate(all_list)

            # Tính giá trị Min/Max thực tế
            self.actual_max = float(np.max(self.all_data))
            self.actual_min = float(np.min(self.all_data))
            unit = self.txt_unit.get()

            # Hiển thị thông tin lên giao diện
            f_display = os.path.basename(files[0]) if len(files) == 1 else f"{len(files)} tệp đã chọn"
            self.lbl_file.config(text=f_display, fg="black")
            self.lbl_years_info.config(
                text=f"Các năm: {', '.join(map(str, sorted_years))} | Tổng số giờ: {len(self.all_data):,} giờ",
                fg="blue"
            )
            self.lbl_actual_stat.config(
                text=f"Số liệu đo thực tế: Max = {self.actual_max:.1f} {unit}  |  Min = {self.actual_min:.1f} {unit}  |  Tổng: {len(self.all_data):,} giờ",
                fg="#d32f2f"
            )

            # TỰ ĐỘNG ĐIỀN CẤP MAX VÀ CẤP MIN VÀO 2 Ô NHẬP LIỆU
            self.auto_fill_min_max()

            messagebox.showinfo(
                "Đọc thành công", 
                f"Đã nạp thành công {len(sorted_years)} năm: {sorted_years}\n"
                f"Tổng số giờ tích lũy: {len(self.all_data):,} giờ.\n\n"
                f"Mực nước Max thực tế: {self.actual_max:.1f} {unit}\n"
                f"Mực nước Min thực tế: {self.actual_min:.1f} {unit}\n\n"
                f"-> Đã tự động điền Cấp Max = {self.txt_hmax.get()} và Cấp Min = {self.txt_hmin.get()}."
            )

        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể đọc file: {e}")

    def compute_multi_year_table(self):
        """Tính bảng tổng hợp nhiều năm chuẩn xác theo cấu trúc ảnh mẫu"""
        if len(self.years_data) == 0:
            messagebox.showwarning("Cảnh báo", "Vui lòng nạp file Excel trước!")
            return False

        try:
            step = float(self.txt_step.get())
            h_max = float(self.txt_hmax.get())
            h_min = float(self.txt_hmin.get())
        except ValueError:
            messagebox.showerror("Lỗi", "Step, Cấp Max, Cấp Min phải là số hợp lệ!")
            return False

        if h_max <= h_min:
            messagebox.showerror("Lỗi", "Cấp Max phải lớn hơn Cấp Min!")
            return False
        if step <= 0:
            messagebox.showerror("Lỗi", "Bước cấp (Step) phải lớn hơn 0!")
            return False

        # Các cấp mực nước (Lower bounds) từ H_max xuống H_min
        lower_bounds = np.arange(h_max, h_min - step / 2.0, -step)
        upper_bounds = lower_bounds + (step - 1)

        sorted_years = sorted(self.years_data.keys())
        total_hours_all = len(self.all_data)

        # Tính số lần xuất hiện cho từng năm
        year_counts = {}
        for yr in sorted_years:
            vals = self.years_data[yr]
            c_list = []
            for i, l in enumerate(lower_bounds):
                if i == 0:
                    c = np.sum(vals >= l)
                elif i == len(lower_bounds) - 1:
                    c = np.sum(vals < lower_bounds[i - 1])  # Bao trọn các giá trị đáy
                else:
                    c = np.sum((vals >= l) & (vals < lower_bounds[i - 1]))
                c_list.append(c)
            year_counts[yr] = np.array(c_list)

        # Tổng số giờ theo từng cấp của tất cả các năm
        total_bracket = np.zeros(len(lower_bounds), dtype=int)
        for yr in sorted_years:
            total_bracket += year_counts[yr]

        freq_pct = (total_bracket / total_hours_all) * 100.0
        cum_hours = np.cumsum(total_bracket)
        cum_freq_pct = (cum_hours / total_hours_all) * 100.0

        # Lưu lại để tra cứu và vẽ biểu đồ
        self.lower_bounds = lower_bounds
        self.cum_freq_pct = cum_freq_pct

        # Lập DataFrame tổng hợp
        df_dict = {
            'TT': range(1, len(lower_bounds) + 1),
            'H_max': [int(u) for u in upper_bounds],
            'H_min': [int(l) for l in lower_bounds]
        }
        for yr in sorted_years:
            df_dict[str(yr)] = year_counts[yr]

        df_dict['TS_TongSo'] = total_bracket
        df_dict['TS_TanSuat'] = np.round(freq_pct, 3)
        df_dict['TSLT_TongSo'] = cum_hours
        df_dict['TSLT_TanSuat'] = np.round(cum_freq_pct, 3)

        self.df_multi = pd.DataFrame(df_dict)
        return True

    def lookup_water_level(self):
        """Tính toán mực nước H tương ứng với tần suất P bất kỳ (nội suy tuyến tính)"""
        if self.cum_freq_pct is None:
            if not self.compute_multi_year_table(): return

        try:
            target_p = float(self.txt_lookup_p.get())
            if not (0.0 <= target_p <= 100.0):
                messagebox.showerror("Lỗi", "Tần suất P phải nằm trong khoảng từ 0% đến 100%!")
                return
        except ValueError:
            messagebox.showerror("Lỗi", "Vui lòng nhập giá trị tần suất là số hợp lệ!")
            return

        # Nội suy tuyến tính H từ P
        p_uniq, u_idx = np.unique(self.cum_freq_pct, return_index=True)
        h_uniq = self.lower_bounds[u_idx]

        h_result = float(np.interp(target_p, p_uniq, h_uniq))
        unit = self.txt_unit.get()
        self.lbl_lookup_result.config(text=f"Mực nước H({target_p:.2f}%) = {h_result:.2f} {unit}")

        # Nếu cửa sổ biểu đồ đang mở, vẽ đường dóng đỏ đánh dấu
        if self.plot_win and self.plot_win.winfo_exists():
            self.draw_lookup_line_on_plot(target_p, h_result)

    def draw_lookup_line_on_plot(self, p_val, h_val):
        """Vẽ đường dóng nét đứt màu đỏ trên biểu đồ khi người dùng tra cứu"""
        if hasattr(self, 'ax_plot') and self.ax_plot:
            if hasattr(self, 'lookup_lines'):
                for item in self.lookup_lines: item.remove()
            self.lookup_lines = []

            l1 = self.ax_plot.axvline(p_val, color='red', linestyle='--', linewidth=1.0, zorder=6)
            l2 = self.ax_plot.axhline(h_val, color='red', linestyle='--', linewidth=1.0, zorder=6)
            pt = self.ax_plot.plot(p_val, h_val, 'ro', markersize=6, zorder=7)[0]
            txt = self.ax_plot.text(p_val + 1.5, h_val + 5, f"P={p_val:.1f}%\nH={h_val:.2f} {self.txt_unit.get()}",
                                    color='red', fontweight='bold', fontsize=9, zorder=8)
            self.lookup_lines.extend([l1, l2, pt, txt])
            self.canvas_plot.draw()

    def show_multi_year_table(self):
        """Hiển thị cửa sổ bảng nhiều năm đúng như ảnh mẫu"""
        if not self.compute_multi_year_table(): return

        if self.tbl_win is not None and self.tbl_win.winfo_exists():
            self.tbl_win.destroy()

        self.tbl_win = tk.Toplevel(self.root)
        station = self.txt_station.get()
        river = self.txt_river.get()
        unit = self.txt_unit.get()
        sorted_years = sorted(self.years_data.keys())
        yr_str = f"{sorted_years[0]} - {sorted_years[-1]}" if len(sorted_years) > 1 else str(sorted_years[0])

        self.tbl_win.title(f"Bảng tổng hợp tần suất mực nước giờ - Trạm {station}")
        self.tbl_win.geometry("980x620")

        tk.Label(self.tbl_win, text=f"BẢNG TÍNH TẦN SUẤT VÀ TẦN SUẤT LŨY TÍCH MỰC NƯỚC GIỜ", 
                 font=("Times New Roman", 13, "bold"), fg="#154360").pack(pady=(8, 2))
        tk.Label(self.tbl_win, text=f"TRẠM {station.upper()} - {river.upper()} (Giai đoạn: {yr_str})", 
                 font=("Times New Roman", 10, "italic bold")).pack(pady=(0, 6))

        tree_frame = tk.Frame(self.tbl_win)
        tree_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=4)

        base_cols = ["tt", "h_max", "h_min"]
        year_cols = [f"y_{y}" for y in sorted_years]
        stat_cols = ["ts_tong", "ts_pct", "tslt_tong", "tslt_pct"]
        all_cols = base_cols + year_cols + stat_cols

        tree = ttk.Treeview(tree_frame, columns=all_cols, show="headings", height=18)

        tree.heading("tt", text="TT")
        tree.heading("h_max", text=f"Cận trên ({unit})")
        tree.heading("h_min", text=f"Cận dưới ({unit})")
        tree.column("tt", width=45, anchor=tk.CENTER)
        tree.column("h_max", width=75, anchor=tk.CENTER)
        tree.column("h_min", width=75, anchor=tk.CENTER)

        for y in sorted_years:
            tree.heading(f"y_{y}", text=str(y))
            tree.column(f"y_{y}", width=60, anchor=tk.E)

        tree.heading("ts_tong", text="Tổng số")
        tree.heading("ts_pct", text="Tần suất %")
        tree.heading("tslt_tong", text="Tổng số L.Tích")
        tree.heading("tslt_pct", text="Tần suất L.Tích %")

        tree.column("ts_tong", width=75, anchor=tk.E)
        tree.column("ts_pct", width=85, anchor=tk.E)
        tree.column("tslt_tong", width=95, anchor=tk.E)
        tree.column("tslt_pct", width=110, anchor=tk.E)

        for _, r in self.df_multi.iterrows():
            row_vals = [
                int(r['TT']),
                int(r['H_max']),
                int(r['H_min'])
            ]
            for y in sorted_years:
                row_vals.append(f"{int(r[str(y)]):,}")
            row_vals.extend([
                f"{int(r['TS_TongSo']):,}",
                f"{r['TS_TanSuat']:.3f}",
                f"{int(r['TSLT_TongSo']):,}",
                f"{r['TSLT_TanSuat']:.3f}"
            ])
            tree.insert("", tk.END, values=row_vals)

        sb_y = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=tree.yview)
        sb_x = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=tree.xview)
        tree.configure(yscroll=sb_y.set, xscroll=sb_x.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb_y.pack(side=tk.RIGHT, fill=tk.Y)
        sb_x.pack(side=tk.BOTTOM, fill=tk.X)

        def export_to_excel_file():
            f_save = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel file", "*.xlsx")])
            if not f_save: return
            self.export_excel_format(f_save, sorted_years, station, river, yr_str, unit)
            messagebox.showinfo("Thành công", f"Đã xuất bảng tính chuẩn xác ra file:\n{f_save}")

        def copy_table():
            out = f"TT\tCan tren\tCan duoi\t" + "\t".join(map(str, sorted_years)) + "\tTong so\tTan suat %\tTong so LT\tTan suat LT %\n"
            for _, r in self.df_multi.iterrows():
                y_str_val = "\t".join([str(int(r[str(y)])) for y in sorted_years])
                out += f"{int(r['TT'])}\t{int(r['H_max'])}\t{int(r['H_min'])}\t{y_str_val}\t{int(r['TS_TongSo'])}\t{r['TS_TanSuat']:.3f}\t{int(r['TSLT_TongSo'])}\t{r['TSLT_TanSuat']:.3f}\n"
            self.root.clipboard_clear()
            self.root.clipboard_append(out)
            messagebox.showinfo("Thành công", "Đã sao chép bảng vào Clipboard! Bạn có thể dán (Ctrl+V) vào Excel.")

        btn_box = tk.Frame(self.tbl_win, pady=8)
        btn_box.pack(fill=tk.X)
        tk.Button(btn_box, text="📋 Sao chép ra Excel (Clipboard)", bg="#1976d2", fg="white", font=("Arial", 9, "bold"),
                  padx=12, pady=5, command=copy_table).pack(side=tk.LEFT, padx=15)
        tk.Button(btn_box, text="💾 Xuất file Excel chuẩn như ảnh (.xlsx)", bg="#27ae60", fg="white", font=("Arial", 9, "bold"),
                  padx=12, pady=5, command=export_to_excel_file).pack(side=tk.RIGHT, padx=15)

    def export_excel_format(self, filepath, years, station, river, yr_str, unit):
        """Xuất file Excel gộp ô 2 tầng đúng định dạng ảnh mẫu"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TanSuat_LuyTich"
        ws.views.sheetView[0].showGridLines = True

        total_cols = 7 + len(years)

        f_title = Font(name="Times New Roman", size=13, bold=True, color="154360")
        f_sub = Font(name="Times New Roman", size=11, bold=True, italic=True)
        f_head = Font(name="Times New Roman", size=10, bold=True)
        f_body = Font(name="Times New Roman", size=10)

        fill_head = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")
        fill_ts = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
        fill_tslt = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='A6ACAF'), right=Side(style='thin', color='A6ACAF'),
            top=Side(style='thin', color='A6ACAF'), bottom=Side(style='thin', color='A6ACAF')
        )

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        ws.cell(row=1, column=1, value=f"BẢNG TÍNH TẦN SUẤT VÀ TẦN SUẤT LŨY TÍCH MỰC NƯỚC GIỜ TRẠM {station.upper()} - {river.upper()}").font = f_title
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        ws.cell(row=2, column=1, value=f"(Giai đoạn: {yr_str})").font = f_sub
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A4:A5")
        ws.cell(row=4, column=1, value="TT\n(No.)").font = f_head

        ws.merge_cells("B4:C4")
        ws.cell(row=4, column=2, value=f"Cấp mực nước\n(Water level)\n({unit})").font = f_head

        start_yr_col = 4
        end_yr_col = 3 + len(years)
        ws.merge_cells(start_row=4, start_column=start_yr_col, end_row=4, end_column=end_yr_col)
        ws.cell(row=4, column=start_yr_col, value="NĂM").font = f_head
        for idx, yr in enumerate(years):
            ws.cell(row=5, column=start_yr_col + idx, value=str(yr)).font = f_head

        ts_col = end_yr_col + 1
        ws.merge_cells(start_row=4, start_column=ts_col, end_row=4, end_column=ts_col+1)
        ws.cell(row=4, column=ts_col, value="Tần suất %\n(Frequence)").font = f_head
        ws.cell(row=5, column=ts_col, value="Tổng số\n(Total)").font = f_head
        ws.cell(row=5, column=ts_col+1, value="Tần suất\n(Frequence)").font = f_head

        tslt_col = ts_col + 2
        ws.merge_cells(start_row=4, start_column=tslt_col, end_row=4, end_column=tslt_col+1)
        ws.cell(row=4, column=tslt_col, value="Tần suất lũy tích %\n(Accumulative frequence)").font = f_head
        ws.cell(row=5, column=tslt_col, value="Tổng số\n(Total)").font = f_head
        ws.cell(row=5, column=tslt_col+1, value="Tần suất\n(Frequence)").font = f_head

        for r in range(4, 6):
            for c in range(1, total_cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
                if c in [ts_col, ts_col+1]: cell.fill = fill_ts
                elif c in [tslt_col, tslt_col+1]: cell.fill = fill_tslt
                else: cell.fill = fill_head

        cur_row = 6
        for _, row in self.df_multi.iterrows():
            ws.cell(row=cur_row, column=1, value=int(row['TT'])).alignment = Alignment(horizontal="center")
            ws.cell(row=cur_row, column=2, value=int(row['H_max'])).alignment = Alignment(horizontal="center")
            ws.cell(row=cur_row, column=3, value=int(row['H_min'])).alignment = Alignment(horizontal="center")

            for i_y, yr in enumerate(years):
                c_cell = ws.cell(row=cur_row, column=start_yr_col + i_y, value=int(row[str(yr)]))
                c_cell.alignment = Alignment(horizontal="right")
                c_cell.number_format = "#,##0"

            ws.cell(row=cur_row, column=ts_col, value=int(row['TS_TongSo'])).alignment = Alignment(horizontal="right")
            ws.cell(row=cur_row, column=ts_col).number_format = "#,##0"
            ws.cell(row=cur_row, column=ts_col+1, value=float(row['TS_TanSuat'])).alignment = Alignment(horizontal="right")
            ws.cell(row=cur_row, column=ts_col+1).number_format = "0.000"

            ws.cell(row=cur_row, column=tslt_col, value=int(row['TSLT_TongSo'])).alignment = Alignment(horizontal="right")
            ws.cell(row=cur_row, column=tslt_col).number_format = "#,##0"
            ws.cell(row=cur_row, column=tslt_col+1, value=float(row['TSLT_TanSuat'])).alignment = Alignment(horizontal="right")
            ws.cell(row=cur_row, column=tslt_col+1).number_format = "0.000"

            for c in range(1, total_cols + 1):
                cell = ws.cell(row=cur_row, column=c)
                cell.font = f_body
                cell.border = thin_border
                if c in [ts_col, ts_col+1]: cell.fill = fill_ts
                elif c in [tslt_col, tslt_col+1]: cell.fill = fill_tslt

            cur_row += 1

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 11
        ws.column_dimensions['A'].width = 7
        ws.column_dimensions['B'].width = 8
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions[get_column_letter(tslt_col+1)].width = 14

        wb.save(filepath)

    def plot_graph(self):
        """Vẽ đường tần suất lũy tích tổng hợp nhiều năm chuẩn như ảnh mẫu"""
        if not self.compute_multi_year_table(): return

        if self.plot_win is not None and self.plot_win.winfo_exists():
            self.plot_win.destroy()
        plt.close('all')

        station = self.txt_station.get()
        river = self.txt_river.get()
        unit = self.txt_unit.get()

        self.plot_win = tk.Toplevel(self.root)
        self.plot_win.title(f"Biểu đồ tần suất lũy tích - Trạm {station}")
        self.plot_win.geometry("1140x700")

        fig, ax = plt.subplots(figsize=(11.5, 6.6), dpi=100)
        self.ax_plot = ax

        x = self.cum_freq_pct
        y = self.lower_bounds

        # Vẽ đường tần suất lũy tích tổng hợp
        ax.plot(x, y, marker='d', markersize=4.2, markeredgewidth=0.7,
                color="#1b4f72", linewidth=1.2, clip_on=False)

        # Tiêu đề
        title_text = f"ĐƯỜNG TẦN SUẤT LŨY TÍCH MỰC NƯỚC GIỜ\nTRẠM {station.upper()} - SÔNG {river.upper()}"
        ax.set_title(title_text, fontsize=13, fontweight='bold', pad=15)

        ax.set_xlabel("Tần suất (%)", fontsize=11, fontweight='bold', labelpad=8)
        ax.set_ylabel(f"Mực nước ({unit})", fontsize=11, fontweight='bold', labelpad=8)

        # Trục hoành từ 0 đến 100%, Trục tung bo tròn bội số 50
        ax.set_xlim(0, 100)
        y_lim_bot = int(np.floor(np.min(y) / 50.0) * 50)
        y_lim_top = int(np.ceil(np.max(y) / 50.0) * 50)
        ax.set_ylim(y_lim_bot, y_lim_top)

        # Vạch chia và hệ lưới ô vuông caro
        ax.xaxis.set_major_locator(MultipleLocator(5))
        ax.xaxis.set_minor_locator(MultipleLocator(1))
        ax.yaxis.set_major_locator(MultipleLocator(50))
        ax.yaxis.set_minor_locator(MultipleLocator(10))

        ax.grid(which="major", linestyle="-", linewidth=0.75, color="#D5D8DC")
        ax.grid(which="minor", linestyle=":", linewidth=0.45, color="#EBEDEF")

        for spine in ax.spines.values():
            spine.set_color("#2C3E50")
            spine.set_linewidth(1.0)

        fig.tight_layout()

        canvas = FigureCanvasTkAgg(fig, master=self.plot_win)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.canvas_plot = canvas

        toolbar = NavigationToolbar2Tk(canvas, self.plot_win)
        toolbar.update()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        # Tra cứu tần suất hiện thời để hiển thị đường dóng mẫu
        self.lookup_water_level()
        self.show_multi_year_table()


if __name__ == "__main__":
    root = tk.Tk()
    app = MultiYearWaterLevelApp(root)
    root.mainloop()