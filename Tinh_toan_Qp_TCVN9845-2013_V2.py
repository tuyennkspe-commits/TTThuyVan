"""Ứng dụng Qp cho 0 < F < 100 km², một tệp độc lập.
Rà soát: kiểm tra số hữu hạn, δ≤1, tra biên có cảnh báo, xuất đúng bộ đầu vào.
Công thức (1), (8)–(10) đã đối chiếu tài liệu TCVN 9845:2013 người dùng cung cấp.
Bảng A.1 dùng đủ cấp mưa, không nội suy theo trục tự đặt. Bảng A.3: 21 ô khác biệt cần xác minh.
Một số bảng Ap có điểm không đơn điệu và sai khác dấu thập phân giữa mã/Word.
Không tự chọn giá trị ở ô chưa thống nhất; cho phép nhập hệ số đã xác minh.
Bảng A.2 vùng 18, Φs=40 thiếu số trong Word: cần nhập ts nếu Φs>35.
Chỉ hỗ trợ lưu vực có lòng dẫn rõ ràng (L, Jl > 0); không áp dụng trường hợp Φl=0 tại 5.2.3.
Cài thư viện: python -m pip install numpy pandas scipy matplotlib openpyxl
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import sys
import json
import numpy as np
import pandas as pd
from scipy.stats import norm, pearson3, skew
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.figure import Figure

plt.rcParams['font.family'] = 'serif'
plt.rcParams['axes.unicode_minus'] = False

# =============================================================================
# CƠ SỞ DỮ LIỆU BẢNG TRA CHUẨN TCVN 9845:2013
# =============================================================================

SOIL_TABLES = {'II': [[0.96, 0.94, 0.93, 0.9, 0.88, 0.85, 0.81, 0.78, 0.76, 0.74, 0.67, 0.65, 0.6], [0.97, 0.96, 0.94, 0.91, 0.9, 0.87, 0.85, 0.78, 0.76, 0.74, 0.67, 0.65, 0.6], [0.97, 0.96, 0.95, 0.93, 0.92, 0.9, 0.89, 0.85, 0.83, 0.81, 0.75, 0.73, 0.7], [0.97, 0.96, 0.96, 0.95, 0.94, 0.93, 0.92, 0.89, 0.89, 0.85, 0.85, 0.85, 0.85], [0.97, 0.96, 0.96, 0.96, 0.95, 0.95, 0.94, 0.93, 0.93, 0.88, 0.88, 0.88, 0.86], [0.97, 0.96, 0.96, 0.96, 0.95, 0.95, 0.95, 0.93, 0.93, 0.91, 0.91, 0.91, 0.91], [0.97, 0.96, 0.96, 0.96, 0.95, 0.95, 0.95, 0.93, 0.93, 0.91, 0.91, 0.91, 0.91]], 'III': [[0.94, 0.89, 0.86, 0.8, 0.77, 0.74, 0.65, 0.6, 0.58, 0.55, 0.53, 0.53, 0.5], [0.95, 0.93, 0.9, 0.85, 0.81, 0.77, 0.72, 0.63, 0.62, 0.6, 0.55, 0.55, 0.55], [0.95, 0.93, 0.91, 0.88, 0.86, 0.82, 0.79, 0.72, 0.68, 0.68, 0.63, 0.63, 0.62], [0.95, 0.93, 0.92, 0.91, 0.9, 0.85, 0.85, 0.75, 0.72, 0.73, 0.73, 0.73, 0.65], [0.95, 0.93, 0.921, 0.91, 0.9, 0.85, 0.85, 0.77, 0.74, 0.74, 0.69, 0.69, 0.67], [0.95, 0.93, 0.921, 0.912, 0.9, 0.855, 0.87, 0.78, 0.76, 0.75, 0.71, 0.71, 0.69], [0.95, 0.93, 0.922, 0.912, 0.902, 0.88, 0.89, 0.79, 0.77, 0.77, 0.73, 0.73, 0.7], [0.95, 0.93, 0.922, 0.913, 0.902, 0.885, 0.895, 0.8, 0.79, 0.78, 0.75, 0.75, 0.71], [0.95, 0.93, 0.923, 0.913, 0.91, 0.89, 0.94, 0.8, 0.8, 0.79, 0.75, 0.75, 0.71], [0.95, 0.93, 0.923, 0.913, 0.91, 0.89, 0.94, 0.8, 0.8, 0.8, 0.76, 0.76, 0.71], [0.95, 0.93, 0.923, 0.913, 0.91, 0.89, 0.94, 0.8, 0.8, 0.8, 0.76, 0.76, 0.71], [0.95, 0.93, 0.923, 0.913, 0.91, 0.89, 0.94, 0.8, 0.8, 0.8, 0.76, 0.76, 0.71]], 'IV': [[0.9, 0.81, 0.76, 0.66, 0.65, 0.6, 0.55, 0.51, 0.5, 0.5, 0.44, 0.4, 0.37], [0.9, 0.84, 0.8, 0.76, 0.68, 0.64, 0.62, 0.58, 0.56, 0.55, 0.52, 0.5, 0.46], [0.9, 0.88, 0.85, 0.82, 0.78, 0.75, 0.72, 0.66, 0.63, 0.6, 0.6, 0.57, 0.55], [0.9, 0.88, 0.822, 0.823, 0.79, 0.78, 0.74, 0.7, 0.67, 0.67, 0.65, 0.6, 0.58], [0.9, 0.88, 0.822, 0.825, 0.79, 0.79, 0.76, 0.74, 0.7, 0.7, 0.69, 0.65, 0.61], [0.9, 0.88, 0.828, 0.828, 0.8, 0.8, 0.78, 0.76, 0.72, 0.71, 0.71, 0.67, 0.64], [0.9, 0.88, 0.828, 0.83, 0.82, 0.82, 0.81, 0.77, 0.74, 0.73, 0.72, 0.69, 0.65], [0.9, 0.88, 0.86, 0.84, 0.84, 0.84, 0.83, 0.77, 0.75, 0.75, 0.73, 0.71, 0.67], [0.9, 0.88, 0.86, 0.85, 0.84, 0.84, 0.83, 0.78, 0.76, 0.77, 0.73, 0.72, 0.68], [0.9, 0.88, 0.87, 0.86, 0.84, 0.84, 0.83, 0.78, 0.76, 0.77, 0.73, 0.72, 0.69], [0.9, 0.88, 0.87, 0.86, 0.84, 0.84, 0.83, 0.78, 0.76, 0.77, 0.73, 0.72, 0.69], [0.9, 0.88, 0.87, 0.86, 0.84, 0.84, 0.83, 0.78, 0.76, 0.77, 0.73, 0.72, 0.69]], 'V': [[0.68, 0.46, 0.35, 0.26, 0.24, 0.22, 0.22, 0.2, 0.18, 0.18, 0.17, 0.16, 0.15], [0.71, 0.56, 0.46, 0.41, 0.4, 0.34, 0.32, 0.28, 0.27, 0.25, 0.23, 0.22, 0.2], [0.75, 0.65, 0.59, 0.5, 0.48, 0.46, 0.46, 0.42, 0.45, 0.38, 0.34, 0.32, 0.3], [0.76, 0.68, 0.63, 0.543, 0.5, 0.5, 0.5, 0.46, 0.49, 0.43, 0.38, 0.36, 0.34], [0.77, 0.71, 0.66, 0.58, 0.58, 0.54, 0.54, 0.49, 0.51, 0.46, 0.41, 0.4, 0.36], [0.77, 0.73, 0.66, 0.58, 0.58, 0.54, 0.56, 0.49, 0.54, 0.46, 0.41, 0.43, 0.37], [0.78, 0.75, 0.7, 0.65, 0.64, 0.57, 0.57, 0.53, 0.55, 0.52, 0.46, 0.46, 0.4], [0.79, 0.76, 0.72, 0.67, 0.67, 0.58, 0.58, 0.54, 0.55, 0.53, 0.47, 0.47, 0.41], [0.79, 0.77, 0.73, 0.68, 0.68, 0.6, 0.6, 0.55, 0.55, 0.53, 0.48, 0.48, 0.41], [0.79, 0.78, 0.73, 0.7, 0.7, 0.6, 0.6, 0.55, 0.55, 0.53, 0.49, 0.5, 0.41], [0.79, 0.78, 0.73, 0.7, 0.7, 0.6, 0.6, 0.55, 0.55, 0.53, 0.5, 0.5, 0.41], [0.79, 0.78, 0.73, 0.7, 0.7, 0.6, 0.6, 0.55, 0.55, 0.53, 0.5, 0.5, 0.41]], 'VI': [[0.25, 0.25, 0.25, 0.25, 0.25, 0.2, 0.2, 0.15, 0.15, 0.15, 0.1, 0.1, 0.1]]}
SOIL_H_LABELS = {'II': ['<100', '101-150', '151-200', '201-250', '251-300', '301-400', '>400'], 'III': ['<100', '101-150', '151-200', '201-250', '251-300', '301-350', '351-400', '401-450', '451-500', '501-550', '551-600', '>600'], 'IV': ['<100', '101-150', '151-200', '201-250', '251-300', '301-350', '351-400', '401-450', '451-500', '501-550', '551-600', '>600'], 'V': ['<100', '101-150', '151-200', '201-250', '251-300', '301-350', '351-400', '401-450', '451-500', '501-550', '551-600', '>600'], 'VI': ['Mọi lượng mưa']}
SOIL_H_BOUNDS = {'II': [100, 150, 200, 250, 300, 400], 'III': [100, 150, 200, 250, 300, 350, 400, 450, 500, 550, 600], 'IV': [100, 150, 200, 250, 300, 350, 400, 450, 500, 550, 600], 'V': [100, 150, 200, 250, 300, 350, 400, 450, 500, 550, 600], 'VI': []}
# The supplied Word table gives area groups but no individual area abscissas within them.
PHI_AREA_GROUPS = [(0,.1,0,5),(.1,1,5,7),(1,10,7,10),(10,100,10,12)]

FS_AXIS = [1.0, 1.5, 2.0, 2.5, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0, 12.0, 15.0, 17.0, 20.0, 25.0, 30.0, 35.0, 40.0]
TS_GRID = [
    [9.6, 9.7, 9.7, 9.0, 9.6, 9.6, 16.0, 8.4, 9.7, 9.8, 9.5, 10.0, 9.8, 8.7, 8.5, 8.7, 9.3, 9.2],
    [10.0, 10.0, 10.0, 9.0, 10.0, 10.0, 18.0, 8.5, 10.0, 10.0, 10.0, 13.0, 10.0, 9.0, 8.7, 9.0, 9.4, 9.3],
    [17.0, 15.0, 17.0, 9.5, 14.0, 10.0, 25.0, 9.0, 13.0, 15.0, 20.0, 20.0, 15.0, 9.3, 9.3, 9.5, 9.7, 9.5],
    [24.0, 22.0, 20.0, 10.0, 20.0, 15.0, 32.0, 10.0, 15.0, 18.0, 28.0, 23.0, 20.0, 9.5, 9.5, 9.6, 10.0, 9.7],
    [35.0, 28.0, 25.0, 18.0, 30.0, 22.0, 37.0, 20.0, 18.0, 25.0, 35.0, 30.0, 25.0, 11.0, 10.0, 12.0, 20.0, 12.0],
    [40.0, 37.0, 32.0, 22.0, 35.0, 30.0, 42.0, 30.0, 25.0, 40.0, 55.0, 35.0, 30.0, 20.0, 20.0, 20.0, 25.0, 20.0],
    [53.0, 45.0, 50.0, 30.0, 44.0, 38.0, 50.0, 40.0, 30.0, 45.0, 65.0, 50.0, 40.0, 30.0, 25.0, 30.0, 35.0, 23.0],
    [62.0, 60.0, 60.0, 45.0, 60.0, 50.0, 55.0, 55.0, 40.0, 60.0, 72.0, 60.0, 55.0, 35.0, 32.0, 37.0, 40.0, 30.0],
    [70.0, 70.0, 72.0, 60.0, 75.0, 70.0, 65.0, 65.0, 65.0, 75.0, 80.0, 75.0, 65.0, 50.0, 50.0, 50.0, 60.0, 40.0],
    [75.0, 78.0, 80.0, 68.0, 85.0, 78.0, 75.0, 70.0, 70.0, 85.0, 90.0, 80.0, 70.0, 70.0, 65.0, 65.0, 70.0, 60.0],
    [80.0, 87.0, 90.0, 80.0, 90.0, 82.0, 85.0, 80.0, 80.0, 90.0, 95.0, 87.0, 82.0, 80.0, 70.0, 78.0, 80.0, 70.0],
    [90.0, 95.0, 100.0, 86.0, 95.0, 88.0, 90.0, 90.0, 95.0, 95.0, 110.0, 105.0, 90.0, 85.0, 80.0, 80.0, 90.0, 80.0],
    [100.0, 115.0, 120.0, 95.0, 100.0, 93.0, 100.0, 115.0, 115.0, 110.0, 130.0, 120.0, 100.0, 90.0, 90.0, 90.0, 97.0, 83.0],
    [130.0, 150.0, 150.0, 120.0, 120.0, 120.0, 125.0, 135.0, 135.0, 135.0, 160.0, 150.0, 125.0, 115.0, 125.0, 115.0, 120.0, 100.0],
    [160.0, 165.0, 180.0, 165.0, 170.0, 150.0, 165.0, 190.0, 170.0, 170.0, 200.0, 190.0, 160.0, 160.0, 150.0, 140.0, 145.0, 130.0],
    [200.0, 220.0, 230.0, 200.0, 200.0, 185.0, 205.0, 235.0, 220.0, 220.0, 230.0, 235.0, 200.0, 200.0, 190.0, 175.0, 190.0, 165.0],
    [260.0, 280.0, 265.0, 235.0, 260.0, 230.0, 250.0, 305.0, 290.0, 265.0, 300.0, 300.0, 250.0, 250.0, 250.0, 225.0, 240.0, 230.0],
    [325.0, 360.0, 365.0, 320.0, 320.0, 310.0, 320.0, 370.0, 370.0, 335.0, 400.0, 380.0, 330.0, 320.0, 320.0, 285.0, 320.0, 300.0],
    [370.0, 430.0, 435.0, 400.0, 370.0, 370.0, 400.0, 480.0, 430.0, 345.0, 470.0, 450.0, 400.0, 400.0, 400.0, 355.0, 380.0, 370.0],
    [470.0, 530.0, 520.0, 470.0, 480.0, 470.0, 570.0, 495.0, 520.0, 410.0, 560.0, 540.0, 510.0, 480.0, 490.0, 425.0, 465.0, 460.0]
]

FL_COLS = [0.0, 1.0, 5.0, 10.0, 20.0, 30.0, 40.0, 50.0, 60.0, 70.0, 80.0, 90.0, 100.0, 150.0, 200.0, 220.0]
TS_LEVELS = [20.0, 30.0, 60.0, 90.0, 180.0]

AP_TABLES = {
    1: [
        [0.2800, 0.2600, 0.2180, 0.1520, 0.1120, 0.0920, 0.0760, 0.0640, 0.0540, 0.0470, 0.0400, 0.0350, 0.0300, 0.0180, 0.0150, 0.0130],
        [0.2100, 0.1900, 0.1600, 0.1360, 0.1040, 0.0850, 0.0720, 0.0610, 0.0520, 0.0450, 0.0380, 0.0330, 0.0290, 0.0170, 0.0140, 0.0125],
        [0.1500, 0.1430, 0.1250, 0.1110, 0.0910, 0.0760, 0.0650, 0.0550, 0.0470, 0.0400, 0.0340, 0.0300, 0.0260, 0.0160, 0.0130, 0.0120],
        [0.1140, 0.1120, 0.1020, 0.0930, 0.0770, 0.0650, 0.0560, 0.0480, 0.0410, 0.0350, 0.0310, 0.0270, 0.0240, 0.0150, 0.0120, 0.0115],
        [0.0720, 0.0710, 0.0570, 0.0630, 0.0550, 0.0480, 0.0430, 0.0370, 0.0330, 0.0290, 0.0250, 0.0220, 0.0210, 0.0140, 0.0115, 0.0110]
    ],
    2: [
        [0.1170, 0.1140, 0.1040, 0.0930, 0.0870, 0.0650, 0.0550, 0.0470, 0.0400, 0.0340, 0.0300, 0.0260, 0.0240, 0.0180, 0.0150, 0.0140],
        [0.1000, 0.0980, 0.0910, 0.0830, 0.0700, 0.0600, 0.0520, 0.0440, 0.0380, 0.0330, 0.0280, 0.0250, 0.0230, 0.0175, 0.0140, 0.0130],
        [0.0820, 0.0810, 0.0760, 0.0700, 0.0600, 0.0520, 0.0450, 0.0390, 0.0340, 0.0300, 0.0270, 0.0240, 0.0220, 0.0160, 0.0130, 0.0125],
        [0.0670, 0.0660, 0.0630, 0.0590, 0.0520, 0.0460, 0.0400, 0.0350, 0.0310, 0.0270, 0.0250, 0.0220, 0.0200, 0.0150, 0.0120, 0.0120],
        [0.0520, 0.0510, 0.0480, 0.0460, 0.0410, 0.0360, 0.0320, 0.0280, 0.0250, 0.0220, 0.0200, 0.0180, 0.0170, 0.0140, 0.0110, 0.0110]
    ],
    3: [
        [0.1590, 0.1530, 0.1370, 0.1220, 0.0985, 0.0831, 0.0708, 0.0618, 0.0544, 0.0492, 0.0450, 0.0410, 0.0378, 0.0281, 0.0218, 0.0183],
        [0.1320, 0.1290, 0.1160, 0.1040, 0.0866, 0.0740, 0.0650, 0.0573, 0.0507, 0.0462, 0.0420, 0.0390, 0.0358, 0.0272, 0.0211, 0.0180],
        [0.0950, 0.0920, 0.0870, 0.0790, 0.0695, 0.0611, 0.0530, 0.0497, 0.0447, 0.0410, 0.0380, 0.0350, 0.0325, 0.0252, 0.0197, 0.0170],
        [0.0730, 0.0680, 0.0659, 0.0612, 0.0549, 0.0500, 0.0443, 0.0414, 0.0384, 0.0355, 0.0330, 0.0307, 0.0292, 0.0228, 0.0185, 0.0160],
        [0.0580, 0.0540, 0.0517, 0.0490, 0.0450, 0.0420, 0.0383, 0.0360, 0.0330, 0.0303, 0.0300, 0.0268, 0.0256, 0.0205, 0.0165, 0.0150]
    ],
    4: [
        [0.2730, 0.2140, 0.1880, 0.1630, 0.1280, 0.1040, 0.0865, 0.0743, 0.0654, 0.0565, 0.0499, 0.0448, 0.0408, 0.0279, 0.0216, 0.0184],
        [0.2000, 0.1840, 0.1630, 0.1420, 0.1153, 0.0950, 0.0816, 0.0703, 0.0615, 0.0545, 0.0479, 0.0429, 0.0390, 0.0269, 0.0212, 0.0182],
        [0.1290, 0.1240, 0.1170, 0.1070, 0.0903, 0.0790, 0.0688, 0.0593, 0.0553, 0.0473, 0.0427, 0.0382, 0.0351, 0.0256, 0.0200, 0.0174],
        [0.1020, 0.0930, 0.0890, 0.0840, 0.0735, 0.0645, 0.0579, 0.0508, 0.0460, 0.0410, 0.0370, 0.0340, 0.0315, 0.0230, 0.0189, 0.0164],
        [0.0720, 0.0710, 0.0670, 0.0630, 0.0555, 0.0503, 0.0456, 0.0413, 0.0378, 0.0328, 0.0315, 0.0310, 0.0275, 0.0210, 0.0178, 0.0155]
    ],
    5: [
        [0.1200, 0.1185, 0.1115, 0.1087, 0.0940, 0.0786, 0.0690, 0.0630, 0.0525, 0.0457, 0.0397, 0.0347, 0.0304, 0.0195, 0.0140, 0.0130],
        [0.1120, 0.1100, 0.1035, 0.0965, 0.0840, 0.0733, 0.0638, 0.0560, 0.0485, 0.0423, 0.0370, 0.0320, 0.0280, 0.0169, 0.0133, 0.0124],
        [0.0980, 0.0965, 0.0855, 0.0815, 0.0748, 0.0655, 0.0577, 0.0506, 0.0445, 0.0393, 0.0345, 0.0304, 0.0268, 0.0163, 0.0126, 0.0119],
        [0.0830, 0.0817, 0.0775, 0.0726, 0.0642, 0.0565, 0.0500, 0.0443, 0.0390, 0.0345, 0.0310, 0.0276, 0.0247, 0.0152, 0.0118, 0.0114],
        [0.0595, 0.0587, 0.0560, 0.0583, 0.0480, 0.0430, 0.0390, 0.0350, 0.0317, 0.0285, 0.0263, 0.0240, 0.0223, 0.0148, 0.0110, 0.0108]
    ],
    6: [
        [0.1215, 0.1195, 0.1130, 0.1053, 0.0916, 0.0803, 0.0703, 0.0617, 0.0543, 0.0478, 0.0417, 0.0377, 0.0324, 0.0195, 0.0150, 0.0140],
        [0.1135, 0.1117, 0.1060, 0.0870, 0.0865, 0.0757, 0.0666, 0.0585, 0.0515, 0.0452, 0.0397, 0.0350, 0.0310, 0.0189, 0.0145, 0.0135],
        [0.1050, 0.0995, 0.0944, 0.0860, 0.0798, 0.0686, 0.0606, 0.0536, 0.0474, 0.0420, 0.0373, 0.0333, 0.0295, 0.0183, 0.0140, 0.0129],
        [0.0863, 0.0858, 0.0816, 0.0770, 0.0690, 0.0617, 0.0553, 0.0490, 0.0440, 0.0390, 0.0350, 0.0310, 0.0278, 0.0172, 0.0135, 0.0124],
        [0.0645, 0.0637, 0.0610, 0.0580, 0.0513, 0.0457, 0.0407, 0.0363, 0.0323, 0.0292, 0.0265, 0.0242, 0.0222, 0.0167, 0.0130, 0.0120]
    ],
    7: [
        [0.1060, 0.1050, 0.1000, 0.0934, 0.0817, 0.0716, 0.0633, 0.0555, 0.0490, 0.0430, 0.0382, 0.0337, 0.0300, 0.0190, 0.0150, 0.0133],
        [0.0970, 0.0960, 0.0910, 0.0786, 0.0763, 0.0677, 0.0603, 0.0534, 0.0474, 0.0417, 0.0370, 0.0327, 0.0290, 0.0181, 0.0142, 0.0129],
        [0.0850, 0.0840, 0.0800, 0.0757, 0.0676, 0.0606, 0.0540, 0.0482, 0.0430, 0.0380, 0.0340, 0.0303, 0.0272, 0.0175, 0.0135, 0.0125],
        [0.0710, 0.0700, 0.0670, 0.0632, 0.0565, 0.0506, 0.0455, 0.0407, 0.0400, 0.0330, 0.0298, 0.0271, 0.0247, 0.0168, 0.0127, 0.0117],
        [0.0570, 0.0560, 0.0540, 0.0510, 0.0460, 0.0408, 0.0365, 0.0326, 0.0293, 0.0265, 0.0238, 0.0218, 0.0200, 0.0160, 0.0121, 0.0110]
    ],
    8: [
        [0.1620, 0.1560, 0.1360, 0.1210, 0.0963, 0.0805, 0.0676, 0.0572, 0.0483, 0.0422, 0.0375, 0.0334, 0.0298, 0.0240, 0.0170, 0.0160],
        [0.1460, 0.1420, 0.1270, 0.1120, 0.0905, 0.0760, 0.0645, 0.0550, 0.0477, 0.0416, 0.0366, 0.0327, 0.0292, 0.0225, 0.0160, 0.0150],
        [0.1190, 0.1160, 0.1040, 0.0933, 0.0773, 0.0656, 0.0560, 0.0486, 0.0435, 0.0386, 0.0345, 0.0309, 0.0280, 0.0210, 0.0150, 0.0140],
        [0.1010, 0.0987, 0.0910, 0.0824, 0.0693, 0.0593, 0.0513, 0.0445, 0.0394, 0.0352, 0.0320, 0.0293, 0.0265, 0.0190, 0.0140, 0.0130],
        [0.0620, 0.0615, 0.0587, 0.0550, 0.0500, 0.0450, 0.0403, 0.0365, 0.0330, 0.0300, 0.0275, 0.0253, 0.0235, 0.0173, 0.0130, 0.0120]
    ],
    9: [
        [0.1923, 0.1825, 0.1570, 0.1430, 0.1152, 0.0956, 0.0810, 0.0705, 0.0616, 0.0549, 0.0489, 0.0443, 0.0407, 0.0290, 0.0220, 0.0200],
        [0.1912, 0.1555, 0.1395, 0.1233, 0.1030, 0.0868, 0.0762, 0.0663, 0.0587, 0.0527, 0.0469, 0.0425, 0.0390, 0.0279, 0.0210, 0.0190],
        [0.1095, 0.1050, 0.1015, 0.0931, 0.0811, 0.0724, 0.0642, 0.0563, 0.0534, 0.0463, 0.0425, 0.0385, 0.0355, 0.0262, 0.0200, 0.0178],
        [0.0905, 0.0820, 0.0800, 0.0756, 0.0740, 0.0607, 0.0553, 0.0493, 0.0452, 0.0407, 0.0372, 0.0345, 0.0322, 0.0233, 0.0190, 0.0165],
        [0.0640, 0.0635, 0.0610, 0.0572, 0.0510, 0.0468, 0.0433, 0.0396, 0.0367, 0.0336, 0.0317, 0.0300, 0.0280, 0.0220, 0.0178, 0.0155]
    ],
    10: [
        [0.0946, 0.0932, 0.0887, 0.0833, 0.0733, 0.0645, 0.0568, 0.0500, 0.0443, 0.0388, 0.0345, 0.0305, 0.0277, 0.0200, 0.0150, 0.0130],
        [0.0893, 0.0880, 0.0836, 0.0788, 0.0690, 0.0608, 0.0537, 0.0473, 0.0417, 0.0370, 0.0330, 0.0293, 0.0263, 0.0192, 0.0145, 0.0128],
        [0.0806, 0.0796, 0.0757, 0.0710, 0.0628, 0.0555, 0.0487, 0.0433, 0.0383, 0.0340, 0.0303, 0.0270, 0.0246, 0.0183, 0.0140, 0.0125],
        [0.0717, 0.0707, 0.0670, 0.0635, 0.0557, 0.0495, 0.0437, 0.0387, 0.0346, 0.0307, 0.0277, 0.0253, 0.0230, 0.0179, 0.0135, 0.0122],
        [0.0525, 0.0520, 0.0500, 0.0472, 0.0425, 0.0382, 0.0345, 0.0313, 0.0283, 0.0262, 0.0243, 0.0224, 0.0216, 0.0173, 0.0130, 0.0115]
    ],
    11: [
        [0.0888, 0.0862, 0.0800, 0.0714, 0.0607, 0.0524, 0.0461, 0.0406, 0.0364, 0.0330, 0.0304, 0.0280, 0.0267, 0.0216, 0.0182, 0.0161],
        [0.0712, 0.0696, 0.0667, 0.0612, 0.0541, 0.0478, 0.0430, 0.0385, 0.0348, 0.0317, 0.0294, 0.0273, 0.0258, 0.0211, 0.0176, 0.0157],
        [0.0631, 0.0615, 0.0582, 0.0542, 0.0480, 0.0431, 0.0388, 0.0360, 0.0315, 0.0286, 0.0268, 0.0251, 0.0234, 0.0196, 0.0164, 0.0149],
        [0.0518, 0.0508, 0.0479, 0.0459, 0.0403, 0.0364, 0.0327, 0.0304, 0.0283, 0.0261, 0.0255, 0.0233, 0.0222, 0.0185, 0.0157, 0.0143],
        [0.0431, 0.0420, 0.0398, 0.0375, 0.0339, 0.0316, 0.0286, 0.0264, 0.0245, 0.0230, 0.0218, 0.0210, 0.0204, 0.0172, 0.0148, 0.0136]
    ],
    12: [
        [0.0900, 0.0880, 0.0807, 0.0727, 0.0600, 0.0503, 0.0423, 0.0360, 0.0307, 0.0270, 0.0242, 0.0225, 0.0218, 0.0185, 0.0150, 0.0138],
        [0.0790, 0.0755, 0.0705, 0.0647, 0.0550, 0.0466, 0.0397, 0.0344, 0.0297, 0.0260, 0.0237, 0.0220, 0.0213, 0.0175, 0.0142, 0.0134],
        [0.0614, 0.0604, 0.0567, 0.0527, 0.0455, 0.0396, 0.0345, 0.0303, 0.0270, 0.0244, 0.0224, 0.0214, 0.0208, 0.0170, 0.0138, 0.0129],
        [0.0520, 0.0510, 0.0487, 0.0460, 0.0406, 0.0357, 0.0317, 0.0283, 0.0253, 0.0232, 0.0217, 0.0205, 0.0197, 0.0165, 0.0130, 0.0122],
        [0.0410, 0.0404, 0.0387, 0.0365, 0.0327, 0.0295, 0.0265, 0.0243, 0.0222, 0.0207, 0.0197, 0.0188, 0.0185, 0.0153, 0.0120, 0.0115]
    ],
    13: [
        [0.1540, 0.1490, 0.1390, 0.1050, 0.0901, 0.0763, 0.0658, 0.0570, 0.0506, 0.0449, 0.0403, 0.0366, 0.0334, 0.0253, 0.0208, 0.0183],
        [0.1290, 0.1260, 0.1120, 0.0990, 0.0834, 0.0713, 0.0624, 0.0539, 0.0476, 0.0428, 0.0382, 0.0350, 0.0319, 0.0241, 0.0198, 0.0177],
        [0.0975, 0.0954, 0.0878, 0.0808, 0.0694, 0.0611, 0.0534, 0.0477, 0.0427, 0.0383, 0.0351, 0.0319, 0.0294, 0.0227, 0.0185, 0.0168],
        [0.0756, 0.0740, 0.0684, 0.0648, 0.0542, 0.0515, 0.0478, 0.0417, 0.0375, 0.0345, 0.0317, 0.0296, 0.0268, 0.0214, 0.0184, 0.0160],
        [0.0543, 0.0530, 0.0513, 0.0491, 0.0448, 0.0415, 0.0378, 0.0315, 0.0320, 0.0297, 0.0278, 0.0257, 0.0246, 0.0200, 0.0175, 0.0152]
    ],
    14: [
        [0.2300, 0.2150, 0.2070, 0.1750, 0.1190, 0.0937, 0.0756, 0.0622, 0.0517, 0.0435, 0.0370, 0.0315, 0.0273, 0.0185, 0.0140, 0.0120],
        [0.1780, 0.1710, 0.1500, 0.1310, 0.1050, 0.0855, 0.0703, 0.0585, 0.0493, 0.0415, 0.0353, 0.0303, 0.0263, 0.0178, 0.0132, 0.0112],
        [0.1370, 0.1340, 0.1220, 0.1100, 0.0920, 0.0757, 0.0633, 0.0533, 0.0437, 0.0383, 0.0326, 0.0284, 0.0250, 0.0170, 0.0125, 0.0103],
        [0.1100, 0.1070, 0.0970, 0.0900, 0.0760, 0.0646, 0.0552, 0.0467, 0.0405, 0.0350, 0.0305, 0.0266, 0.0236, 0.0160, 0.0118, 0.0095],
        [0.0860, 0.0660, 0.0630, 0.0510, 0.0530, 0.0464, 0.0410, 0.0363, 0.0317, 0.0280, 0.0247, 0.0220, 0.0197, 0.0140, 0.0100, 0.0085]
    ],
    15: [
        [0.2610, 0.2510, 0.2330, 0.2100, 0.1530, 0.1210, 0.0965, 0.0786, 0.0719, 0.0630, 0.0508, 0.0440, 0.0375, 0.0259, 0.0211, 0.0191],
        [0.2250, 0.2200, 0.1910, 0.1660, 0.1330, 0.1060, 0.0875, 0.0730, 0.0632, 0.0590, 0.0478, 0.0420, 0.0370, 0.0252, 0.0206, 0.0189],
        [0.1580, 0.1170, 0.1360, 0.1100, 0.0990, 0.0840, 0.0723, 0.0620, 0.0548, 0.0485, 0.0430, 0.0390, 0.0354, 0.0234, 0.0195, 0.0181],
        [0.1050, 0.1030, 0.0940, 0.0870, 0.0755, 0.0660, 0.0590, 0.0520, 0.0463, 0.0418, 0.0383, 0.0345, 0.0313, 0.0215, 0.0185, 0.0166],
        [0.0740, 0.0730, 0.0687, 0.0640, 0.0570, 0.0514, 0.0463, 0.0421, 0.0386, 0.0350, 0.0321, 0.0295, 0.0274, 0.0202, 0.0172, 0.0155]
    ],
    16: [
        [0.3000, 0.2900, 0.2490, 0.2290, 0.1840, 0.1550, 0.1290, 0.1060, 0.0900, 0.0768, 0.0674, 0.0593, 0.0530, 0.0403, 0.0298, 0.0231],
        [0.2520, 0.2430, 0.2150, 0.2000, 0.1660, 0.1380, 0.1140, 0.0960, 0.0820, 0.0717, 0.0627, 0.0555, 0.0507, 0.0368, 0.0287, 0.0227],
        [0.1940, 0.1890, 0.1730, 0.1550, 0.1300, 0.1100, 0.0920, 0.0790, 0.0692, 0.0617, 0.0552, 0.0493, 0.0445, 0.0324, 0.0270, 0.0218],
        [0.1480, 0.1430, 0.1300, 0.1190, 0.0990, 0.0870, 0.0740, 0.0660, 0.0590, 0.0530, 0.0469, 0.0428, 0.0392, 0.0290, 0.0242, 0.0205],
        [0.0940, 0.0920, 0.0890, 0.0810, 0.0710, 0.0630, 0.0570, 0.0520, 0.0473, 0.0433, 0.0397, 0.0357, 0.0330, 0.0265, 0.0228, 0.0193]
    ],
    17: [
        [0.2000, 0.1900, 0.1660, 0.1460, 0.1170, 0.0960, 0.0800, 0.0680, 0.0575, 0.0490, 0.0420, 0.0360, 0.0305, 0.0160, 0.0140, 0.0125],
        [0.1800, 0.1720, 0.1540, 0.1370, 0.1120, 0.0920, 0.0770, 0.0650, 0.0560, 0.0470, 0.0400, 0.0345, 0.0295, 0.0155, 0.0135, 0.0122],
        [0.1500, 0.1470, 0.1340, 0.1210, 0.1000, 0.0840, 0.0700, 0.0539, 0.0500, 0.0430, 0.0370, 0.0315, 0.0270, 0.0150, 0.0130, 0.0118],
        [0.1300, 0.1280, 0.1270, 0.1050, 0.0860, 0.0780, 0.0620, 0.0530, 0.0455, 0.0387, 0.0335, 0.0295, 0.0250, 0.0145, 0.0125, 0.0115],
        [0.0850, 0.0840, 0.0780, 0.0720, 0.0600, 0.0510, 0.0440, 0.0375, 0.0325, 0.0290, 0.0262, 0.0235, 0.0210, 0.0140, 0.0120, 0.0110]
    ],
    18: [
        [0.3020, 0.2760, 0.2360, 0.2210, 0.1670, 0.1390, 0.1140, 0.0963, 0.0819, 0.0707, 0.0615, 0.0543, 0.0478, 0.0329, 0.0254, 0.0223],
        [0.2360, 0.2290, 0.2020, 0.1810, 0.1500, 0.1250, 0.1050, 0.0878, 0.0765, 0.0660, 0.0580, 0.0513, 0.0433, 0.0312, 0.0246, 0.0213],
        [0.1840, 0.1790, 0.1380, 0.1420, 0.1180, 0.1000, 0.0857, 0.0746, 0.0647, 0.0567, 0.0505, 0.0451, 0.0409, 0.0285, 0.0228, 0.0200],
        [0.1290, 0.1260, 0.1140, 0.0980, 0.0880, 0.0770, 0.0670, 0.0596, 0.0534, 0.0477, 0.0431, 0.0396, 0.0357, 0.0264, 0.0213, 0.0182],
        [0.0920, 0.0890, 0.0820, 0.0750, 0.0652, 0.0580, 0.0513, 0.0467, 0.0428, 0.0390, 0.0357, 0.0326, 0.0303, 0.0232, 0.0190, 0.0172]
    ]
}

AP_SOURCE_DISCREPANCIES = {(1, 3, 4): (0.077, 0.017), (3, 0, 3): (0.122, 0.112), (4, 1, 6): (0.0816, 0.816), (10, 4, 6): (0.0345, 0.0435), (10, 4, 11): (0.0224, 0.0242), (11, 0, 6): (0.0461, 0.461), (13, 0, 1): (0.149, 0.0149), (13, 2, 10): (0.0351, 0.0315), (16, 0, 7): (0.106, 0.0106), (18, 0, 4): (0.167, 0.0167), (18, 0, 5): (0.139, 0.0139), (18, 0, 6): (0.114, 0.0114), (18, 0, 7): (0.0963, 0.963), (18, 1, 4): (0.15, 0.015), (18, 1, 5): (0.125, 0.0125), (18, 1, 6): (0.105, 0.0105), (18, 1, 7): (0.0878, 0.0978), (18, 2, 4): (0.118, 0.0118), (18, 2, 5): (0.1, 0.01), (18, 2, 11): (0.0451, 0.0541), (18, 4, 12): (0.0303, 0.303)}

def finite_number(value, name):
    value = float(str(value).strip().replace(',', '.'))
    if not np.isfinite(value): raise ValueError(f"{name} phải là số hữu hạn, không dùng NaN/vô cực.")
    return value


def validate_region(value):
    region = finite_number(value, 'Vùng mưa')
    if region != int(region) or not 1 <= region <= 18:
        raise ValueError('Vùng mưa phải là số nguyên từ 1 đến 18.')
    return int(region)


def peak_flow_parameters(F, L, Sli, Jl, Js, ms, ml, delta, H, region, soil, k_bs=1.8, confirmed_phi=None, confirmed_ap=None, confirmed_ts=None):
    values = dict(F=F,L=L,Sli=Sli,Jl=Jl,Js=Js,ms=ms,ml=ml,delta=delta,H=H,k_bs=k_bs)
    values = {key:finite_number(value,key) for key,value in values.items()}
    F,L,Sli,Jl,Js,ms,ml,delta,H,k_bs=(values[k] for k in values)
    if not 0<F<100:raise ValueError('Phạm vi phần mềm: 0 < F < 100 km².')
    if Sli<0 or any(v<=0 for v in (L,Jl,Js,ms,ml,H)):raise ValueError('L, Jl, Js, ms, ml, Hp phải dương; tổng dòng nhánh không âm.')
    if not 0<delta<=1:raise ValueError('Hệ số giảm hồ đầm δ phải trong khoảng 0 < δ ≤ 1.')
    if k_bs not in (.9,1.8):raise ValueError('Mẫu số b sườn dốc phải là 0,90 hoặc 1,80.')
    region=validate_region(region)
    phi=lookup_phi(soil,H,F,confirmed_phi)
    runoff=phi*H
    bs=1000*F/(k_bs*(L+Sli))
    Fs=bs**.6/(ms*Js**.3*runoff**.4)
    ts=lookup_ts(region,Fs,confirmed_ts)
    Fl=1000*L/(ml*Jl**(1/3)*(F*runoff)**.25)
    Ap=lookup_Ap(region,ts,Fl,confirmed_ap)
    Q=Ap*runoff*F*delta
    result=dict(phi=phi,bs=bs,Fs=Fs,ts=ts,Fl=Fl,Ap=Ap,Qmax=Q)
    if not all(np.isfinite(v) and v>0 for v in result.values()):raise ValueError('Kết quả không hữu hạn/dương; kiểm tra dữ liệu và bảng tra.')
    return result


def table_consistency_notes(region,soil):
    notes=[]
    mat=np.asarray(AP_TABLES[region],dtype=float)
    if np.any(np.diff(mat,axis=1)>0) or np.any(np.diff(mat,axis=0)>0):
        notes.append(f'Bảng Ap vùng {region} có điểm không đơn điệu; cần kiểm tra bản chuẩn khi dùng cho thiết kế.')
    return notes


def phi_reference_values(soil_type,H,F):
    soil_type=soil_type.strip().upper()
    H=finite_number(H,'Hp');F=finite_number(F,'F')
    if soil_type not in SOIL_TABLES or H<=0 or not 0<F<100:raise ValueError('Tra φ cần cấp đất II–VI, Hp > 0 và 0 < F < 100.')
    row=min(int(np.searchsorted(SOIL_H_BOUNDS[soil_type],H,side='left')),len(SOIL_TABLES[soil_type])-1)
    start,end=next((start,end) for lo,hi,start,end in PHI_AREA_GROUPS if lo<=F<hi)
    return SOIL_TABLES[soil_type][row][start:end]

def lookup_phi(soil_type,H,F,confirmed_phi=None):
    reference=phi_reference_values(soil_type,H,F)
    if confirmed_phi is not None:
        phi=finite_number(confirmed_phi,'φ đã xác minh')
        if not 0<phi<=1:raise ValueError('Hệ số dòng chảy phải trong khoảng 0 < φ ≤ 1.')
        if not min(reference)-1e-9<=phi<=max(reference)+1e-9:
            raise ValueError(f'φ={phi:.3f} ngoài khoảng {min(reference):.3f}–{max(reference):.3f} của cấp đất/mưa/diện tích trong bảng A.1 được cung cấp.')
        return phi
    if max(reference)-min(reference)<1e-12:return float(reference[0])
    raise ValueError(f'Bảng A.1 Word thiếu mốc F cho từng cột trong nhóm này. Hãy nhập φ đã xác minh (khoảng tham khảo {min(reference):.3f}–{max(reference):.3f}), không nội suy theo trục F tự đặt.')


def lookup_ts(vung_id, Fs, confirmed_ts=None):
    vung_id = validate_region(vung_id)
    if not 1 <= vung_id <= 18:
        raise ValueError("Vùng mưa phải từ 1 đến 18.")
    Fs=finite_number(Fs,"Φs")
    if Fs <= 0:
        raise ValueError("Phi_s phải lớn hơn 0.")
    if confirmed_ts is not None:
        value=finite_number(confirmed_ts,"ts đã xác minh")
        if value<=0:raise ValueError("ts đã xác minh phải dương.")
        return value
    if vung_id==18 and Fs>FS_AXIS[-2]:
        raise ValueError("Ô A.2 vùng 18, Φs=40 bị bỏ trống trong bản Word. Hãy nhập ts đã xác minh khi Φs > 35.")
    col_idx = vung_id - 1
    col_vals = [TS_GRID[i][col_idx] for i in range(len(FS_AXIS))]
    return float(np.interp(Fs, FS_AXIS, col_vals))


def lookup_Ap(vung_id, ts_val, fl_val, confirmed_ap=None):
    vung_id = validate_region(vung_id)
    if not 1 <= vung_id <= 18:
        raise ValueError("Vùng mưa phải từ 1 đến 18.")
    if vung_id not in AP_TABLES:
        raise ValueError(f"Chưa có bảng A.3 cho vùng mưa {vung_id}.")
    ts_val=finite_number(ts_val,"ts");fl_val=finite_number(fl_val,"Φl")
    if ts_val <= 0 or fl_val < 0:
        raise ValueError("ts phải > 0 và Phi_l phải >= 0.")
    if confirmed_ap is not None:
        value=finite_number(confirmed_ap,'Ap đã xác minh')
        if value<=0:raise ValueError('Ap đã xác minh phải dương.')
        return value
    def active_indices(axis,value):
        clamped=float(np.clip(value,axis[0],axis[-1]))
        exact=np.flatnonzero(np.isclose(axis,clamped,rtol=0,atol=1e-12))
        if len(exact):return [int(exact[0])]
        upper=int(np.searchsorted(axis,clamped));return [upper-1,upper]
    disputed=[(row,col) for row in active_indices(TS_LEVELS,ts_val) for col in active_indices(FL_COLS,fl_val) if (vung_id,row,col) in AP_SOURCE_DISCREPANCIES]
    if disputed:
        detail='; '.join(f'ts={TS_LEVELS[r]}, Φl={FL_COLS[c]}: mã {AP_SOURCE_DISCREPANCIES[vung_id,r,c][0]}, Word {AP_SOURCE_DISCREPANCIES[vung_id,r,c][1]}' for r,c in disputed)
        raise ValueError('Bảng Ap có ô chưa thống nhất giữa mã và tài liệu: '+detail+'. Cần nhập Ap đã đối chiếu bản chuẩn trước khi tính.')
    table = AP_TABLES[vung_id]
    mat = np.array(table, dtype=float)
    ap_at_ts = [np.interp(fl_val, FL_COLS, mat[i, :]) for i in range(len(TS_LEVELS))]
    return float(np.interp(ts_val, TS_LEVELS, ap_at_ts))


# =============================================================================
# GIAO DIỆN ĐỒ HỌA HOÀN CHỈNH
# =============================================================================
class FullTCVNApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"TÍNH TOÁN LƯU LƯỢNG LŨ THIẾT KẾ Qp - TCVN 9845:2013")
        self.root.geometry("1040x980")
        self.root.minsize(960,700)
        self.root.resizable(True, True)

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.calc_results = []
        self._result_signature = None
        self.mean_x = 100.0
        self.rain_series = None
        self.phi_win = None
        self.ts_win = None
        self.ap_win = None
        self.freq_table_win = None
        self.freq_plot_win = None

        self.setup_ui()

    def on_close(self):
        try:
            for w in [self.phi_win, self.ts_win, self.ap_win, self.freq_table_win, self.freq_plot_win]:
                if w and w.winfo_exists(): w.destroy()
            self.root.quit()
            self.root.destroy()
        except Exception:
            pass
        finally:
            os._exit(0)

    def setup_ui(self):
        menubar = tk.Menu(self.root)
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="💾 Lưu dự án (Save Project)", command=self.save_project)
        file_menu.add_command(label="📂 Mở dự án cũ (Open Project)", command=self.open_project)
        file_menu.add_separator()
        file_menu.add_command(label="Thoát", command=self.on_close)
        menubar.add_cascade(label="Tệp dự án (File)", menu=file_menu)
        self.root.config(menu=menubar)

        hdr = tk.Frame(self.root, pady=4)
        hdr.pack(fill=tk.X)
        tk.Label(hdr, text=f"TÍNH TOÁN LƯU LƯỢNG LŨ THIẾT KẾ THEO TCVN 9845:2013",
                 font=("Times New Roman", 14, "bold"), fg="#154360").pack()

        r_proj = tk.Frame(self.root, padx=15, pady=2)
        r_proj.pack(fill=tk.X)
        tk.Label(r_proj, text="Tên công trình / Vị trí tính:", font=("Arial", 9, "bold"), fg="#b71c1c").pack(side=tk.LEFT)
        self.txt_project = tk.Entry(r_proj, font=("Arial", 10, "bold"), fg="#b71c1c", width=34)
        self.txt_project.insert(0, "CẦU SUỐI SẢI")
        self.txt_project.pack(side=tk.LEFT, padx=10)

        btn_tables_fr = tk.Frame(self.root, padx=15, pady=2)
        btn_tables_fr.pack(fill=tk.X)
        tk.Label(btn_tables_fr, text="Tra cứu bảng chuẩn:", font=("Arial", 8, "italic")).pack(side=tk.LEFT)
        tk.Button(btn_tables_fr, text="📑 Bảng tra hệ số φ", font=("Arial", 8, "bold"), bg="#e8f5e9", fg="#2e7d32", command=self.show_phi_table_window).pack(side=tk.LEFT, padx=6)
        tk.Button(btn_tables_fr, text="📑 Bảng II-2: Tra ts", font=("Arial", 8, "bold"), bg="#e1f5fe", fg="#0277bd", command=self.show_ts_table_window).pack(side=tk.LEFT, padx=6)
        tk.Button(btn_tables_fr, text="📑 Bảng tra mô-đun Ap", font=("Arial", 8, "bold"), bg="#f3e5f5", fg="#6a1b9a", command=self.show_ap_table_window).pack(side=tk.LEFT, padx=6)

        # 2. KHUNG THÔNG SỐ ĐẶC TRƯNG LƯU VỰC
        p_fr = tk.LabelFrame(self.root, text=" 1. Thông số đặc trưng hình thái lưu vực ", font=("Arial", 9, "bold"), padx=10, pady=4)
        p_fr.pack(fill=tk.X, padx=15, pady=2)

        r1 = tk.Frame(p_fr); r1.pack(fill=tk.X, pady=2)
        tk.Label(r1, text="Diện tích lưu vực F (km²):", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_F = tk.Entry(r1, width=10); self.txt_F.insert(0, "65.79"); self.txt_F.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r1, text="Chiều dài sông chính L (km):", width=24, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_L = tk.Entry(r1, width=10); self.txt_L.insert(0, "11.51"); self.txt_L.pack(side=tk.LEFT)

        r2 = tk.Frame(p_fr); r2.pack(fill=tk.X, pady=2)
        tk.Label(r2, text="Tổng dài dòng nhánh ΣLi (km):", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_Sli = tk.Entry(r2, width=10); self.txt_Sli.insert(0, "21.87"); self.txt_Sli.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r2, text="Độ dốc lòng sông Jl (‰):", width=24, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_Jl = tk.Entry(r2, width=10); self.txt_Jl.insert(0, "2.45"); self.txt_Jl.pack(side=tk.LEFT)

        r2b = tk.Frame(p_fr); r2b.pack(fill=tk.X, pady=2)
        tk.Label(r2b, text="Dạng sườn dốc lưu vực:", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.cbo_slope = ttk.Combobox(
            r2b, width=28, state="readonly",
            values=["Hai sườn dốc - mẫu số 1,80", "Một sườn dốc - mẫu số 0,90"]
        )
        self.cbo_slope.set("Hai sườn dốc - mẫu số 1,80")
        self.cbo_slope.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r2b, text="Hệ số bs tự động theo dạng lưu vực", fg="#555", font=("Arial", 8, "italic")).pack(side=tk.LEFT)

        r3 = tk.Frame(p_fr); r3.pack(fill=tk.X, pady=2)
        tk.Label(r3, text="Độ dốc sườn dốc Js (‰):", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_Js = tk.Entry(r3, width=10); self.txt_Js.insert(0, "63.12"); self.txt_Js.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r3, text="Vùng mưa (1 - 18):", width=24, anchor=tk.W).pack(side=tk.LEFT)
        self.cbo_vung = ttk.Combobox(r3, width=8, state="readonly", values=[str(i) for i in range(1, 19)])
        self.cbo_vung.set("18"); self.cbo_vung.pack(side=tk.LEFT)

        r4 = tk.Frame(p_fr); r4.pack(fill=tk.X, pady=2)
        tk.Label(r4, text="Cấp đất lưu vực:", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.cbo_soil = ttk.Combobox(r4, width=8, state="readonly", values=["II", "III", "IV", "V", "VI"])
        self.cbo_soil.set("IV"); self.cbo_soil.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r4, text="Hệ số chiết giảm hồ đầm δ:", width=24, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_d = tk.Entry(r4, width=10); self.txt_d.insert(0, "1.0"); self.txt_d.pack(side=tk.LEFT)

        r5 = tk.Frame(p_fr); r5.pack(fill=tk.X, pady=2)
        tk.Label(r5, text="Độ nhám sườn dốc ms:", width=22, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_ms = tk.Entry(r5, width=10); self.txt_ms.insert(0, "0.15"); self.txt_ms.pack(side=tk.LEFT, padx=(0, 20))
        tk.Label(r5, text="Độ nhám lòng chính ml:", width=24, anchor=tk.W).pack(side=tk.LEFT)
        self.txt_ml = tk.Entry(r5, width=10); self.txt_ml.insert(0, "7.0"); self.txt_ml.pack(side=tk.LEFT)

        # 3. PHÂN TÍCH THỐNG KÊ MƯA & BĐKH
        stat_fr = tk.LabelFrame(self.root, text=" 2. Thống kê mưa Pearson III & Hệ số Biến đổi khí hậu (BĐKH) ", font=("Arial", 9, "bold"), fg="#b71c1c", padx=10, pady=5)
        stat_fr.pack(fill=tk.X, padx=15, pady=2)

        r_st1 = tk.Frame(stat_fr); r_st1.pack(fill=tk.X, pady=2)
        tk.Button(r_st1, text="📁 Nhập tệp trạm mưa (.txt)", font=("Arial", 9, "bold"), bg="#e8f5e9", fg="#2e7d32", command=self.load_rain_txt_file).pack(side=tk.LEFT, padx=(0, 10))

        tk.Label(r_st1, text="Cv:").pack(side=tk.LEFT, padx=2)
        self.txt_cv = tk.Entry(r_st1, width=6, justify=tk.CENTER); self.txt_cv.insert(0, "0.35"); self.txt_cv.pack(side=tk.LEFT, padx=2)

        tk.Label(r_st1, text="Cs:").pack(side=tk.LEFT, padx=2)
        self.txt_cs = tk.Entry(r_st1, width=6, justify=tk.CENTER); self.txt_cs.insert(0, "1.05"); self.txt_cs.pack(side=tk.LEFT, padx=2)

        tk.Button(r_st1, text="⚡ Cập nhật Hp", font=("Arial", 8, "bold"), bg="#fff3e0", fg="#e65100", command=self.recalculate_hp_from_cv_cs).pack(side=tk.LEFT, padx=6)

        # NÚT XEM BẢNG TẦN SUẤT & NÚT XEM ĐỒ THỊ TẦN SUẤT
        tk.Button(r_st1, text="📊 Xem Đồ thị tần suất", font=("Arial", 8, "bold"), bg="#ede7f6", fg="#4527a0", command=self.show_frequency_plot_window).pack(side=tk.RIGHT, padx=2)
        tk.Button(r_st1, text="📈 Xem Bảng tần suất", font=("Arial", 8, "bold"), bg="#e1f5fe", fg="#0277bd", command=self.show_frequency_table_window).pack(side=tk.RIGHT, padx=2)

        r_rain = tk.Frame(stat_fr); r_rain.pack(fill=tk.X, pady=4)
        self.rain_entries = {}
        self.cc_entries = {}
        self.ts_entries = {}
        self.phi_hint_labels = {}
        self.phi_entries = {}
        self.ap_entries = {}

        rain_defaults = [
            ("P=0.5%", "235.4"),
            ("P=1%", "210.8"),
            ("P=2%", "178.6"),
            ("P=4%", "167.0"),
            ("P=5%", "160.0"),
            ("P=10%", "148.11"),
            ("P=20%", "124.5")
        ]
        for p_label, val in rain_defaults:
            box = tk.Frame(r_rain, bd=1, relief=tk.GROOVE, padx=2, pady=2)
            box.pack(side=tk.LEFT, expand=True, padx=2)
            tk.Label(box, text=p_label, font=("Arial", 8, "bold"), fg="#154360").pack()

            tk.Label(box, text="Hp gốc:", font=("Arial", 7)).pack()
            ent = tk.Entry(box, width=7, justify=tk.CENTER)
            ent.insert(0, val)
            ent.pack()
            self.rain_entries[p_label] = ent

            tk.Label(box, text="Hệ số BĐKH:", font=("Arial", 7), fg="#b71c1c").pack()
            ent_cc = tk.Entry(box, width=7, justify=tk.CENTER)
            ent_cc.insert(0, "1.0")
            ent_cc.pack()
            self.cc_entries[p_label] = ent_cc
            tk.Label(box,text='φ đã xác minh:',font=('Arial',7)).pack()
            ent_phi=tk.Entry(box,width=7,justify=tk.CENTER);ent_phi.pack();self.phi_entries[p_label]=ent_phi
            hint=tk.Label(box,text='',font=('Arial',7),fg='#155e75',wraplength=125,justify=tk.CENTER)
            hint.pack(pady=(2,3));self.phi_hint_labels[p_label]=hint
            tk.Label(box,text='ts (nếu cần, phút):',font=('Arial',7)).pack()
            ent_ts=tk.Entry(box,width=7,justify=tk.CENTER);ent_ts.pack();self.ts_entries[p_label]=ent_ts
            tk.Label(box,text='Ap (nếu cần):',font=('Arial',7)).pack()
            ent_ap=tk.Entry(box,width=7,justify=tk.CENTER);ent_ap.pack();self.ap_entries[p_label]=ent_ap

        tk.Label(stat_fr,text='Để trống φ chỉ khi các cột trong nhóm A.1 cho cùng giá trị. Ap tự tra nếu không chạm ô khác biệt; ô chưa thống nhất cần nhập giá trị đã xác minh.',wraplength=900,fg='#8b4513',font=('Arial',8)).pack()
        # 4. KẾT QUẢ TÍNH TOÁN LŨ Qp
        res_fr = tk.LabelFrame(self.root, text=" 3. Kết quả tính toán lưu lượng lũ thiết kế Qp [m³/s] ", font=("Arial", 9, "bold"), padx=10, pady=5)
        res_fr.pack(fill=tk.BOTH, expand=True, padx=15, pady=2)

        cols = ("stt", "p", "hp_goc", "cc", "hp_bdkh", "phi", "bs", "fs", "ts", "fl", "ap", "qp")
        self.tree = ttk.Treeview(res_fr, columns=cols, show="headings", height=7)
        self.tree.heading("stt", text="STT"); self.tree.heading("p", text="Tần suất P")
        self.tree.heading("hp_goc", text="Hp gốc"); self.tree.heading("cc", text="Hệ số BĐKH")
        self.tree.heading("hp_bdkh", text="Hp tính toán"); self.tree.heading("phi", text="Hệ số phi")
        self.tree.heading("bs", text="bs (m)"); self.tree.heading("fs", text="Φs")
        self.tree.heading("ts", text="ts (phút)"); self.tree.heading("fl", text="Φl")
        self.tree.heading("ap", text="Mô-đun Ap"); self.tree.heading("qp", text="Qmax (m³/s)")

        self.tree.column("stt", width=35, anchor=tk.CENTER); self.tree.column("p", width=65, anchor=tk.CENTER)
        self.tree.column("hp_goc", width=65, anchor=tk.E); self.tree.column("cc", width=75, anchor=tk.CENTER)
        self.tree.column("hp_bdkh", width=80, anchor=tk.E); self.tree.column("phi", width=65, anchor=tk.CENTER)
        self.tree.column("bs", width=65, anchor=tk.CENTER); self.tree.column("fs", width=55, anchor=tk.CENTER)
        self.tree.column("ts", width=65, anchor=tk.CENTER); self.tree.column("fl", width=60, anchor=tk.CENTER)
        self.tree.column("ap", width=75, anchor=tk.CENTER); self.tree.column("qp", width=100, anchor=tk.E)
        self.tree.pack(fill=tk.BOTH, expand=True, pady=2)

        # 5. CÁC NÚT THAO TÁC
        btn_fr = tk.Frame(self.root, pady=8)
        btn_fr.pack(fill=tk.X, padx=15)

        tk.Button(btn_fr, text="🚀 TÍNH TOÁN LẠI TẤT CẢ", font=("Arial", 10, "bold"), bg="#154360", fg="white", padx=15, pady=5, command=self.calculate_all).pack(side=tk.LEFT)
        tk.Button(btn_fr, text="📋 Sao chép ra Excel", font=("Arial", 9, "bold"), bg="#27ae60", fg="white", padx=10, pady=5, command=self.copy_table).pack(side=tk.LEFT, padx=10)
        tk.Button(btn_fr, text="💾 Xuất file Excel (.xlsx)", font=("Arial", 9, "bold"), bg="#2980b9", fg="white", padx=10, pady=5, command=self.export_excel).pack(side=tk.RIGHT)

        # Cập nhật gợi ý cả khi nhập tay và khi mở dự án/cập nhật Hp.
        self._phi_hint_signature=None
        self.refresh_phi_hints()
        self.root.after(250,self.poll_phi_hints)
        # Chờ người dùng nhập/xác minh hệ số rồi bấm tính.

    def refresh_phi_hints(self):
        """Khoảng tham khảo A.1; không tự chọn hệ số thay người dùng."""
        for label,hint in self.phi_hint_labels.items():
            try:
                soil=self.cbo_soil.get().strip().upper()
                F=finite_number(self.txt_F.get(),'Diện tích')
                H=finite_number(self.rain_entries[label].get(),'Hp')*finite_number(self.cc_entries[label].get(),'Hệ số biến đổi khí hậu')
                if finite_number(self.cc_entries[label].get(),'Hệ số biến đổi khí hậu')<=0:raise ValueError('Hệ số phải dương')
                values=phi_reference_values(soil,H,F)
                row=min(int(np.searchsorted(SOIL_H_BOUNDS[soil],H,side='left')),len(SOIL_TABLES[soil])-1)
                lo,hi,_,_=next(group for group in PHI_AREA_GROUPS if group[0]<=F<group[1])
                context=f'Đất {soil}; mưa {SOIL_H_LABELS[soil][row]} mm\nF: {lo:g}–{hi:g} km²'
                if max(values)-min(values)<1e-12:
                    text=f'φ = {values[0]:.3f} (tự tra)\n'+context
                else:
                    text=f'Gợi ý φ: {min(values):.3f}–{max(values):.3f}\n'+context+'\nBạn chọn và nhập φ'
                hint.configure(text=text,fg='#155e75')
            except (ValueError,KeyError,StopIteration):
                hint.configure(text='Nhập F, cấp đất, Hp\nvà hệ số khí hậu hợp lệ\nđể xem khoảng φ',fg='#8b4513')

    def poll_phi_hints(self):
        signature=(self.txt_F.get(),self.cbo_soil.get())+tuple(entry.get() for entry in self.rain_entries.values())+tuple(entry.get() for entry in self.cc_entries.values())
        if signature!=self._phi_hint_signature:
            self._phi_hint_signature=signature
            self.refresh_phi_hints()
        self.root.after(250,self.poll_phi_hints)

    # --- TÍNH NĂNG LƯU VÀ MỞ DỰ ÁN (SAVE / OPEN) ---
    def save_project(self):
        f_path = filedialog.asksaveasfilename(
            defaultextension=".tcvn",
            filetypes=[("TCVN Hydrology Project files", "*.tcvn"), ("JSON files", "*.json"), ("All files", "*.*")]
        )
        if not f_path: return
        try:
            data = {
                "project_name": self.txt_project.get(),
                "F": self.txt_F.get(),
                "L": self.txt_L.get(),
                "Sli": self.txt_Sli.get(),
                "Jl": self.txt_Jl.get(),
                "slope_type": self.cbo_slope.get(),
                "Js": self.txt_Js.get(),
                "vung": self.cbo_vung.get(),
                "soil": self.cbo_soil.get(),
                "d": self.txt_d.get(),
                "ms": self.txt_ms.get(),
                "ml": self.txt_ml.get(),
                "cv": self.txt_cv.get(),
                "cs": self.txt_cs.get(),
                "mean_x": self.mean_x,
                "rain_series": self.rain_series.tolist() if self.rain_series is not None else None,
                "rain_data": {p: self.rain_entries[p].get() for p in self.rain_entries},
                "ts_data": {p:self.ts_entries[p].get() for p in self.ts_entries},
                "phi_data": {p:self.phi_entries[p].get() for p in self.phi_entries},
                "ap_data": {p:self.ap_entries[p].get() for p in self.ap_entries},
                "cc_data": {p: self.cc_entries[p].get() for p in self.cc_entries}
            }
            with open(f_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=4)
            messagebox.showinfo("Thành công", f"Đã lưu dự án thành công tại:\n{f_path}")
        except Exception as e:
            messagebox.showerror("Lỗi lưu file", f"Không thể lưu dự án: {e}")

    def open_project(self):
        f_path = filedialog.askopenfilename(
            filetypes=[("TCVN Hydrology Project files", "*.tcvn *.json"), ("All files", "*.*")]
        )
        if not f_path: return
        try:
            with open(f_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            self.txt_project.delete(0, tk.END); self.txt_project.insert(0, data.get("project_name", "CẦU"))
            self.txt_F.delete(0, tk.END); self.txt_F.insert(0, data.get("F", "65.79"))
            self.txt_L.delete(0, tk.END); self.txt_L.insert(0, data.get("L", "11.51"))
            self.txt_Sli.delete(0, tk.END); self.txt_Sli.insert(0, data.get("Sli", "21.87"))
            self.txt_Jl.delete(0, tk.END); self.txt_Jl.insert(0, data.get("Jl", "2.45"))
            self.cbo_slope.set(data.get("slope_type", "Hai sườn dốc - mẫu số 1,80"))
            self.txt_Js.delete(0, tk.END); self.txt_Js.insert(0, data.get("Js", "63.12"))
            self.cbo_vung.set(data.get("vung", "18"))
            self.cbo_soil.set(data.get("soil", "IV"))
            self.txt_d.delete(0, tk.END); self.txt_d.insert(0, data.get("d", "1.0"))
            self.txt_ms.delete(0, tk.END); self.txt_ms.insert(0, data.get("ms", "0.15"))
            self.txt_ml.delete(0, tk.END); self.txt_ml.insert(0, data.get("ml", "7.0"))
            self.txt_cv.delete(0, tk.END); self.txt_cv.insert(0, data.get("cv", "0.35"))
            self.txt_cs.delete(0, tk.END); self.txt_cs.insert(0, data.get("cs", "1.05"))
            self.mean_x = finite_number(data.get("mean_x",100.0),"Trung bình mưa")
            stored_series=data.get('rain_series')
            self.rain_series=None
            if stored_series is not None:
                series=np.asarray(stored_series,dtype=float)
                if series.ndim!=1 or len(series)<3 or not np.all(np.isfinite(series)) or np.any(series<0) or np.mean(series)<=0:
                    raise ValueError('Chuỗi mưa lưu trong dự án không hợp lệ.')
                self.rain_series=series;self.mean_x=float(np.mean(series))

            rain_data = data.get("rain_data", {})
            for p, val in rain_data.items():
                if p in self.rain_entries:
                    self.rain_entries[p].delete(0, tk.END)
                    self.rain_entries[p].insert(0, val)

            cc_data = data.get("cc_data", {})
            for p, val in cc_data.items():
                if p in self.cc_entries:
                    self.cc_entries[p].delete(0, tk.END)
                    self.cc_entries[p].insert(0, val)

            for key,entries in (('ts_data',self.ts_entries),('phi_data',self.phi_entries),('ap_data',self.ap_entries)):
                for label,entry in entries.items():
                    entry.delete(0,tk.END);entry.insert(0,data.get(key,{}).get(label,''))
            self.calculate_all()
            messagebox.showinfo("Thành công", f"Đã mở lại dự án từ tệp:\n{f_path}")
        except Exception as e:
            messagebox.showerror("Lỗi mở file", f"Không thể đọc tệp dự án: {e}")

    def input_signature(self):
        fields=(self.txt_project,self.txt_F,self.txt_L,self.txt_Sli,self.txt_Jl,self.txt_Js,self.txt_ms,self.txt_ml,self.txt_d,self.cbo_vung,self.cbo_soil,self.cbo_slope)
        return tuple(w.get() for w in fields)+tuple(w.get() for w in self.rain_entries.values())+tuple(w.get() for w in self.cc_entries.values())+tuple(w.get() for w in self.phi_entries.values())+tuple(w.get() for w in self.ap_entries.values())+tuple(w.get() for w in self.ts_entries.values())

    def require_current_results(self):
        if not self.calc_results:
            messagebox.showwarning('Chưa có kết quả','Hãy tính lại lưu lượng trước khi sao chép/xuất.');return False
        if self._result_signature != self.input_signature():
            self.calc_results=[];self.tree.delete(*self.tree.get_children())
            messagebox.showwarning('Đầu vào đã thay đổi','Kết quả cũ đã xóa. Hãy bấm tính lại trước khi sao chép/xuất.');return False
        return True

    def get_kp(self, p_percent, cv, cs):
        p_percent=finite_number(p_percent,'P');cv=finite_number(cv,'Cv');cs=finite_number(cs,'Cs')
        if not 0<p_percent<100 or cv<0:raise ValueError('0 < P < 100%; Cv không âm.')
        if cv==0:return 1.0
        factor=float(1+cv*pearson3.ppf(1-p_percent/100,skew=cs))
        if not np.isfinite(factor) or factor<=0:raise ValueError('Phân phối Pearson III cho hệ số mưa không dương/hữu hạn; cần kiểm tra tham số, không tự ép thành 0.')
        return factor

    def recalculate_hp_from_cv_cs(self):
        try:
            cv = finite_number(self.txt_cv.get(),"Cv")
            cs = finite_number(self.txt_cs.get(),"Cs")
            if finite_number(self.mean_x,"Trung bình mưa") <= 0:raise ValueError("Trung bình mưa phải dương.")
            p_vals = [0.5, 1.0, 2.0, 4.0, 5.0, 10.0, 20.0]

            new_hp = []
            for p in p_vals:
                kp = self.get_kp(p, cv, cs)
                new_hp.append(self.mean_x * kp)

            keys = list(self.rain_entries.keys())
            for k_label, hp_val in zip(keys, new_hp):
                if k_label in self.rain_entries:
                    self.rain_entries[k_label].delete(0, tk.END)
                    self.rain_entries[k_label].insert(0, f"{hp_val:.2f}")

            messagebox.showinfo("Thành công", f"Đã cập nhật Hp gốc chuẩn Pearson III với X_tb = {self.mean_x:.2f}, Cv = {cv}, Cs = {cs}!")
            self.calculate_all()
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không tính được Hp: {e}")

    def load_rain_txt_file(self):
        f_path = filedialog.askopenfilename(filetypes=[("Text files", "*.txt"), ("All files", "*.*")])
        if not f_path: return
        try:
            with open(f_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = f.readlines()

            vals = []
            for line_number,line in enumerate(lines,1):
                line=line.strip()
                if not line or line.startswith('#'):continue
                if not vals and any(word in line.lower() for word in ('năm','year','rain','mưa')):continue
                parts=line.replace(';',' ').split()
                if len(parts)==1 and ',' in line:
                    pieces=line.split(',')
                    if len(pieces)==2 and pieces[0].isdigit() and len(pieces[0])==4:parts=pieces
                if len(parts) not in (1,2):raise ValueError(f'Dòng {line_number}: cần một cột mưa hoặc hai cột năm và mưa.')
                if len(parts)==2:
                    year=finite_number(parts[0],f'Năm dòng {line_number}')
                    if year!=int(year):raise ValueError(f'Dòng {line_number}: năm phải là số nguyên.')
                rain=finite_number(parts[-1],f'Mưa dòng {line_number}')
                if rain<0:raise ValueError(f'Dòng {line_number}: mưa không được âm.')
                vals.append(rain)

            if len(vals) < 3:
                messagebox.showwarning("Cảnh báo", "Chuỗi số liệu mưa phải có ít nhất 3 năm.")
                return

            arr_vals = np.array(vals)
            self.rain_series = arr_vals
            self.mean_x = float(np.mean(arr_vals))
            std_x = float(np.std(arr_vals, ddof=1))
            if self.mean_x<=0:raise ValueError("Trung bình chuỗi mưa phải dương.")
            cv_calc = std_x / self.mean_x
            cs_calc = float(skew(arr_vals, bias=False)) if std_x>0 else 0.0

            self.txt_cv.delete(0, tk.END); self.txt_cv.insert(0, f"{cv_calc:.3f}")
            self.txt_cs.delete(0, tk.END); self.txt_cs.insert(0, f"{cs_calc:.3f}")

            messagebox.showinfo("Thành công", f"Đã nạp {len(vals)} năm từ file trạm mưa.\nX_tb = {self.mean_x:.2f} mm | Cv = {cv_calc:.3f} | Cs = {cs_calc:.3f}")
            self.recalculate_hp_from_cv_cs()
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không đọc được file: {e}")

    def show_frequency_table_window(self):
        if self.freq_table_win and self.freq_table_win.winfo_exists(): self.freq_table_win.destroy()
        self.freq_table_win = tk.Toplevel(self.root)
        self.freq_table_win.title("BẢNG TẦN SUẤT LÝ LUẬN MƯA PEARSON III")
        self.freq_table_win.geometry("820x560")

        tk.Label(self.freq_table_win, text="BẢNG TẦN SUẤT LÝ LUẬN MƯA (PHÂN PHỐI PEARSON III)", font=("Times New Roman", 13, "bold"), fg="#0277bd").pack(pady=(10, 4))

        try:
            cv = finite_number(self.txt_cv.get(),"Cv")
            cs = finite_number(self.txt_cs.get(),"Cs")
            if finite_number(self.mean_x,"Trung bình mưa") <= 0:raise ValueError("Trung bình mưa phải dương.")
        except ValueError:
            cv, cs = 0.35, 1.05

        p_list = [0.01, 0.02, 0.05, 0.1, 0.2, 0.5, 1.0, 2.0, 3.0, 4.0, 5.0, 10.0, 20.0, 30.0, 50.0, 70.0, 80.0, 90.0, 95.0, 99.0]

        cols = ("p", "kp", "hp")
        tree = ttk.Treeview(self.freq_table_win, columns=cols, show="headings", height=14)
        tree.heading("p", text="Tần suất P (%)"); tree.heading("kp", text="Hệ số Kp"); tree.heading("hp", text="Lượng mưa Hp (mm)")
        tree.column("p", width=150, anchor=tk.CENTER); tree.column("kp", width=180, anchor=tk.CENTER); tree.column("hp", width=220, anchor=tk.E)

        table_records = []
        for p in p_list:
            kp = self.get_kp(p, cv, cs)
            hp = self.mean_x * kp
            tree.insert("", tk.END, values=(f"{p} %", f"{kp:.4f}", f"{hp:.2f}"))
            table_records.append({"P (%)": p, "Kp": kp, "Hp (mm)": hp})

        tree.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)

        def export_freq_excel():
            f_s = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
            if not f_s: return
            pd.DataFrame(table_records).to_excel(f_s, index=False)
            messagebox.showinfo("Thành công", f"Đã xuất bảng tần suất ra file:\n{f_s}")

        tk.Button(self.freq_table_win, text="💾 Xuất bảng tần suất ra Excel", font=("Arial", 9, "bold"), bg="#27ae60", fg="white", padx=10, pady=5, command=export_freq_excel).pack(pady=8)

    # --- TÍNH NĂNG XEM ĐỒ THỊ ĐƯỜNG TẦN SUẤT CHUẨN THỦY VĂN ---
    def show_frequency_plot_window(self):
        if self.freq_plot_win and self.freq_plot_win.winfo_exists():
            self.freq_plot_win.destroy()
        self.freq_plot_win = tk.Toplevel(self.root)
        self.freq_plot_win.title("ĐỒ THỊ ĐƯỜNG TẦN SUẤT MƯA PEARSON III (LƯỚI XÁC SUẤT)")
        self.freq_plot_win.geometry("980x680")

        ctrl_fr = tk.Frame(self.freq_plot_win, padx=12, pady=6)
        ctrl_fr.pack(fill=tk.X)

        tk.Label(ctrl_fr, text="ĐƯỜNG TẦN SUẤT LÝ LUẬN MƯA (PHÂN PHỐI PEARSON III)",
                 font=("Times New Roman", 13, "bold"), fg="#1a237e").pack(side=tk.LEFT)

        scale_var = tk.StringVar(value="Prob")

        plot_container = tk.Frame(self.freq_plot_win)
        plot_container.pack(fill=tk.BOTH, expand=True)

        def render_plot():
            for widget in plot_container.winfo_children():
                widget.destroy()

            try:
                cv = finite_number(self.txt_cv.get(),"cv")
                cs = finite_number(self.txt_cs.get(),"cs")
            except ValueError:
                cv, cs = 0.35, 1.05

            fig = Figure(figsize=(9.2, 5.6), dpi=100)
            ax = fig.add_subplot(111)

            mode = scale_var.get()

            # Các mức tần suất làm lưới chuẩn thủy văn
            prob_ticks = [0.05, 0.1, 0.2, 0.5, 1.0, 2.0, 5.0, 10.0, 20.0, 30.0, 50.0, 70.0, 80.0, 90.0, 95.0, 98.0, 99.0]

            def p_to_x(p_arr):
                p_clamped = np.clip(np.array(p_arr, dtype=float), 0.01, 99.99)
                if mode == "Prob":
                    # Biến đổi Gauss chuẩn: norm.ppf(P/100) để P nhỏ ở bên trái, P lớn ở bên phải
                    return norm.ppf(p_clamped / 100.0)
                return p_clamped

            # 1. Đường cong tần suất lý luận
            if mode == "Prob":
                p_curve = np.concatenate([
                    np.linspace(0.04, 1.0, 100),
                    np.linspace(1.0, 10.0, 100),
                    np.linspace(10.0, 99.5, 150)
                ])
            elif mode == "Log":
                p_curve = np.geomspace(0.05, 99.0, 300)
            else:
                p_curve = np.linspace(0.05, 99.0, 300)

            hp_curve = [self.mean_x * self.get_kp(p, cv, cs) for p in p_curve]
            x_curve = p_to_x(p_curve)
            ax.plot(x_curve, hp_curve, color="#1565c0", lw=2.2, zorder=3,
                    label=f"Lý luận Pearson III (Xtb={self.mean_x:.1f} mm, Cv={cv:.2f}, Cs={cs:.2f})")

            # 2. Điểm kinh nghiệm trạm mưa thực đo (nếu có nạp file)
            if self.rain_series is not None and len(self.rain_series) > 0:
                sorted_x = np.sort(self.rain_series)[::-1]
                n = len(sorted_x)
                p_emp = np.array([(m / (n + 1.0)) * 100.0 for m in range(1, n + 1)])
                x_emp = p_to_x(p_emp)
                ax.scatter(x_emp, sorted_x, color="#ff9800", edgecolors="#b71c1c", s=36, zorder=4,
                           label=f"Điểm kinh nghiệm trạm mưa (n={n} năm)")

            # 3. Các điểm tần suất thiết kế lấy từ giao diện
            p_design = []
            hp_design = []
            for p_label, ent in self.rain_entries.items():
                try:
                    p_num = float(p_label.replace("P=", "").replace("%", ""))
                    h_val = float(ent.get())
                    p_design.append(p_num)
                    hp_design.append(h_val)
                except ValueError:
                    pass

            if p_design:
                x_des = p_to_x(p_design)
                ax.scatter(x_des, hp_design, color="#d32f2f", s=55, zorder=5, label="Tần suất thiết kế (Hp gốc)")
                for px, hy, orig_p in zip(x_des, hp_design, p_design):
                    ax.annotate(f"{orig_p}%: {hy:.1f}", (px, hy), textcoords="offset points",
                                xytext=(6, 8), fontsize=8.5, fontweight="bold", color="#b71c1c")

            # Định dạng trục tọa độ
            proj_name = self.txt_project.get()
            ax.set_title(f"ĐỒ THỊ TẦN SUẤT MƯA PEARSON III - {proj_name}", fontsize=11, fontweight="bold", pad=12)
            ax.set_ylabel("Lượng mưa ngày lớn nhất Hp (mm)", fontsize=10, fontweight="bold")
            ax.grid(True, which="both", linestyle="--", linewidth=0.5, alpha=0.7)

            if mode == "Prob":
                ax.set_xlabel("Tần suất thiết kế P (%) - [Lưới xác suất Gauss]", fontsize=10, fontweight="bold")
                tick_x = p_to_x(prob_ticks)
                ax.set_xticks(tick_x)
                ax.set_xticklabels([f"{p}" for p in prob_ticks], fontsize=8.5, rotation=45)
                ax.set_xlim(p_to_x(0.04), p_to_x(99.5))
            elif mode == "Log":
                ax.set_xlabel("Tần suất P (%) [Thang Logarit]", fontsize=10, fontweight="bold")
                ax.set_xscale("log")
                ax.set_xlim(0.04, 105)
            else:
                ax.set_xlabel("Tần suất P (%) [Thang Tuyến tính]", fontsize=10, fontweight="bold")
                ax.set_xlim(0, 100)

            ax.legend(loc="upper right", fontsize=9, framealpha=0.92)
            fig.tight_layout()

            canvas = FigureCanvasTkAgg(fig, master=plot_container)
            canvas.draw()
            toolbar = NavigationToolbar2Tk(canvas, plot_container)
            toolbar.update()
            canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        scale_box = tk.Frame(ctrl_fr)
        scale_box.pack(side=tk.RIGHT)
        tk.Label(scale_box, text="Dạng lưới vẽ:", font=("Arial", 9, "bold")).pack(side=tk.LEFT, padx=3)
        tk.Radiobutton(scale_box, text="Lưới Gauss (Chuẩn Thủy văn)", variable=scale_var, value="Prob", command=render_plot).pack(side=tk.LEFT)
        tk.Radiobutton(scale_box, text="Logarit", variable=scale_var, value="Log", command=render_plot).pack(side=tk.LEFT, padx=3)
        tk.Radiobutton(scale_box, text="Tuyến tính", variable=scale_var, value="Linear", command=render_plot).pack(side=tk.LEFT)

        render_plot()

    def show_phi_table_window(self):
        if self.phi_win and self.phi_win.winfo_exists():self.phi_win.destroy()
        self.phi_win=tk.Toplevel(self.root);self.phi_win.title('Bảng A.1 — hệ số dòng chảy từ tài liệu cung cấp')
        self.phi_win.geometry('1180x620')
        tk.Label(self.phi_win,text='Cột 3–7: F<0,1; 8–9: 0,1≤F<1; 10–12: 1≤F<10; 13–14: 10≤F<100; 15: F>100.\nTài liệu Word không ghi mốc F riêng cho từng cột. Không tự gán các cột này thành điểm nội suy.',wraplength=1100).pack(pady=10)
        frame=ttk.Frame(self.phi_win);frame.pack(fill='both',expand=True)
        cols=['soil','rain']+[f'c{i}' for i in range(3,16)]
        tree=ttk.Treeview(frame,columns=cols,show='headings')
        for key in cols:
            tree.heading(key,text='Cấp đất' if key=='soil' else 'Cấp mưa (mm)' if key=='rain' else 'Cột '+key[1:])
            tree.column(key,width=105 if key=='rain' else 72,anchor='center')
        for grade in SOIL_TABLES:
            for index,row in enumerate(SOIL_TABLES[grade]):tree.insert('',tk.END,values=[grade,SOIL_H_LABELS[grade][index]]+[f'{v:.3f}' for v in row])
        tree.grid(row=0,column=0,sticky='nsew');frame.rowconfigure(0,weight=1);frame.columnconfigure(0,weight=1)
        sy=ttk.Scrollbar(frame,command=tree.yview);sy.grid(row=0,column=1,sticky='ns')
        sx=ttk.Scrollbar(frame,orient='horizontal',command=tree.xview);sx.grid(row=1,column=0,sticky='ew')
        tree.configure(yscrollcommand=sy.set,xscrollcommand=sx.set)

    def show_ts_table_window(self):
        if self.ts_win and self.ts_win.winfo_exists(): self.ts_win.destroy()
        self.ts_win = tk.Toplevel(self.root)
        self.ts_win.title("BẢNG II-2: THỜI GIAN NƯỚC CHẢY TRÊN SƯỜN DỐC ts")
        self.ts_win.geometry("980x520")
        tk.Label(self.ts_win, text="BẢNG II-2: THỜI GIAN NƯỚC CHẢY TRÊN SƯỜN DỐC  ts = f(Φs, Vùng mưa)", font=("Times New Roman", 13, "bold"), fg="#0277bd").pack(pady=(10, 4))
        cols = ["fs"] + [f"v_{i}" for i in range(1, 19)]
        tree = ttk.Treeview(self.ts_win, columns=cols, show="headings", height=18)
        tree.heading("fs", text="Φs"); tree.column("fs", width=55, anchor=tk.CENTER)
        for i in range(1, 19):
            tree.heading(f"v_{i}", text=f"Vùng {i}"); tree.column(f"v_{i}", width=50, anchor=tk.CENTER)
        for i, fs_val in enumerate(FS_AXIS):
            tree.insert("", tk.END, values=[f"{fs_val:.1f}"] + ["Thiếu" if i==len(FS_AXIS)-1 and j==17 else value for j,value in enumerate(TS_GRID[i])])
        tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

    def show_ap_table_window(self):
        if self.ap_win and self.ap_win.winfo_exists(): self.ap_win.destroy()
        self.ap_win = tk.Toplevel(self.root)
        self.ap_win.title("BẢNG MÔ-ĐUN TƯƠNG ĐỐI Ap = f(Vùng mưa, ts, Φl)")
        self.ap_win.geometry("980x360")
        vung = self.cbo_vung.get()
        tk.Label(self.ap_win, text=f"BẢNG MÔ-ĐUN TƯƠNG ĐỐI  Ap = f(Vùng mưa {vung}, ts, Φl)", font=("Times New Roman", 13, "bold"), fg="#6a1b9a").pack(pady=(10, 4))
        cols = ["ts"] + [f"fl_{c}" for c in FL_COLS]
        tree = ttk.Treeview(self.ap_win, columns=cols, show="headings", height=8)
        tree.heading("ts", text="ts (phút)"); tree.column("ts", width=70, anchor=tk.CENTER)
        for c in FL_COLS:
            tree.heading(f"fl_{c}", text=f"Φl={c}"); tree.column(f"fl_{c}", width=55, anchor=tk.CENTER)
        table_data = AP_TABLES.get(int(vung))
        if table_data is None:
            tk.Label(self.ap_win, text=f"Chưa tích hợp Bảng A.3 cho vùng mưa {vung} trong bộ dữ liệu hiện tại.", fg="#b71c1c", font=("Arial", 10, "bold")).pack(pady=20)
            return
        for idx, ts_val in enumerate(TS_LEVELS):
            tree.insert("", tk.END, values=[int(ts_val)] + table_data[idx])
        tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

    # --- TÍNH TOÁN LŨ Qp ---
    def calculate_all(self):
        self.refresh_phi_hints()
        self.tree.delete(*self.tree.get_children())
        self.calc_results=[];self._result_signature=None
        try:
            F = finite_number(self.txt_F.get(),"F")
            L = finite_number(self.txt_L.get(),"L")
            Sli = finite_number(self.txt_Sli.get(),"Sli")
            Jl = finite_number(self.txt_Jl.get(),"Jl")
            Js = finite_number(self.txt_Js.get(),"Js")
            ms = finite_number(self.txt_ms.get(),"ms")
            ml = finite_number(self.txt_ml.get(),"ml")
            d = finite_number(self.txt_d.get(),"d")
            vung = validate_region(self.cbo_vung.get())
            soil = self.cbo_soil.get()
            proj_name = self.txt_project.get()

            if F <= 0 or L <= 0 or Sli < 0 or Jl <= 0 or Js <= 0 or ms <= 0 or ml <= 0:
                raise ValueError("F, L, Jl, Js, ms, ml phải > 0; SigmaLi không được âm.")
            if F >= 100:
                raise ValueError(
                    "Phần mềm này chỉ tính lưu vực F < 100 km². Với F ≥ 100 km² cần chọn phương pháp phù hợp theo tiêu chuẩn. "
                    "Không dùng công thức cường độ giới hạn cho trường hợp này."
                )
            if not 1 <= vung <= 18:
                raise ValueError("Vùng mưa phải từ 1 đến 18.")
            if not 0<d<=1:
                raise ValueError("Hệ số giảm hồ đầm δ phải trong khoảng 0 < δ ≤ 1.")

            slope_text = self.cbo_slope.get()
            k_bs = 0.90 if slope_text.startswith("Một sườn") else 1.80
            if soil not in SOIL_TABLES:raise ValueError("Cấp đất phải là II, III, IV, V hoặc VI.")

            self.tree.delete(*self.tree.get_children())
            self.calc_results = []
            warnings = table_consistency_notes(vung,soil)
            result_rows=[];pending_results=[]

            for idx, (p_label, ent) in enumerate(self.rain_entries.items(), start=1):
                hp_goc = finite_number(ent.get(),f"Hp {p_label}")
                cc_val = finite_number(self.cc_entries[p_label].get(),f"Hệ số BĐKH {p_label}")
                if hp_goc <= 0:
                    raise ValueError(f"Hp gốc tại {p_label} phải > 0.")
                if cc_val <= 0:
                    raise ValueError(f"Hệ số BĐKH tại {p_label} phải > 0.")

                hp_val = hp_goc * cc_val

                phi_input=self.phi_entries[p_label].get().strip()
                ap_input=self.ap_entries[p_label].get().strip()
                confirmed_phi=finite_number(phi_input,'φ') if phi_input else None
                confirmed_ap=finite_number(ap_input,'Ap') if ap_input else None
                ts_input=self.ts_entries[p_label].get().strip()
                confirmed_ts=finite_number(ts_input,'ts') if ts_input else None
                computed=peak_flow_parameters(F,L,Sli,Jl,Js,ms,ml,d,hp_val,vung,soil,k_bs,confirmed_phi,confirmed_ap,confirmed_ts)
                phi,bs,Fs,ts,Fl,Ap,Qmax=(computed[k] for k in ('phi','bs','Fs','ts','Fl','Ap','Qmax'))
                for value,axis,name in ((Fs,FS_AXIS,'Φs'),(ts,TS_LEVELS,'ts'),(Fl,FL_COLS,'Φl')):
                    if not axis[0]<=value<=axis[-1]:
                        warnings.append(f'{p_label}: {name}={value:.3f} ngoài miền [{axis[0]}, {axis[-1]}]; tra giá trị biên, cần kiểm tra phạm vi bảng.')

                result_rows.append((
                    idx, p_label.replace("P=", "") + " %", f"{hp_goc:.2f}", f"{cc_val:.2f}",
                    f"{hp_val:.2f}", f"{phi:.3f}", f"{bs:.1f}", f"{Fs:.2f}", f"{ts:.1f}",
                    f"{Fl:.2f}", f"{Ap:.5f}", f"{Qmax:.2f}"
                ))

                pending_results.append({
                    "Công trình": proj_name, "STT": idx, "Tần suất P": p_label.replace("P=", "") + " %",
                    "Hp gốc (mm)": hp_goc, "Hệ số BĐKH": cc_val, "Hp tính toán (mm)": hp_val,
                    "Hệ số phi": phi, "bs (m)": bs, "Mẫu số bs": k_bs,
                    "Phi_s": Fs, "ts (phút)": ts, "Phi_l": Fl,
                    "Mô-đun Ap": Ap, "Lưu lượng Qmax (m3/s)": Qmax,
                    "Nguồn phi": "Người dùng xác minh" if confirmed_phi is not None else "Bảng A.1",
                    "Nguồn ts": "Người dùng xác minh" if confirmed_ts is not None else "Bảng A.2",
                    "Nguồn Ap": "Người dùng xác minh" if confirmed_ap is not None else "Bảng A.3"
                })

            for row in result_rows:self.tree.insert("",tk.END,values=row)
            self.calc_results=pending_results;self._result_signature=self.input_signature()

            if warnings:
                uniq = list(dict.fromkeys(warnings))
                messagebox.showwarning("Cảnh báo tra bảng", "\n".join(uniq[:12]) + ("\n..." if len(uniq) > 12 else ""))

        except ValueError as e:
            messagebox.showerror("Lỗi dữ liệu/tính toán", str(e))
        except Exception as e:
            messagebox.showerror("Lỗi tính toán", f"Không thể tính Qp: {e}")

    def copy_table(self):
        if not self.require_current_results():return
        if not self.calc_results: return
        proj = self.txt_project.get()
        out = f"BẢNG KẾT QUẢ TÍNH LŨ THIẾT KẾ: {proj}\n"
        out += "STT\tTan suat P\tHp goc\tHe so BDKH\tHp tinh toan\tHe so phi\tbs (m)\tPhi_s\tts (phut)\tPhi_l\tAp\tQmax (m3/s)\n"
        for r in self.calc_results:
            out += f"{r['STT']}\t{r['Tần suất P']}\t{r['Hp gốc (mm)']:.3f}\t{r['Hệ số BĐKH']:.3f}\t{r['Hp tính toán (mm)']:.3f}\t{r['Hệ số phi']:.3f}\t{r['bs (m)']:.3f}\t{r['Phi_s']:.3f}\t{r['ts (phút)']:.3f}\t{r['Phi_l']:.3f}\t{r['Mô-đun Ap']:.3f}\t{r['Lưu lượng Qmax (m3/s)']:.3f}\n"
        self.root.clipboard_clear(); self.root.clipboard_append(out)
        messagebox.showinfo("Thành công", "Đã sao chép bảng kết quả vào Clipboard!")

    def export_excel(self):
        if not self.require_current_results():return
        if not self.calc_results:
            messagebox.showwarning("Cảnh báo", "Không có dữ liệu để xuất!")
            return

        proj = self.txt_project.get()
        f_save = filedialog.asksaveasfilename(
            initialfile=f"Thuyet_Minh_Qp_{proj.replace(' ', '_')}.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not f_save:
            return

        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side
            import pandas as pd

            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            # 1. Sheet Tổng hợp
            ws_summary = wb.create_sheet(title="Bảng Tổng Hợp")
            df = pd.DataFrame(self.calc_results)

            for c_idx, col_name in enumerate(df.columns, 1):
                cell = ws_summary.cell(row=1, column=c_idx, value=col_name)
                cell.font = Font(bold=True)

            for r_idx, row in enumerate(df.values, 2):
                for c_idx, value in enumerate(row, 1):
                    ws_summary.cell(row=r_idx, column=c_idx, value=value)

            # 2. Định dạng
            font_bold_italic = Font(name='Times New Roman', size=12, bold=True, italic=True)
            font_normal = Font(name='Times New Roman', size=12)
            font_bold = Font(name='Times New Roman', size=12, bold=True)
            align_left = Alignment(horizontal='left', vertical='center')
            align_center = Alignment(horizontal='center', vertical='center')
            align_right = Alignment(horizontal='right', vertical='center')
            border_bottom = Border(bottom=Side(style='thin'))

            F_val = self.txt_F.get()
            L_val = self.txt_L.get()
            Sli_val = self.txt_Sli.get()
            Jl_val = self.txt_Jl.get()
            Js_val = self.txt_Js.get()
            soil = self.cbo_soil.get()
            ms_val = self.txt_ms.get()
            ml_val = self.txt_ml.get()
            d_val = self.txt_d.get()
            vung = self.cbo_vung.get()

            k_bs_str = "1.8" if self.cbo_slope.get().startswith("Hai sườn") else "0.9"

            # 3. Sheet Thuyết minh từng tần suất
            for res in self.calc_results:
                p_str = str(res['Tần suất P']).replace("%", "").strip()
                ws = wb.create_sheet(title=f"Thuyết minh P={p_str}%")

                ws.column_dimensions['A'].width = 45
                ws.column_dimensions['B'].width = 15
                ws.column_dimensions['C'].width = 5
                ws.column_dimensions['D'].width = 35
                ws.column_dimensions['E'].width = 5
                ws.column_dimensions['F'].width = 12
                ws.column_dimensions['G'].width = 10

                hp_val = res['Hp tính toán (mm)']
                phi_val = res['Hệ số phi']
                bs_val = res['bs (m)']
                fs_val = res['Phi_s']
                ts_val = res['ts (phút)']
                fl_val = res['Phi_l']
                ap_val = res['Mô-đun Ap']
                qmax_val = res['Lưu lượng Qmax (m3/s)']

                ws['A1'] = "I- Số liệu"
                ws['A1'].font = font_bold_italic

                data_I = [
                    ("Diện tích lưu vực", "Fl.v.", "=", F_val, "", "km²"),
                    ("Chiều dài lòng chủ", "Lch", "=", L_val, "", "km"),
                    ("Tổng chiều dài dòng nhánh", "Σli", "=", Sli_val, "", "km"),
                    ("Độ dốc lòng chủ", "Il", "=", Jl_val, "", "‰"),
                    ("Độ dốc sườn dốc", "Is", "=", Js_val, "", "‰"),
                    (f"Lượng mưa tính toán với tần suất X{p_str}%", f"X{p_str}%", "=", hp_val, "", "mm"),
                    ("Cấp đất trong lưu vực", "", "", soil, "", ""),
                    ("Hệ số dòng chảy lũ φ = f(cấp đất; Hp%; F)", "", "", "", "", ""),
                    (f"φ = f( {soil} ; {hp_val} ; {F_val} )", "", "=", phi_val, "", ""),
                    ("Độ nhám sườn dốc", "ms", "=", ms_val, "", ""),
                    ("Độ nhám lòng chính", "ml", "=", ml_val, "", ""),
                    ("Hệ số triết giảm do ao hồ", "δ", "=", d_val, "", ""),
                    ("Vùng mưa:", "", "", vung, "", ""),
                    ("Lưu lượng lũ thiết kế", f"Qmax.{p_str}%", "=", "?", "", "(m³/s)")
                ]

                r = 2
                for row_data in data_I:
                    ws.cell(row=r, column=1, value=row_data[0]).font = font_normal
                    if row_data[1]:
                        ws.cell(row=r, column=2, value=row_data[1]).font = font_normal
                        ws.cell(row=r, column=2).alignment = align_left
                    if row_data[2]:
                        ws.cell(row=r, column=3, value=row_data[2]).font = font_normal
                        ws.cell(row=r, column=3).alignment = align_center
                    if row_data[3]:
                        ws.cell(row=r, column=4, value=row_data[3]).font = font_normal
                        ws.cell(row=r, column=4).alignment = align_center
                    if row_data[5]:
                        ws.cell(row=r, column=6, value=row_data[5]).font = font_normal
                        ws.cell(row=r, column=6).alignment = align_left
                    r += 1

                r += 1
                ws.cell(row=r, column=1, value="II- Tính toán").font = font_bold_italic
                r += 1

                # 1. bs
                ws.cell(row=r, column=1, value="Chiều dài trung bình sườn dốc:").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value="bs").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value="1000.Fl.v.").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=4).border = border_bottom
                ws.cell(row=r+1, column=4, value=f"{k_bs_str} × (L + Σli)").font = font_normal; ws.cell(row=r+1, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=bs_val).font = font_normal; ws.cell(row=r, column=6).alignment = align_center
                r += 2

                # 2. Φs
                ws.cell(row=r, column=1, value="Đặc trưng địa mạo sườn dốc:").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value="Φs").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value="bs^0.6").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=4).border = border_bottom
                ws.cell(row=r+1, column=4, value="ms × Is^0.3 × (φ × Hp%)^0.4").font = font_normal; ws.cell(row=r+1, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=fs_val).font = font_normal; ws.cell(row=r, column=6).alignment = align_center
                r += 2

                # 3. ts
                ws.cell(row=r, column=1, value="Thời gian chảy trên sườn dốc τs:").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value="τs").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value=f"f( {fs_val} ; {vung} )").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=ts_val).font = font_normal; ws.cell(row=r, column=6).alignment = align_center
                ws.cell(row=r, column=7, value="phút").font = font_normal; ws.cell(row=r, column=7).alignment = align_left
                r += 1

                # 4. Φl
                ws.cell(row=r, column=1, value="Đặc trưng địa mạo lòng sông:").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value="Φl").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value="1000 × L").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=4).border = border_bottom
                ws.cell(row=r+1, column=4, value="ml × Il^0.333 × (F × φ × Hp%)^0.25").font = font_normal; ws.cell(row=r+1, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=fl_val).font = font_normal; ws.cell(row=r, column=6).alignment = align_center
                r += 2

                # 5. Ap
                ws.cell(row=r, column=1, value="Mô đun tương đối của dòng chảy lớn nhất:").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value="Ap").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value=f"f( {vung} ; {ts_val} ; {fl_val} )").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=ap_val).font = font_normal; ws.cell(row=r, column=6).alignment = align_center
                r += 1

                # 6. Qmax
                ws.cell(row=r, column=1, value="Lưu lượng lũ lớn nhất theo tần suất thiết kế (m³/s)").font = font_normal
                r += 1
                ws.cell(row=r, column=2, value=f"Qmax.{p_str}%").font = font_normal; ws.cell(row=r, column=2).alignment = align_center
                ws.cell(row=r, column=3, value="=").font = font_normal; ws.cell(row=r, column=3).alignment = align_center
                ws.cell(row=r, column=4, value="Ap × φ × Hp × F × δ").font = font_normal; ws.cell(row=r, column=4).alignment = align_center
                ws.cell(row=r, column=5, value="=").font = font_normal; ws.cell(row=r, column=5).alignment = align_center
                ws.cell(row=r, column=6, value=qmax_val).font = font_bold; ws.cell(row=r, column=6).alignment = align_center
                ws.cell(row=r, column=7, value="(m³/s)").font = font_normal
                r += 1

            for sheet in wb.worksheets:
                for cells in sheet.iter_rows():
                    for cell in cells:
                        if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):cell.number_format='0.000'
            wb.save(f_save)
            messagebox.showinfo("Thành công", f"Đã xuất file báo cáo tính toán:\n{f_save}")

        except ImportError:
            messagebox.showerror("Thiếu thư viện", "Vui lòng chạy lệnh 'pip install openpyxl' trong Terminal để sử dụng định dạng xuất Excel này.")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Có lỗi khi xuất Excel: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = FullTCVNApp(root)
    root.mainloop()
