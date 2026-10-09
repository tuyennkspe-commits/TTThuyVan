#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Qp_Sokolovsky v2 - Tính đỉnh lũ thiết kế theo phương pháp Xô-kô-lôp-ski (F > 100 km²)

Cấu trúc:
  1. Cơ sở dữ liệu 8 bảng tra (giữ nguyên số liệu bản gốc, có bổ sung kiểm tra bất thường)
  2. Lõi tính toán thuần (không phụ thuộc Tkinter -> kiểm thử được)
  3. Xuất Excel có công thức sống (kiểm toán được)
  4. Giao diện Tkinter (tự tính lại khi đổi số liệu, không bật hộp thoại liên tục)

Yêu cầu: Python >= 3.9, numpy, scipy, openpyxl, matplotlib (tkinter đi kèm Python)
"""
from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass

import numpy as np
from scipy.stats import norm, pearson3, skew

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:  # giao diện là tùy chọn để lõi tính toán vẫn import được trên máy không có Tk
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    HAS_TK = True
except ImportError:  # pragma: no cover
    HAS_TK = False

try:
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
    HAS_MPL = True
except ImportError:  # pragma: no cover
    HAS_MPL = False

APP_TITLE = "TÍNH ĐỈNH LŨ THIẾT KẾ - PHƯƠNG PHÁP XÔ-KÔ-LÔP-SKI (F > 100 km²)"

# =============================================================================
# 1. CƠ SỞ DỮ LIỆU 8 BẢNG TRA (số liệu giữ nguyên từ bản gốc - CẦN đối chiếu lại với tiêu chuẩn)
# =============================================================================

# 1. Bảng 2-10: Tra alpha và Ho theo phân vùng lưu vực
BANG_2_10 = {
    1: {"name": "Lưu vực sông Nậm Rốn và thượng nguồn sông Mã", "alpha": 0.65, "Ho": 20},
    2: {"name": "Lưu vực sông Đà, sông Thao", "alpha": 0.81, "Ho": 22},
    3: {"name": "Các Lưu vực thượng nguồn sông Lô, sông Chảy", "alpha": 0.82, "Ho": 20},
    4: {"name": "Sông Gâm, hạ Lưu sông Lô, sông Phó Đáy", "alpha": 0.66, "Ho": 26},
    5: {"name": "Lưu vực sông Cầu, sông Thương, sông Trung, sông Bằng Giang", "alpha": 0.77, "Ho": 22},
    6: {"name": "Lưu vực sông Kỳ Cùng, sông Lục Nam", "alpha": 0.86, "Ho": 19},
    7: {"name": "Lưu vực các sông Quảng Ninh", "alpha": 0.89, "Ho": 15},
    8: {"name": "Lưu vực các sông từ sông Chu - sông Hương", "alpha": 0.92, "Ho": 21},
    9: {"name": "Lưu vực các sông từ Thu Bồn - sông Cái", "alpha": 0.86, "Ho": 16},
    10: {"name": "Lưu vực các sông Sê San và sông Srêpốk", "alpha": 0.76, "Ho": 21},
    11: {"name": "Lưu vực các sông Đồng Nai, sông Bé", "alpha": 0.64, "Ho": 25}
}

# 2. Bảng C.13: Tra hệ số hình dạng lũ f
BANG_C13 = {
    "I": {"name": "Các lưu vực sông nhỏ tỉnh Quảng Ninh", "f_range": "1.2 - 1.5", "f_def": 1.35},
    "II": {"name": "Lưu vực sông Thái Bình", "f_range": "0.6 - 0.8", "f_def": 0.70},
    "III": {"name": "Lưu vực sông Bằng Giang - Kỳ Cùng, sông Gâm", "f_range": "0.6 - 0.8", "f_def": 0.70},
    "IV": {"name": "Lưu vực sông Lô - Chảy", "f_range": "0.7 - 0.9", "f_def": 0.80},
    "V": {"name": "Lưu vực sông Thao, sông Đà", "f_range": "0.7 - 0.9", "f_def": 0.80},
    "VI": {"name": "Thượng nguồn sông Mã, thượng nguồn sông Mê Kông", "f_range": "0.6 - 0.7", "f_def": 0.65},
    "VII": {"name": "Trung & hạ lưu sông Mã + Lưu vực sông Cả", "f_range": "0.6 - 0.8", "f_def": 0.70},
    "VIII": {"name": "Các lưu vực sông nhỏ tỉnh Quảng Bình", "f_range": "0.8 - 1.2", "f_def": 1.00},
    "IX": {"name": "Từ Quảng Trị đến Quảng Ngãi", "f_range": "0.5 - 0.7", "f_def": 0.60},
    "X": {"name": "Khu vực Bắc Tây Nguyên", "f_range": "1.3 - 1.5", "f_def": 1.40},
    "XI": {"name": "Từ Bình Định đến Bình Thuận & Nam Tây Nguyên", "f_range": "0.5 - 0.9", "f_def": 0.70},
    "XII": {"name": "Lưu vực sông Đồng Nai + ĐBSCL", "f_range": "0.5 - 0.7", "f_def": 0.60}
}

# 3. Tọa độ đường cong mưa rào psi(T) của 18 vùng mưa Việt Nam
TIME_STEPS = [10, 15, 20, 30, 45, 60, 90, 120, 240, 480, 540, 720, 1080, 1440]
RAINFALL_CURVE_COORDS = {
    1: [0.180, 0.220, 0.260, 0.340, 0.430, 0.490, 0.610, 0.660, 0.800, 0.940, 0.950, 0.960, 0.980, 1.070],
    2: [0.130, 0.180, 0.220, 0.250, 0.330, 0.350, 0.400, 0.440, 0.580, 0.770, 0.790, 0.880, 0.900, 1.090],
    3: [0.070, 0.090, 0.120, 0.140, 0.200, 0.220, 0.270, 0.300, 0.440, 0.630, 0.680, 0.780, 0.830, 1.070],
    4: [0.150, 0.210, 0.240, 0.320, 0.380, 0.470, 0.550, 0.600, 0.920, 0.820, 0.830, 0.880, 0.930, 1.060],
    5: [0.1005, 0.120, 0.150, 0.226, 0.300, 0.378, 0.460, 0.537, 0.700, 0.924, 0.935, 0.952, 0.985, 1.055],
    6: [0.120, 0.140, 0.180, 0.260, 0.300, 0.380, 0.470, 0.590, 0.780, 0.920, 0.950, 0.990, 1.030, 1.200],
    7: [0.098, 0.110, 0.176, 0.214, 0.240, 0.322, 0.419, 0.508, 0.682, 0.857, 0.890, 0.912, 0.950, 1.110],
    8: [0.125, 0.160, 0.200, 0.268, 0.320, 0.408, 0.504, 0.594, 0.734, 0.890, 0.920, 0.994, 1.040, 1.160],
    9: [0.100, 0.120, 0.150, 0.220, 0.250, 0.320, 0.390, 0.460, 0.590, 0.810, 0.830, 0.890, 0.930, 1.050],
    10: [0.080, 0.110, 0.130, 0.190, 0.230, 0.300, 0.380, 0.460, 0.640, 0.820, 0.835, 0.900, 0.965, 1.160],
    11: [0.060, 0.080, 0.102, 0.130, 0.170, 0.187, 0.260, 0.305, 0.415, 0.617, 0.670, 0.827, 0.935, 1.040],
    12: [0.078, 0.102, 0.118, 0.115, 0.2054, 0.240, 0.3025, 0.335, 0.500, 0.660, 0.710, 0.825, 1.060, 1.095],
    13: [0.098, 0.280, 0.145, 0.795, 0.245, 0.302, 0.380, 0.440, 0.630, 0.770, 0.830, 0.870, 0.970, 1.090],
    14: [0.160, 0.232, 0.295, 0.360, 0.420, 0.590, 0.665, 0.680, 0.790, 0.890, 0.960, 0.940, 0.965, 1.005],
    15: [0.255, 0.310, 0.463, 0.510, 0.540, 0.570, 0.610, 0.690, 0.766, 0.820, 0.840, 0.905, 0.960, 1.020],
    16: [0.230, 0.320, 0.417, 0.530, 0.700, 0.780, 0.830, 0.850, 0.870, 0.950, 0.965, 0.980, 0.990, 1.030],
    17: [0.205, 0.220, 0.250, 0.330, 0.380, 0.480, 0.580, 0.660, 0.730, 0.890, 0.910, 1.035, 1.045, 1.050],
    18: [0.190, 0.285, 0.330, 0.430, 0.520, 0.610, 0.715, 0.935, 0.780, 0.880, 0.900, 0.980, 1.030, 1.150]
}

# 4. Bảng tra hệ số km theo vùng khí hậu
BANG_KM = {
    "Đông Bắc": 0.55,
    "Tây Bắc": 0.60,
    "Bắc Trung Bộ": 0.55,
    "Nam Trung Bộ": 0.60,
    "Tây Nguyên": 0.70,
    "Nam Bộ": 0.80
}

# 5. Bảng tra vận tốc lớn nhất trên sườn dốc Vmax
BANG_VMAX = {
    "Đồng bằng rất bằng phẳng": 0.5,
    "Đồng bằng": 1.0,
    "Trung du thấp": 1.5,
    "Đồi núi thấp": 2.0,
    "Núi trung bình": 2.5,
    "Núi cao, dốc lớn": 3.0,
    "Núi rất dốc": 4.0
}

# 6. Bảng tra hệ số truyền lũ lòng sông kv (22TCN 220)
BANG_KV = {
    "Sông quanh co nhiều, lòng dẫn kém phát triển": 0.50,
    "Phát triển trung bình": 0.55,
    "Phát triển khá": 0.60,
    "Phát triển tốt": 0.65,
    "Phát triển rất tốt, lòng dẫn rõ ràng": 0.70
}

# 7. Bảng hệ số nhám lòng sông thiên nhiên
BANG_NHAM_SONG = [
    (1, "Sông bờ nhẵn nhụi, dòng thẳng không trở ngại, nước chảy dễ dàng", 40.0, 0.025, 1.20),
    (2, "Sông vùng đồng bằng luôn có nước, lòng sông và nước chảy đặc biệt tốt", 30.0, 0.035, 2.00),
    (3, "Sông đồng bằng nước tương đối sạch, đáy có bãi nổi, hố xói, đá lác đác", 25.0, 0.040, 2.75),
    (4, "Sông quanh co, có nhiều trở ngại cục bộ, cây cỏ mọc hoặc đá tròn to", 20.0, 0.050, 3.75),
    (5, "Sông cực kỳ trở ngại, khúc khuỷu, bãi sông không phẳng, nhiều đá cuội to", 15.0, 0.060, 5.50),
    (6, "Sông có bãi, cây cỏ mọc đặc biệt rậm rạp, nước chảy chậm, có vực sâu", 12.5, 0.080, 7.00),
    (7, "Sông miền núi có nhiều đá lớn, nước chảy sinh bọt tung toé, khúc khuỷu", 12.5, 0.080, 7.00),
    (8, "Sông miền núi có thác, lòng sông khúc khuỷu, đá to, nước chảy réo vang", 10.0, 0.100, 9.00),
    (9, "Sông núi cây cối rậm rạp, nhiều chỗ nước ứ đọng, khúc chết sâu rộng", 7.5, 0.133, 12.00),
    (10, "Sông có bùn đá trôi, bãi sông cây lớn mọc rậm rạp dày đặc", 5.0, 0.200, 20.00)
]

# 8. Bảng 3: Ranh giới phân vùng mưa rào Việt Nam
BANG_PHAN_VUNG_MUA = {
    1: "Lưu vực thượng nguồn sông Mã, sông Chu, sông Cả",
    2: "Vùng thượng nguồn sông Đà từ biên giới đến Nghĩa Lộ",
    3: "Tâm mưa Hoàng Liên Sơn hữu ngạn sông Thao, từ biên giới đến Ngòi Hút",
    4: "Vùng lưu vực sông Kỳ Cùng, sông Bằng Giang, thượng nguồn sông Hồng",
    5: "Lưu vực sông Gâm, tả ngạn sông Lô",
    6: "Thung lũng sông Thao, sông Chảy, hạ lưu sông Lô - Gâm",
    7: "Các lưu vực bắt nguồn từ dãy Yên Tử đổ ra biển",
    8: "Vùng ven biển từ Hải Phòng đến Thanh Hóa",
    9: "Các lưu vực phần trung du sông Mã, sông Chu ra đến biển",
    10: "Vùng ven biển từ Thanh Hóa đến Đồng Hới",
    11: "Vùng ven biển từ Đồng Hới đến Đà Nẵng",
    12: "Vùng ven biển từ Đà Nẵng đến Quảng Ngãi",
    13: "Vùng ven biển từ Quảng Ngãi đến Phan Rang",
    14: "Các lưu vực sông phía bắc Tây Nguyên",
    15: "Các lưu vực sông phía nam Tây Nguyên",
    16: "Các lưu vực sông từ Ban Mê Thuột tới Bảo Lộc",
    17: "Vùng ven biển từ Phan Rang đến Vũng Tàu",
    18: "Vùng đồng bằng Nam Bộ"
}

# =============================================================================
# 2. LÕI TÍNH TOÁN (thuần Python - không phụ thuộc giao diện)
# =============================================================================
KP_MANUAL = "(nhập tay)"
DEFAULT_P = [1.0, 1.5, 4.0, 5.0, 10.0, 50.0]
PLOTTING_POSITIONS = {  # m: thứ hạng giảm dần (1 = lớn nhất), n: số năm -> tần suất vượt (0..1)
    "(m−0,25)/(n+0,5)  [như bản gốc]": lambda m, n: (m - 0.25) / (n + 0.5),
    "Weibull  m/(n+1)": lambda m, n: m / (n + 1.0),
    "Hazen  (m−0,5)/n": lambda m, n: (m - 0.5) / n,
    "Cunnane  (m−0,4)/(n+0,2)": lambda m, n: (m - 0.4) / (n + 0.2),
}
CS_MODES = ["Theo mẫu (n−3)", "Cs = 2 Cv", "Cs = 2,5 Cv", "Cs = 3 Cv", "Nhập tay", "Như bản cũ"]


def parse_float(text, name="giá trị") -> float:
    """Chấp nhận cả dấu phẩy thập phân (78,98)."""
    try:
        value = float(str(text).strip().replace(",", "."))
        if not math.isfinite(value): raise ValueError("Số phải hữu hạn")
        return value
    except ValueError:
        raise ValueError(f"Ô «{name}» không phải số hợp lệ: “{text}”") from None


def parse_key(text) -> str:
    """'13 – Vùng ven biển ...' -> '13'."""
    return str(text).split(" – ")[0].strip()


def parse_p_list(text) -> list[float]:
    vals = [parse_float(t, "danh sách P") for t in re.split(r"[\s,;]+", text.strip()) if t]
    if not vals:
        raise ValueError("Danh sách tần suất P đang trống.")
    if any(not (0 < p < 100) for p in vals):
        raise ValueError("Tần suất P phải nằm trong khoảng (0; 100) %.")
    return sorted(set(vals))


def p_label(p: float) -> str:
    return f"P={p:g}%"


# ---- Kiểm tra bảng tọa độ ψ(T) -------------------------------------------------
def _step_issue(y0: float, y1: float, t1: int, max_jump: float = 0.25):
    if y1 < y0:
        return f"ψ giảm tại {t1}′ ({y0:.3f} → {y1:.3f})"
    if y1 - y0 > max_jump:
        return f"ψ tăng đột ngột tại {t1}′ ({y0:.3f} → {y1:.3f})"
    return None


def find_curve_anomalies() -> dict[int, list[str]]:
    """ψ(T) phải tăng đơn điệu theo thời đoạn. Trả về các vùng có số liệu đáng ngờ (nghi nhập sai)."""
    out = {}
    for z, y in RAINFALL_CURVE_COORDS.items():
        msgs = [m for i in range(1, len(y)) if (m := _step_issue(y[i - 1], y[i], TIME_STEPS[i]))]
        if msgs:
            out[z] = msgs
    return out


CURVE_ANOMALIES = find_curve_anomalies()


def curve_bracket_issue(zone: int, t_min: float):
    """Nếu đoạn nội suy chứa t_min rơi đúng chỗ số liệu đáng ngờ thì trả về mô tả."""
    y = RAINFALL_CURVE_COORDS[zone]
    t = min(max(t_min, TIME_STEPS[0]), TIME_STEPS[-1])
    i = max(1, min(int(np.searchsorted(TIME_STEPS, t)), len(TIME_STEPS) - 1))
    return _step_issue(y[i - 1], y[i], TIME_STEPS[i])


def psi_of(T_hours: float, zone: int, log_time: bool = False):
    """Nội suy ψ(T). Trả về (ψ, bị_chặn_biên). Ngoài [10′; 1440′] giá trị bị chặn ở biên (không ngoại suy)."""
    t = T_hours * 60.0
    clamped = t < TIME_STEPS[0] or t > TIME_STEPS[-1]
    t = min(max(t, TIME_STEPS[0]), TIME_STEPS[-1])
    y = RAINFALL_CURVE_COORDS[zone]
    if log_time:
        return float(np.interp(math.log(t), np.log(TIME_STEPS), y)), clamped
    return float(np.interp(t, TIME_STEPS, y)), clamped


# ---- Thống kê mưa (Pearson III) ------------------------------------------------
def kp_pearson3(p_percent: float, cv: float, cs: float) -> float:
    """Kp = 1 + Cv·Φ(P, Cs) với Φ là phân vị chuẩn hóa CHÍNH XÁC của Pearson III (tương đương bảng Foster-Rưpkin)."""
    if not all(math.isfinite(v) for v in (p_percent,cv,cs)) or not 0<p_percent<100 or cv<0:raise ValueError("P, Cv, Cs không hợp lệ")
    phi = float(pearson3.ppf(1.0 - p_percent / 100.0, skew=cs))
    value=1.0+cv*phi
    if not math.isfinite(value) or value<0:raise ValueError("Phân vị mưa âm hoặc không hữu hạn; kiểm tra phân phối")
    return value


def kp_legacy(p_percent: float, cv: float, cs: float) -> float:
    if not all(math.isfinite(v) for v in (p_percent,cv,cs)) or not 0<p_percent<100 or cv<0:raise ValueError("P, Cv, Cs không hợp lệ")
    """Công thức XẤP XỈ của bản cũ - chỉ để đối chiếu với kết quả cũ (sai lệch tới vài % ở P nhỏ)."""
    u = float(norm.ppf(1.0 - p_percent / 100.0))
    if abs(cs) < 0.01:
        return 1.0 + cv * u
    return max(1.0 + cv * (u + (u ** 2 - 1.0) * cs / 6.0 + (u ** 3 - 3.0 * u) * cs ** 2 / 27.0), 0.0)


def legacy_stats(x):
    """Cách bản cũ tính khi nạp tệp: Cv, Cs làm tròn 2 số lẻ; Cs = scipy skew(bias=False), nếu < 0,1 thì Cs = 3Cv."""
    x = np.asarray(x, float)
    mean = float(x.mean())
    cv = round(float(x.std(ddof=1)) / mean, 2)
    cs = round(float(skew(x, bias=False)), 2)
    if cs < 0.1:
        cs = round(3.0 * cv, 2)
    return mean, cv, cs


def sample_stats(x):
    """X̄, Cv (n−1), Cs = Σ(K−1)³ / ((n−3)·Cv³) - công thức thường dùng trong thủy văn Việt Nam."""
    x = np.asarray(x, float)
    n = len(x)
    if n<4 or not np.all(np.isfinite(x)) or np.any(x<0) or x.mean()<=0:raise ValueError("Cần ít nhất 4 giá trị mưa hữu hạn, không âm, trung bình dương")
    mean = float(x.mean())
    k = x / mean
    cv = math.sqrt(float(((k - 1) ** 2).sum()) / (n - 1))
    cs = float(((k - 1) ** 3).sum()) / ((n - 3) * cv ** 3) if (n > 3 and cv > 0) else 0.0
    return mean, cv, cs, n


def parse_series(text: str) -> list[float]:
    """Mỗi dòng: 'giá_trị' hoặc 'năm giá_trị' (lấy cột 2 nếu có ≥ 2 cột, như bản gốc). Bỏ qua dòng tiêu đề.
    Chấp nhận dấu phẩy thập phân; dùng dấu chấm phẩy hoặc khoảng trắng giữa năm và mưa."""
    vals=[]
    for number,line in enumerate(text.splitlines(),1):
        line=line.strip()
        if not line:continue
        if not vals and re.match(r'^(năm|year|rain|mưa)',line,re.I):continue
        parts=line.split(';') if ';' in line else line.split()
        if len(parts)==1 and re.match(r'^\d{4},',line):parts=line.split(',')
        if len(parts) not in (1,2):raise ValueError(f'Dòng {number}: dùng một lượng mưa hoặc năm và lượng mưa.')
        value=parse_float(parts[-1],f'mưa dòng {number}')
        if value<0:raise ValueError(f'Dòng {number}: mưa không được âm.')
        vals.append(value)
    return vals


def delta_from_areas(fa: float, fl: float, fr: float) -> float:
    if not all(math.isfinite(v) and 0<=v<=100 for v in (fa,fl,fr)):raise ValueError("Tỷ lệ diện tích phải trong 0–100% và hữu hạn")
    return 1.0 - 0.6 * math.log10(1.0 + fa + 0.2 * fl + 0.05 * fr)


# ---- Lưu lượng đỉnh lũ ---------------------------------------------------------
@dataclass
class Basin:
    F: float
    Ls: float
    Ho: float
    alpha: float
    f: float
    delta: float
    Vmax: float
    kv: float
    km: float
    kt: float
    zone: int
    log_time: bool = False
    Qng: float = 0.0

    def validate(self):
        errs, warns = [], []
        for name in ("F","Ls","Ho","alpha","f","delta","Vmax","kv","km","kt","Qng"):
            if not math.isfinite(getattr(self,name)):errs.append(f"{name} phải là số hữu hạn")
        if self.Qng<0:errs.append("Qng phải ≥ 0")
        for name, val in [("Flv", self.F), ("Ls", self.Ls), ("Vmax", self.Vmax), ("kv", self.kv),
                          ("km", self.km), ("kt", self.kt), ("f", self.f)]:
            if not val > 0:
                errs.append(f"{name} phải > 0")
        if self.Ho < 0:
            errs.append("Ho phải ≥ 0")
        if not (0 < self.alpha <= 1):
            errs.append("α phải trong (0; 1]")
        if not (0 < self.delta <= 1):
            errs.append("δ phải trong (0; 1]")
        if self.zone not in RAINFALL_CURVE_COORDS:
            errs.append(f"Không có vùng mưa {self.zone}")
        if not errs and self.F <= 100:
            warns.append(f"Flv = {self.F:g} km² ≤ 100 km²: ngoài phạm vi áp dụng của phương pháp")
        return errs, warns


def derive(b: Basin) -> dict:
    errs,_=b.validate()
    if errs:raise ValueError("; ".join(errs))
    v_tb = b.kv * b.Vmax                      # m/s
    ts = b.Ls / (v_tb * 3.6)                  # giờ (Ls km, v m/s)
    tl = ts                                   # như bản gốc: thời gian lũ lên = ts
    T = ts * b.km * b.kt                      # giờ
    psi, clamped = psi_of(T, b.zone, b.log_time)
    return dict(v_tb=v_tb, ts=ts, tl=tl, T=T, T_min=T * 60.0, psi=psi, clamped=clamped,
                reduce=1.0 + 0.001 * b.F ** 0.8, issue=curve_bracket_issue(b.zone, T * 60.0))


def run_case(b: Basin, d: dict, rows, use_cc: bool):
    """rows: [(label, p, h24_goc, cc)]. Trả về danh sách dict kết quả."""
    out = []
    for label, p, h_goc, cc in rows:
        if not all(math.isfinite(v) for v in (p,h_goc,cc)) or not 0<p<100 or h_goc<=0 or cc<=0:
            raise ValueError(f"{label}: tần suất phải trong (0;100), mưa và hệ số khí hậu phải hữu hạn và dương.")
        h24 = h_goc * (cc if use_cc else 1.0)
        htp = h24 * d["psi"]
        hred = htp / d["reduce"]
        eff = max(hred - b.Ho, 0.0)
        q = 0.278 * b.alpha * eff * b.f * b.F * b.delta / d["tl"] + b.Qng
        if not math.isfinite(q):raise ValueError("Kết quả Q không hữu hạn")
        out.append(dict(label=label, p=p, h24_goc=h_goc, cc=cc if use_cc else 1.0, h24=h24,
                        psi=d["psi"], htp=htp, hred=hred, q=q, no_runoff=hred <= b.Ho))
    return out


# =============================================================================
# 3. XUẤT EXCEL - CÓ CÔNG THỨC SỐNG
# =============================================================================
def _safe_name(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_") or "DuAn"


def build_sheet(ws, title: str, res: dict, with_cc: bool):
    b, d = res["basin"], res["derived"]
    cases = res["bdkh"] if with_cc else res["hien_trang"]
    n = len(cases)
    fn = "Times New Roman"
    f_t, f_r = Font(name=fn, size=12, bold=True), Font(name=fn, size=11)
    f_b, f_i = Font(name=fn, size=11, bold=True), Font(name=fn, size=11, italic=True)
    f_h = Font(name=fn, size=11, bold=True, color="FFFFFF")
    fill_h = PatternFill("solid", start_color="203764", end_color="203764")
    fill_in = PatternFill("solid", start_color="FFFDE7", end_color="FFFDE7")
    s = Side(border_style="thin", color="000000")
    box = Border(left=s, right=s, top=s, bottom=s)
    ctr, lft, rgt = (Alignment(horizontal=h, vertical="center") for h in ("center", "left", "right"))

    ws["A1"] = f"{title} - {res['project']}"
    ws["A1"].font = f_t
    ws["A2"] = "I. THÔNG SỐ ĐẦU VÀO LƯU VỰC VÀ THỦY VĂN (ô vàng: số liệu nhập; ô còn lại: công thức)"
    ws["A2"].font = f_t

    # Bố cục cố định: hàng 3..18 -> C3..C18
    params = [
        ("Diện tích lưu vực", "Flv", b.F, "km²", "0.00"),                          # C3
        ("Chiều dài sông chính", "Ls", b.Ls, "km", "0.00"),                        # C4
        ("Lớp nước tổn thất ban đầu", "Ho", b.Ho, "mm", "0.0"),                    # C5
        ("Vận tốc lớn nhất trên sườn dốc", "Vmax", b.Vmax, "m/s", "0.00"),         # C6
        ("H/S truyền lũ lòng sông", "kv", b.kv, "", "0.000"),                      # C7
        ("H/S thời gian theo vùng khí hậu", "km", b.km, "", "0.000"),              # C8
        ("H/S thời gian mưa rào", "kt", b.kt, "", "0.000"),                        # C9
        ("Hệ số dòng chảy lũ", "α", b.alpha, "", "0.000"),                         # C10
        ("Hệ số hình dạng lũ", "f", b.f, "", "0.000"),                             # C11
        ("Hệ số triết giảm ao hồ, đầm lầy", "δ", b.delta, "", "0.0000"),           # C12
        ("Thời gian tập trung nước trong sông", "ts", "=C4/(C7*C6*3.6)", "giờ", "0.00"),   # C13
        ("Thời gian lũ lên", "tl", "=C13", "giờ", "0.00"),                         # C14
        ("Thời gian mưa tính toán", "T", "=C13*C8*C9", "giờ", "0.00"),             # C15
        ("Phân vùng mưa rào", "R", b.zone, "", "0"),                               # C16
        ("Tung độ ψ(T) (nội suy bảng tọa độ vùng R)", "ψ", d["psi"], "", "0.0000"),  # C17
        ("Hệ số triết giảm diện tích 1+0,001·F^0,8", "ϑ", "=1+0.001*C3^0.8", "", "0.0000"),  # C18
    ]
    params.append(("Lưu lượng nước ngầm", "Qng", b.Qng, "m³/s", "0.000"))
    for r, (name, sym, val, unit, fmt) in enumerate(params, start=3):
        cells = [ws.cell(r, 1, name), ws.cell(r, 2, sym), ws.cell(r, 3, val), ws.cell(r, 4, unit)]
        for c in cells:
            c.border = box
        cells[0].font, cells[0].alignment = f_r, lft
        cells[1].font, cells[1].alignment = f_i, ctr
        cells[2].font, cells[2].alignment, cells[2].number_format = f_b, ctr, fmt
        if not (isinstance(val, str) and val.startswith("=")) and sym != "ψ":
            cells[2].fill = fill_in
        cells[3].font, cells[3].alignment = f_r, lft

    r0 = 20
    ws.cell(r0, 1, "II. KẾT QUẢ TÍNH TOÁN THEO TẦN SUẤT THIẾT KẾ P (%)").font = f_t
    hdr = ["Thông số tính toán", "Ký hiệu (Đơn vị)"] + [c["label"].replace("P=", "") for c in cases]
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(r0 + 1, j, h)
        c.font, c.fill, c.alignment, c.border = f_h, fill_h, ctr, box

    rows = [("Tần suất thiết kế", "P (%)", "p")]
    if with_cc:
        rows += [("Lượng mưa ngày (số liệu gốc)", "H24p gốc (mm)", "h24_goc"),
                 ("Hệ số biến đổi khí hậu", "Kbđkh", "cc"),
                 ("Lượng mưa ngày thiết kế xét BĐKH", "H24p (mm)", "h24")]
    else:
        rows += [("Lượng mưa ngày thiết kế", "H24p (mm)", "h24_goc")]
    rows += [("Tung độ đường cong mưa rào", "ψ(T)", "psi"),
             ("Lượng mưa rào thiết kế", "Htp (mm)", "htp"),
             ("Lượng mưa rào thiết kế triết giảm", "H'tp (mm)", "hred"),
             ("Lưu lượng đỉnh lũ thiết kế", "Qmax,p (m³/s)", "q")]
    R = {key: r0 + 2 + i for i, (_, _, key) in enumerate(rows)}
    h24_row = R["h24"] if with_cc else R["h24_goc"]
    fmts = dict(p='0.0"%"', h24_goc="0.0", cc="0.000", h24="0.0", psi="0.0000", htp="0.0", hred="0.0", q="#,##0")

    for name, sym, key in rows:
        r = R[key]
        bold = key == "q"
        a, s_ = ws.cell(r, 1, name), ws.cell(r, 2, sym)
        for c, al in ((a, lft), (s_, ctr)):
            c.font, c.alignment, c.border = (f_b if bold else f_r), al, box
        for j, case in enumerate(cases):
            col = get_column_letter(3 + j)
            if key == "p":
                v = case["p"]
            elif key in ("h24_goc", "cc") and (with_cc or key == "h24_goc"):
                v = case[key]
            elif key == "h24" and with_cc:
                v = f"={col}{R['h24_goc']}*{col}{R['cc']}"
            elif key == "psi":
                v = "=$C$17"
            elif key == "htp":
                v = f"={col}{h24_row}*$C$17"
            elif key == "hred":
                v = f"={col}{R['htp']}/$C$18"
            else:  # q
                v = (f"=0.278*$C$10*MAX({col}{R['hred']}-$C$5,0)*$C$11*$C$3*$C$12/$C$14+$C$19")
            c = ws.cell(r, 3 + j, v)
            c.font, c.alignment, c.border, c.number_format = (f_b if bold else f_r), rgt, box, fmts[key]
            if key in ("p", "h24_goc", "cc") and not str(v).startswith("="):
                c.fill = fill_in

    rn = R["q"] + 2
    ws.cell(rn, 1, "Công thức: Qmax,p = 0,278·α·(H'tp − Ho)·f·F·δ / tl + Qng  (thành phần dòng chảy mặt = 0 nếu H'tp ≤ Ho)").font = f_i
    for k, w in enumerate(res["warnings"], start=1):
        ws.cell(rn + k, 1, f"⚠ {w}").font = Font(name=fn, size=10, italic=True, color="B71C1C")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    for j in range(3, 3 + max(n, 2)):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.freeze_panes = "C3"


def export_workbook(path: str, res: dict):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Hiện trạng"
    build_sheet(ws1, "KẾT QUẢ TÍNH LŨ THEO XÔ-KÔ-LÔP-SKI - HIỆN TRẠNG", res, with_cc=False)
    ws2 = wb.create_sheet("Xét BĐKH")
    build_sheet(ws2, "KẾT QUẢ TÍNH LŨ THEO XÔ-KÔ-LÔP-SKI - XÉT BĐKH", res, with_cc=True)
    wb.save(path)


# =============================================================================
# 4. GIAO DIỆN TKINTER
# =============================================================================
MANUAL_BG, AUTO_BG = "#fffde7", "#e1f5fe"
P_TICKS = [0.01, 0.1, 0.5, 1, 2, 3, 5, 10, 15, 20, 30, 40, 50, 60, 70, 80, 85, 90, 95, 97, 99, 99.9, 99.99]
FREQ_TABLE_P = [0.01, 0.05, 0.1, 0.2, 0.5, 1, 1.5, 2, 3, 4, 5, 10, 20, 25, 30, 50, 75, 80, 90, 95, 99, 99.9]

DEFAULTS = dict(
    project="Cầu Phú Kiểng", Flv="1855", Ls="78.98", kt="1.0",
    b210=None, alpha="0.86", Ho="16",            # vùng 9 (Thu Bồn - sông Cái) cho đúng α, Ho
    c13=None, f="0.7",                           # vùng XI (Bình Định - Bình Thuận) cho đúng f
    vmax_sel="Đồi núi thấp", Vmax="2.0",
    kv_sel="Phát triển tốt", kv="0.65",
    km_sel=KP_MANUAL, km="0.565",
    vung=None,
    Qng="0.0", fa="0.1", fl="0.1", fr="6.9", delta="0.9005",
    mean_x="107.2", cv="0.24", cs="0.72", cs_mode=CS_MODES[4], pp=list(PLOTTING_POSITIONS)[0],
    p_list="1, 1.5, 4, 5, 10, 50", cc_all="1.135", log_time=False, auto_calc=True, kp_legacy=False,
    drawer="", checker="", year="",
)
DEFAULT_H24 = {1.0: 398.77, 1.5: 373.86, 4.0: 314.74, 5.0: 298.64, 10.0: 253.87, 50.0: 137.68}


def lab_b210(k): return f"{k} – {BANG_2_10[k]['name']}"
def lab_c13(k): return f"{k} – {BANG_C13[k]['name']}"
def lab_zone(k): return f"{k} – {BANG_PHAN_VUNG_MUA[k]}"


class SokolovskyApp:
    def __init__(self, root):
        self.root = root
        root.title(APP_TITLE)
        root.geometry("1060x760")
        root.minsize(900, 620)
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.v = {}
        self.sample = None            # np.ndarray số liệu mưa ngày lớn nhất năm (nếu đã nạp)
        self.rain_rows = []
        self.results = None
        self._after = None
        self._internal = False
        self._init_vars()
        self._build_menu()
        self._build_ui()
        self.apply_p_list([float(p) for p in DEFAULT_H24], list(DEFAULT_H24.values()),
                          [DEFAULTS["cc_all"]] * len(DEFAULT_H24))
        self.calculate()

    # ------------------------------------------------------------------ vars
    def _init_vars(self):
        d = dict(DEFAULTS, b210=lab_b210(9), c13=lab_c13("XI"), vung=lab_zone(13))
        for k, val in d.items():
            var = tk.BooleanVar(value=val) if isinstance(val, bool) else tk.StringVar(value=str(val))
            var.trace_add("write", self._changed)
            self.v[k] = var

    def _changed(self, *_):
        if not self._internal:self.results=None
        if self._internal or not self.v["auto_calc"].get():
            return
        if self._after:
            self.root.after_cancel(self._after)
        self._after = self.root.after(200, self.calculate)   # gộp nhiều thay đổi liên tiếp

    def _set(self, key, value, silent=False):
        if silent:
            self._internal = True
        try:
            self.v[key].set(value)
        finally:
            self._internal = False

    def num(self, key, name=None):
        return parse_float(self.v[key].get(), name or key)

    # ------------------------------------------------------------------ menu
    def _build_menu(self):
        mb = tk.Menu(self.root)
        m = tk.Menu(mb, tearoff=0)
        m.add_command(label="💾 Lưu dự án…", command=self.save_project)
        m.add_command(label="📂 Mở dự án…", command=self.open_project)
        m.add_command(label="📊 Xuất Excel (có công thức)…", command=self.export_excel)
        m.add_separator()
        m.add_command(label="Thoát", command=self.on_close)
        mb.add_cascade(label="Tệp", menu=m)

        t = tk.Menu(mb, tearoff=0)
        items = [
            ("Bảng 2-10 (α, Ho)", lambda: self.show_table(
                "Bảng 2-10: α và Ho theo phân vùng lưu vực", ["Vùng", "Lưu vực sông / địa danh", "α", "Ho (mm)"],
                [(k, v["name"], v["alpha"], v["Ho"]) for k, v in BANG_2_10.items()], [60, 500, 80, 80])),
            ("Bảng C.13 (hệ số f)", lambda: self.show_table(
                "Bảng C.13: hệ số hình dạng lũ f", ["Vùng", "Lưu vực sông / địa danh", "Khoảng f", "f mặc định"],
                [(k, v["name"], v["f_range"], v["f_def"]) for k, v in BANG_C13.items()], [60, 500, 90, 90])),
            ("Tọa độ đường cong mưa rào ψ(T)", self.show_curve_table),
            ("Bảng 3: phân vùng mưa rào", lambda: self.show_table(
                "Bảng 3: ranh giới phân vùng mưa rào", ["Vùng", "Ranh giới"],
                list(BANG_PHAN_VUNG_MUA.items()), [60, 760])),
            ("Hệ số nhám lòng sông (1/n, N, γ)", lambda: self.show_table(
                "Hệ số nhám lòng sông thiên nhiên", ["TT", "Đặc điểm lòng sông", "1/n", "N", "γ"],
                BANG_NHAM_SONG, [40, 600, 60, 60, 60])),
            ("Hệ số km", lambda: self.show_table("Hệ số km theo vùng khí hậu", ["Vùng khí hậu", "km"],
                                                  list(BANG_KM.items()), [300, 100])),
            ("Vận tốc Vmax trên sườn dốc", lambda: self.show_table(
                "Vận tốc lớn nhất trên sườn dốc", ["Địa hình", "Vmax (m/s)"], list(BANG_VMAX.items()), [360, 100])),
            ("Hệ số truyền lũ kv (22TCN 220)", lambda: self.show_table(
                "Hệ số truyền lũ lòng sông kv", ["Đặc điểm lòng dẫn", "kv"], list(BANG_KV.items()), [460, 100])),
        ]
        for label, cmd in items:
            t.add_command(label=label, command=cmd)
        mb.add_cascade(label="Tra cứu bảng chuẩn", menu=t)
        h = tk.Menu(mb, tearoff=0)
        h.add_command(label="Kiểm tra số liệu bảng ψ(T)", command=self.show_curve_check)
        mb.add_cascade(label="Kiểm tra", menu=h)
        self.root.config(menu=mb)

    # -------------------------------------------------------------------- UI
    def _build_ui(self):
        nb = self.nb = ttk.Notebook(self.root)
        nb.pack(fill="both", expand=True, padx=8, pady=(8, 0))
        self.tab1, self.tab2, self.tab3 = (ttk.Frame(nb) for _ in range(3))
        nb.add(self.tab1, text="  1. Lưu vực & thông số  ")
        nb.add(self.tab2, text="  2. Mưa thiết kế (Pearson III)  ")
        nb.add(self.tab3, text="  3. Kết quả  ")
        self._build_basin_tab()
        self._build_rain_tab()
        self._build_result_tab()
        self.status = tk.Label(self.root, anchor="w", justify="left", wraplength=1020, padx=8, pady=3, relief="sunken")
        self.status.pack(fill="x", side="bottom")
        bar = ttk.Frame(self.root, padding=(8, 6))
        bar.pack(fill="x", side="bottom")
        tk.Button(bar, text="🚀 TÍNH TOÁN LŨ SOKOLOVSKY", font=("Arial", 10, "bold"), bg="#154360", fg="white",
                  padx=16, pady=4, command=self.run_clicked).pack(side="left")
        ttk.Checkbutton(bar, text="Tự động tính lại khi sửa số liệu", variable=self.v["auto_calc"]).pack(side="left", padx=14)
        ttk.Button(bar, text="📋 Sao chép kết quả", command=self.copy_results).pack(side="right", padx=4)
        ttk.Button(bar, text="📊 Xuất Excel", command=self.export_excel).pack(side="right", padx=4)

    def _entry(self, parent, key, width=10, auto=False, **kw):
        return tk.Entry(parent, textvariable=self.v[key], width=width, bg=AUTO_BG if auto else MANUAL_BG,
                        fg="#0277bd" if auto else "#b71c1c", font=("Arial", 10, "bold"), **kw)

    def _build_basin_tab(self):
        g = ttk.Frame(self.tab1, padding=12)
        g.pack(fill="x")
        g.columnconfigure(1, weight=1)
        self._r = 0

        def row(label, key, unit="", auto=False, combo=None, handler=None, width=10, entry=True):
            r = self._r
            ttk.Label(g, text=label).grid(row=r, column=0, sticky="w", pady=3)
            if combo is not None:
                cb = ttk.Combobox(g, textvariable=self.v[combo[0]], values=combo[1], state="readonly", width=58)
                cb.grid(row=r, column=1, sticky="we", padx=8)
                if handler:
                    cb.bind("<<ComboboxSelected>>", lambda e, h=handler: h())
            if entry:
                self._entry(g, key, width, auto).grid(row=r, column=2, sticky="w")
            ttk.Label(g, text=unit).grid(row=r, column=3, sticky="w", padx=6)
            self._r += 1

        ttk.Label(g, text="Công trình / vị trí:").grid(row=0, column=0, sticky="w")
        tk.Entry(g, textvariable=self.v["project"], width=40, font=("Arial", 10, "bold")).grid(row=0, column=1, sticky="w", padx=8)
        self._r = 1
        ttk.Label(g, text="Màu vàng: nhập tay  |  màu xanh: tra bảng/tính tự động (vẫn sửa được)",
                  foreground="#555").grid(row=self._r, column=0, columnspan=3, sticky="w", pady=(2, 8))
        self._r += 1
        row("Diện tích lưu vực Flv", "Flv", "km²")
        row("Chiều dài sông chính Ls", "Ls", "km")
        row("H/S thời gian mưa rào kt", "kt", "")
        row("α, Ho - Bảng 2-10", "alpha", "α", True, ("b210", [lab_b210(k) for k in BANG_2_10]), self.pick_b210)
        row("Tổn thất ban đầu Ho", "Ho", "mm", True)
        row("Lưu lượng nước ngầm Qng", "Qng", "m³/s")
        row("Hệ số hình dạng lũ f - Bảng C.13", "f", "f", True, ("c13", [lab_c13(k) for k in BANG_C13]), self.pick_c13)
        row("Địa hình sườn dốc → Vmax", "Vmax", "m/s", True,
            ("vmax_sel", [KP_MANUAL] + list(BANG_VMAX)), lambda: self.pick_simple("vmax_sel", BANG_VMAX, "Vmax"))
        row("Đặc điểm lòng sông → kv", "kv", "", True,
            ("kv_sel", [KP_MANUAL] + list(BANG_KV)), lambda: self.pick_simple("kv_sel", BANG_KV, "kv"))
        row("Vùng khí hậu → km", "km", "", True,
            ("km_sel", [KP_MANUAL] + list(BANG_KM)), lambda: self.pick_simple("km_sel", BANG_KM, "km"))
        row("Phân vùng mưa rào - Bảng 3 & ψ(T)", "vung", "", False,
            ("vung", [lab_zone(k) for k in BANG_PHAN_VUNG_MUA]), None, entry=False)

        r = self._r
        ttk.Label(g, text="Triết giảm ao hồ, rừng δ").grid(row=r, column=0, sticky="w", pady=6)
        fr = ttk.Frame(g)
        fr.grid(row=r, column=1, sticky="w", padx=8)
        for lab, k in (("fa", "fa"), ("fl", "fl"), ("fr", "fr")):
            ttk.Label(fr, text=f"{lab}:").pack(side="left")
            tk.Entry(fr, textvariable=self.v[k], width=6, bg=MANUAL_BG).pack(side="left", padx=(2, 8))
        ttk.Button(fr, text="⚡ Tính δ từ fa, fl, fr", command=self.recalc_delta).pack(side="left", padx=8)
        self._entry(g, "delta", 10, True).grid(row=r, column=2, sticky="w")
        ttk.Label(g, text="δ").grid(row=r, column=3, sticky="w", padx=6)
        ttk.Checkbutton(g, text="Nội suy ψ(T) theo log(thời gian) (mặc định: tuyến tính như bản gốc)",
                        variable=self.v["log_time"]).grid(row=r + 1, column=0, columnspan=3, sticky="w", pady=4)

        box = ttk.LabelFrame(self.tab1, text=" Đại lượng trung gian ", padding=10)
        box.pack(fill="x", padx=12, pady=8)
        self.lbl_derived = ttk.Label(box, text="", font=("Consolas", 10), justify="left")
        self.lbl_derived.pack(anchor="w")

    def _build_rain_tab(self):
        t = self.tab2
        top = ttk.LabelFrame(t, text=" Thống kê chuỗi mưa ngày lớn nhất năm ", padding=8)
        top.pack(fill="x", padx=10, pady=8)
        r1 = ttk.Frame(top); r1.pack(fill="x", pady=2)
        ttk.Button(r1, text="📁 Nạp tệp trạm mưa (.txt/.csv)", command=self.load_rain_file).pack(side="left")
        self.lbl_sample = ttk.Label(r1, text="  (chưa nạp chuỗi số liệu)")
        self.lbl_sample.pack(side="left")
        r2 = ttk.Frame(top); r2.pack(fill="x", pady=4)
        for lab, k, w in (("X̄ (mm):", "mean_x", 8), ("Cv:", "cv", 6), ("Cs:", "cs", 6)):
            ttk.Label(r2, text=lab).pack(side="left", padx=(0, 2))
            self._entry(r2, k, w).pack(side="left", padx=(0, 10))
        ttk.Label(r2, text="Cs xác định theo:").pack(side="left")
        cb = ttk.Combobox(r2, textvariable=self.v["cs_mode"], values=CS_MODES, state="readonly", width=16)
        cb.pack(side="left", padx=4)
        cb.bind("<<ComboboxSelected>>", lambda e: self.apply_cs_mode())
        r3 = ttk.Frame(top); r3.pack(fill="x", pady=2)
        ttk.Label(r3, text="Công thức tần suất kinh nghiệm (vẽ điểm):").pack(side="left")
        ttk.Combobox(r3, textvariable=self.v["pp"], values=list(PLOTTING_POSITIONS), state="readonly", width=32).pack(side="left", padx=4)
        ttk.Checkbutton(r3, text="Kp xấp xỉ như bản cũ (chỉ để đối chiếu)", variable=self.v["kp_legacy"]).pack(side="left", padx=12)
        ttk.Button(r3, text="📋 Bảng tần suất lý luận", command=self.show_freq_table).pack(side="right", padx=4)
        ttk.Button(r3, text="📈 Vẽ đường tần suất", command=self.plot_frequency).pack(side="right", padx=4)

        mid = ttk.LabelFrame(t, text=" Lượng mưa ngày thiết kế H24p theo tần suất ", padding=8)
        mid.pack(fill="both", expand=True, padx=10, pady=4)
        r4 = ttk.Frame(mid); r4.pack(fill="x", pady=2)
        ttk.Label(r4, text="Các tần suất P (%):").pack(side="left")
        tk.Entry(r4, textvariable=self.v["p_list"], width=34, bg=MANUAL_BG).pack(side="left", padx=4)
        ttk.Button(r4, text="Áp dụng", command=lambda: self.apply_p_list()).pack(side="left", padx=4)
        ttk.Label(r4, text="   Hệ số BĐKH chung:").pack(side="left")
        tk.Entry(r4, textvariable=self.v["cc_all"], width=7, bg=MANUAL_BG).pack(side="left", padx=4)
        ttk.Button(r4, text="Gán cho mọi P", command=self.apply_cc_all).pack(side="left")
        ttk.Button(r4, text="⚡ Gán H24p từ X̄, Cv, Cs", command=self.fill_from_theory).pack(side="right")
        self.rain_grid = ttk.Frame(mid)
        self.rain_grid.pack(fill="x", pady=6)
        ttk.Label(mid, text="Cột «H24p theo X̄, Cv, Cs» để đối chiếu: nếu lệch nhiều so với H24p đang dùng thì số liệu "
                            "trạm mưa và bảng nhập chưa khớp nhau.", foreground="#555", wraplength=980).pack(anchor="w")

        sig = ttk.LabelFrame(t, text=" Ghi chú bản vẽ đường tần suất ", padding=6)
        sig.pack(fill="x", padx=10, pady=6)
        for lab, k, w in (("Người vẽ:", "drawer", 22), ("Người kiểm tra:", "checker", 22), ("Năm:", "year", 8)):
            ttk.Label(sig, text=lab).pack(side="left")
            tk.Entry(sig, textvariable=self.v[k], width=w).pack(side="left", padx=(2, 12))

    def _build_result_tab(self):
        t = self.tab3
        self.trees = {}
        specs = [
            ("ht", " II. KẾT QUẢ HIỆN TRẠNG ", ("p", "h24", "psi", "htp", "hred", "q"),
             ("P", "H24p (mm)", "ψ(T)", "Htp (mm)", "H'tp triết giảm (mm)", "Qmax,p (m³/s)"), (90, 110, 90, 110, 150, 130)),
            ("cc", " III. KẾT QUẢ XÉT BIẾN ĐỔI KHÍ HẬU ", ("p", "goc", "cc", "h24", "psi", "htp", "hred", "q"),
             ("P", "H24p gốc", "Hệ số BĐKH", "H24p BĐKH", "ψ(T)", "Htp (mm)", "H'tp (mm)", "Qmax,p (m³/s)"),
             (80, 90, 90, 100, 80, 90, 90, 120)),
        ]
        for key, title, cols, heads, widths in specs:
            fr = ttk.LabelFrame(t, text=title, padding=6)
            fr.pack(fill="both", expand=True, padx=10, pady=6)
            tv = ttk.Treeview(fr, columns=cols, show="headings", height=6)
            for c, h, w in zip(cols, heads, widths):
                tv.heading(c, text=h)
                tv.column(c, width=w, anchor="e" if c != "p" else "center")
            tv.tag_configure("zero", foreground="#b71c1c")
            tv.pack(fill="both", expand=True)
            self.trees[key] = tv

    # -------------------------------------------------------------- handlers
    def recalc_delta(self):
        """Chỉ đổi δ khi bấm nút (giống bản gốc); gõ tay δ thì giữ nguyên số đã gõ."""
        try:
            self._set("delta", f"{delta_from_areas(self.num('fa'), self.num('fl'), self.num('fr')):.4f}")
        except ValueError as e:
            self.show_status(str(e), "error")

    def kp(self, p, cv, cs):
        return (kp_legacy if self.v["kp_legacy"].get() else kp_pearson3)(p, cv, cs)

    def pick_b210(self):
        k = int(parse_key(self.v["b210"].get()))
        self._set("alpha", str(BANG_2_10[k]["alpha"]))
        self._set("Ho", str(BANG_2_10[k]["Ho"]))

    def pick_c13(self):
        k = parse_key(self.v["c13"].get())
        self._set("f", str(BANG_C13[k]["f_def"]))

    def pick_simple(self, sel_key, table, target):
        sel = self.v[sel_key].get()
        if sel in table:
            self._set(target, str(table[sel]))

    def apply_cc_all(self):
        for r in self.rain_rows:
            r["cc"].set(self.v["cc_all"].get())

    def apply_p_list(self, ps=None, hs=None, ccs=None):
        """Dựng lại lưới nhập H24p theo danh sách P. Giữ giá trị cũ nếu P đã có; P mới lấy theo lý thuyết."""
        try:
            if ps is None:
                ps = parse_p_list(self.v["p_list"].get())
                old = {r["p"]: (r["h"].get(), r["cc"].get()) for r in self.rain_rows}
                hs, ccs = [], []
                for p in ps:
                    if p in old:
                        hs.append(old[p][0]); ccs.append(old[p][1])
                    else:
                        hs.append(self._theory_h(p)); ccs.append(self.v["cc_all"].get())
            self._set("p_list", ", ".join(f"{p:g}" for p in ps), silent=True)
        except ValueError as e:
            self.show_status(str(e), "error")
            return
        for w in self.rain_grid.winfo_children():
            w.destroy()
        for j, h in enumerate(("P (%)", "H24p gốc (mm)", "Hệ số BĐKH", "H24p theo X̄, Cv, Cs")):
            ttk.Label(self.rain_grid, text=h, font=("Arial", 9, "bold")).grid(row=0, column=j, padx=8, pady=2)
        self.rain_rows = []
        for i, (p, h, cc) in enumerate(zip(ps, hs, ccs), start=1):
            hv, cv_ = tk.StringVar(value=str(h)), tk.StringVar(value=str(cc))
            hv.trace_add("write", self._changed)
            cv_.trace_add("write", self._changed)
            ttk.Label(self.rain_grid, text=f"{p:g}").grid(row=i, column=0)
            tk.Entry(self.rain_grid, textvariable=hv, width=10, justify="center", bg=MANUAL_BG).grid(row=i, column=1, padx=8, pady=1)
            tk.Entry(self.rain_grid, textvariable=cv_, width=10, justify="center", bg=MANUAL_BG).grid(row=i, column=2, padx=8)
            lt = ttk.Label(self.rain_grid, text="-", width=14, anchor="center")
            lt.grid(row=i, column=3)
            self.rain_rows.append(dict(p=p, h=hv, cc=cv_, theory=lt))
        self._changed()

    def _theory_h(self, p):
        try:
            return f"{self.num('mean_x', 'X̄') * self.kp(p, self.num('cv', 'Cv'), self.num('cs', 'Cs')):.2f}"
        except ValueError:
            return "0"

    def fill_from_theory(self):
        try:
            self.num("mean_x", "X̄"); self.num("cv", "Cv"); self.num("cs", "Cs")
        except ValueError as e:
            self.show_status(str(e), "error")
            return
        for r in self.rain_rows:
            r["h"].set(self._theory_h(r["p"]))

    def apply_cs_mode(self):
        mode = self.v["cs_mode"].get()
        try:
            cv = self.num("cv", "Cv")
            if mode.startswith("Cs = "):
                k = float(mode.split("=")[1].replace("Cv", "").replace(",", ".").strip())
                self._set("cs", f"{k * cv:.2f}")
            elif mode == "Như bản cũ":
                if self.sample is None:
                    self.show_status("Chưa nạp chuỗi số liệu.", "warn")
                    return
                _, cv_l, cs_l = legacy_stats(self.sample)
                self._set("cv", f"{cv_l:.2f}", silent=True)
                self._set("cs", f"{cs_l:.2f}")
            elif mode.startswith("Theo mẫu"):
                if self.sample is None:
                    self.show_status("Chưa nạp chuỗi số liệu để tính Cs theo mẫu.", "warn")
                    return
                cs = sample_stats(self.sample)[2]
                if math.isnan(cs):
                    self.show_status("Cần n > 3 năm để tính Cs theo mẫu.", "warn")
                    return
                self._set("cs", f"{cs:.2f}")
        except ValueError as e:
            self.show_status(str(e), "error")

    def load_rain_file(self):
        path = filedialog.askopenfilename(filetypes=[("Text/CSV", "*.txt *.csv"), ("Tất cả", "*.*")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                vals = parse_series(fh.read())
            if len(vals) < 4:
                messagebox.showwarning("Cảnh báo", "Chuỗi mưa cần ít nhất 4 năm (để tính được Cs theo mẫu).")
                return
            self.sample = np.array(vals)
            mean, cv, cs, n = sample_stats(self.sample)
            self._set("mean_x", f"{mean:.4f}", silent=True)
            self._set("cv", f"{cv:.3f}", silent=True)
            self._set("cs", f"{cs:.2f}", silent=True)
            self._set("cs_mode", CS_MODES[0], silent=True)
            self.lbl_sample.config(text=f"  Đã nạp n = {n} năm | X̄ = {mean:.2f} | Cv = {cv:.3f} | Cs(mẫu) = {cs:.2f}  "
                                        f"(Cs/Cv = {cs / cv if cv else 0:.1f}; với n nhỏ Cs mẫu kém ổn định, cân nhắc Cs = 2-3 Cv)")
            if messagebox.askyesno("Nạp thành công", "Gán lại H24p theo X̄, Cv, Cs vừa tính?"):
                self.fill_from_theory()
            else:
                self._changed()
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không đọc được tệp: {e}")

    # ------------------------------------------------------------- calculate
    def show_status(self, msg, level="ok"):
        colors = {"ok": ("#e8f5e9", "#1b5e20"), "warn": ("#fff8e1", "#e65100"), "error": ("#ffebee", "#b71c1c")}
        bg, fg = colors[level]
        self.status.config(text=msg, bg=bg, fg=fg)

    def read_basin(self) -> Basin:
        zone = int(parse_key(self.v["vung"].get()))
        return Basin(F=self.num("Flv"), Ls=self.num("Ls"), Ho=self.num("Ho"), alpha=self.num("alpha", "α"),
                     f=self.num("f"), delta=self.num("delta", "δ"), Vmax=self.num("Vmax"), kv=self.num("kv"),
                     km=self.num("km"), kt=self.num("kt"), zone=zone, log_time=self.v["log_time"].get(), Qng=self.num("Qng"))

    def calculate(self):
        if self._after:
            self.root.after_cancel(self._after)
            self._after = None
        for tv in self.trees.values():
            tv.delete(*tv.get_children())
        self.results = None
        try:
            b = self.read_basin()
            rows = [(p_label(r["p"]), r["p"], parse_float(r["h"].get(), f"H24p (P={r['p']:g}%)"),
                     parse_float(r["cc"].get(), f"BĐKH (P={r['p']:g}%)")) for r in self.rain_rows]
        except (ValueError, KeyError) as e:
            self.show_status(f"✖ {e}", "error")
            return
        errs, warns = b.validate()
        if errs:
            self.show_status("✖ " + "; ".join(errs), "error")
            return
        if any(not math.isfinite(v) for r in rows for v in r[1:]) or any(not 0<r[1]<100 or r[2]<=0 or r[3]<=0 for r in rows):
            self.show_status("Mưa và hệ số khí hậu phải hữu hạn, dương; P trong (0;100).", "error");return
        d = derive(b)
        if d["clamped"]:
            warns.append(f"T = {d['T_min']:.0f} phút nằm ngoài bảng ψ(T) [10′; 1440′] → ψ bị chặn ở biên, không ngoại suy")
        if d["issue"]:
            warns.append(f"Vùng mưa {b.zone}: đoạn nội suy ψ đang dùng có số liệu đáng ngờ - {d['issue']} "
                         f"(xem menu Kiểm tra → bảng ψ(T))")
        ht, cc = run_case(b, d, rows, False), run_case(b, d, rows, True)
        for case, name in ((ht, "hiện trạng"), (cc, "xét BĐKH")):
            zero = [c["label"] for c in case if c["no_runoff"]]
            if zero:
                warns.append(f"H'tp ≤ Ho ở {', '.join(zero)} ({name}) → Qmax = Qng (mưa không đủ vượt tổn thất)")
        self.results = dict(project=self.v["project"].get(), basin=b, derived=d, hien_trang=ht, bdkh=cc, warnings=warns)

        for c in ht:
            self.trees["ht"].insert("", "end", tags=("zero",) if c["no_runoff"] else (),
                                    values=(c["label"], f"{c['h24']:.2f}", f"{c['psi']:.4f}", f"{c['htp']:.2f}",
                                            f"{c['hred']:.2f}", f"{c['q']:,.1f}"))
        for c in cc:
            self.trees["cc"].insert("", "end", tags=("zero",) if c["no_runoff"] else (),
                                    values=(c["label"], f"{c['h24_goc']:.2f}", f"{c['cc']:.3f}", f"{c['h24']:.2f}",
                                            f"{c['psi']:.4f}", f"{c['htp']:.2f}", f"{c['hred']:.2f}", f"{c['q']:,.1f}"))
        self.lbl_derived.config(text=(
            f"v_tb = kv·Vmax = {d['v_tb']:.3f} m/s      ts = Ls/(3,6·v_tb) = {d['ts']:.2f} giờ      tl = {d['tl']:.2f} giờ\n"
            f"T = ts·km·kt = {d['T']:.2f} giờ = {d['T_min']:.0f} phút      ψ(T) = {d['psi']:.4f}      "
            f"triết giảm = 1+0,001·F^0,8 = {d['reduce']:.4f}"))
        for r in self.rain_rows:
            r["theory"].config(text=self._theory_h(r["p"]))
        self.show_status("⚠ " + " | ".join(warns) if warns else "✔ Đã tính xong, không có cảnh báo.", "warn" if warns else "ok")

    def run_clicked(self):
        self.calculate()
        if self.results is not None:
            self.nb.select(self.tab3)      # chuyển sang tab Kết quả

    def _ensure_results(self):
        self.calculate()
        if self.results is None:
            messagebox.showerror("Lỗi", "Số liệu đầu vào chưa hợp lệ - xem thanh trạng thái phía dưới.")
            return None
        return self.results

    # ---------------------------------------------------------- table windows
    def show_table(self, title, cols, rows, widths=None, note=None, tagged=(), button=None):
        w = tk.Toplevel(self.root)
        w.title(title)
        w.geometry(f"{min(sum(widths or [800]) + 60, 1200)}x520")
        ttk.Label(w, text=title, font=("Arial", 12, "bold")).pack(pady=6)
        ids = [f"c{i}" for i in range(len(cols))]
        tv = ttk.Treeview(w, columns=ids, show="headings")
        for i, c in enumerate(cols):
            long_text = any(len(str(r[i])) > 25 for r in rows)
            tv.heading(ids[i], text=c)
            tv.column(ids[i], width=(widths or [120] * len(cols))[i], anchor="w" if long_text else "center")
        tv.tag_configure("warn", background="#fff3cd")
        for k, r in enumerate(rows):
            tv.insert("", "end", values=list(r), tags=("warn",) if k in tagged else ())
        sb = ttk.Scrollbar(w, orient="vertical", command=tv.yview)
        tv.configure(yscroll=sb.set)
        if note:
            ttk.Label(w, text=note, wraplength=900, foreground="#b71c1c").pack(side="bottom", anchor="w", padx=10, pady=4)
        if button:
            ttk.Button(w, text=button[0], command=button[1]).pack(side="bottom", pady=4)
        sb.pack(side="right", fill="y")
        tv.pack(fill="both", expand=True, padx=(10, 0), pady=4)

    def show_curve_table(self):
        rows = [[k] + [f"{x:.3f}" for x in v] for k, v in RAINFALL_CURVE_COORDS.items()]
        tagged = [i for i, k in enumerate(RAINFALL_CURVE_COORDS) if k in CURVE_ANOMALIES]
        note = ("Hàng tô vàng: có ψ giảm hoặc tăng đột ngột theo thời đoạn - nghi nhập sai so với tiêu chuẩn. "
                + "; ".join(f"Vùng {k}: {', '.join(v)}" for k, v in CURVE_ANOMALIES.items()))
        self.show_table("Tọa độ đường cong mưa rào ψ(T)", ["Vùng"] + [f"{t}′" for t in TIME_STEPS], rows,
                        [50] + [58] * len(TIME_STEPS), note if tagged else None, tagged)

    def show_curve_check(self):
        if not CURVE_ANOMALIES:
            messagebox.showinfo("Kiểm tra", "Bảng ψ(T) tăng đơn điệu ở mọi vùng.")
            return
        msg = "\n".join(f"Vùng {k}: " + "; ".join(v) for k, v in CURVE_ANOMALIES.items())
        messagebox.showwarning("Số liệu ψ(T) đáng ngờ",
                               msg + "\n\nHãy đối chiếu với bảng gốc trong tiêu chuẩn và sửa RAINFALL_CURVE_COORDS.")

    def show_freq_table(self):
        try:
            mean, cv, cs = self.num("mean_x", "X̄"), self.num("cv", "Cv"), self.num("cs", "Cs")
        except ValueError as e:
            messagebox.showerror("Lỗi", str(e))
            return
        rows = [(f"{p:g} %", f"{self.kp(p, cv, cs):.4f}", f"{mean * self.kp(p, cv, cs):.2f}") for p in FREQ_TABLE_P]

        def export():
            path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
            if path:
                wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tần suất lý luận"
                ws.append(["P (%)", "Kp", "H24p (mm)"])
                for p in FREQ_TABLE_P:
                    ws.append([p, round(self.kp(p, cv, cs), 4), round(mean * self.kp(p, cv, cs), 2)])
                wb.save(path)
                messagebox.showinfo("Thành công", f"Đã xuất: {path}")
        self.show_table(f"Tần suất lý luận Pearson III (X̄ = {mean:g}, Cv = {cv:g}, Cs = {cs:g})",
                        ["P (%)", "Kp", "H24p (mm)"], rows, [160, 160, 160], button=("💾 Xuất Excel", export))

    # ------------------------------------------------------------------ plot
    def plot_frequency(self):
        if not HAS_MPL:
            messagebox.showwarning("Thiếu matplotlib", "Cần cài matplotlib để vẽ đường tần suất.")
            return
        try:
            mean, cv, cs = self.num("mean_x", "X̄"), self.num("cv", "Cv"), self.num("cs", "Cs")
        except ValueError as e:
            messagebox.showerror("Lỗi", str(e))
            return
        w = tk.Toplevel(self.root)
        w.title("Đường tần suất lượng mưa ngày lớn nhất năm (Pearson III)")
        w.geometry("1100x720")
        fig = Figure(figsize=(11, 7), dpi=100)          # Figure trực tiếp: không rò rỉ pyplot khi đóng cửa sổ
        ax = fig.add_subplot(111)
        to_x = lambda p: norm.ppf(np.asarray(p, float) / 100.0)
        pd_ = np.concatenate([np.linspace(0.01, 1, 80), np.linspace(1, 99, 400), np.linspace(99, 99.99, 80)])
        ax.plot(to_x(pd_), mean * (1 + cv * pearson3.ppf(1 - pd_ / 100, skew=cs)), "k-", lw=1.8, label="Đường Pearson III lý luận")
        n_obs = 0
        if self.sample is not None:
            xs = np.sort(self.sample)[::-1]
            n_obs = len(xs)
            m = np.arange(1, n_obs + 1)
            pp = PLOTTING_POSITIONS[self.v["pp"].get()](m, n_obs) * 100
            ax.scatter(to_x(pp), xs, c="black", s=16, zorder=5, label="Điểm thực đo")
        pts = [(r["p"], parse_float(r["h"].get(), "H24p")) for r in self.rain_rows]
        if pts:
            ax.scatter(to_x([p for p, _ in pts]), [h for _, h in pts], c="#c62828", marker="D", s=40, zorder=6,
                       label="H24p đang dùng để tính lũ")
        ax.set_xticks(to_x(P_TICKS))
        ax.set_xticklabels([f"{p:g}" for p in P_TICKS], fontsize=8.5)
        ax.set_xlim(to_x(0.005), to_x(99.995))
        ax.grid(True, color="#999", lw=0.5, alpha=0.7)
        ax.set_xlabel("Tần suất P (%)")
        ax.set_ylabel("H (mm)")
        ax.set_title(f"ĐƯỜNG TẦN SUẤT LƯỢNG MƯA 1 NGÀY LỚN NHẤT NĂM - {self.v['project'].get().upper()}", fontsize=11, fontweight="bold")
        info = f"X̄ = {mean:.2f} mm   Cv = {cv:.3f}   Cs = {cs:.3f}" + (f"   n = {n_obs}" if n_obs else "")
        ax.text(0.98, 0.95, info, transform=ax.transAxes, ha="right", va="top", fontsize=9,
                bbox=dict(boxstyle="square,pad=0.5", fc="white", ec="black"))
        ax.legend(loc="lower left", fontsize=8.5)
        sig = "   |   ".join(f"{a}: {self.v[k].get()}" for a, k in (("Người vẽ", "drawer"), ("Người kiểm tra", "checker"), ("Năm", "year"))
                             if self.v[k].get().strip())
        if sig:
            fig.text(0.99, 0.01, sig, ha="right", fontsize=8.5)
        fig.tight_layout(rect=(0, 0.03, 1, 1))
        canvas = FigureCanvasTkAgg(fig, master=w)
        canvas.draw()
        NavigationToolbar2Tk(canvas, w).update()
        canvas.get_tk_widget().pack(fill="both", expand=True, padx=6, pady=4)

    # ----------------------------------------------------------- copy / excel
    def copy_results(self):
        res = self._ensure_results()
        if not res:
            return
        out = ["=== II. KẾT QUẢ HIỆN TRẠNG ===", "P\tH24p (mm)\tψ(T)\tHtp (mm)\tH'tp (mm)\tQmax,p (m³/s)"]
        out += [f"{c['label']}\t{c['h24']:.2f}\t{c['psi']:.4f}\t{c['htp']:.2f}\t{c['hred']:.2f}\t{c['q']:.1f}" for c in res["hien_trang"]]
        out += ["", "=== III. KẾT QUẢ XÉT BĐKH ===", "P\tH24p gốc\tK bđkh\tH24p bđkh\tψ(T)\tHtp (mm)\tH'tp (mm)\tQmax,p (m³/s)"]
        out += [f"{c['label']}\t{c['h24_goc']:.2f}\t{c['cc']:.3f}\t{c['h24']:.2f}\t{c['psi']:.4f}\t{c['htp']:.2f}\t{c['hred']:.2f}\t{c['q']:.1f}"
                for c in res["bdkh"]]
        self.root.clipboard_clear()
        self.root.clipboard_append("\n".join(out))
        self.show_status("✔ Đã sao chép kết quả vào Clipboard.", "ok")

    def export_excel(self):
        res = self._ensure_results()
        if not res:
            return
        path = filedialog.asksaveasfilename(initialfile=f"Sokolovsky_{_safe_name(res['project'])}.xlsx",
                                            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            export_workbook(path, res)
            messagebox.showinfo("Thành công", f"Đã xuất file Excel (có công thức sống):\n{path}")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất Excel: {e}")

    # -------------------------------------------------------- save / open
    KEYS = ["Qng", "project", "Flv", "Ls", "kt", "b210", "alpha", "Ho", "c13", "f", "vmax_sel", "Vmax", "kv_sel", "kv",
            "km_sel", "km", "vung", "fa", "fl", "fr", "delta", "mean_x", "cv", "cs", "cs_mode", "pp",
            "p_list", "cc_all", "log_time", "auto_calc", "kp_legacy", "drawer", "checker", "year"]

    def save_project(self):
        path = filedialog.asksaveasfilename(defaultextension=".soko", filetypes=[("Dự án Sokolovsky", "*.soko"), ("JSON", "*.json")])
        if not path:
            return
        try:
            data = {"version": 2, "vars": {k: self.v[k].get() for k in self.KEYS},
                    "rows": [[r["p"], r["h"].get(), r["cc"].get()] for r in self.rain_rows],
                    "sample": None if self.sample is None else self.sample.tolist()}
            with open(path, "w", encoding="utf-8") as fh:
                json.dump(data, fh, ensure_ascii=False, indent=2)
            self.show_status(f"✔ Đã lưu dự án: {path}", "ok")
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không lưu được: {e}")

    def open_project(self):
        path = filedialog.askopenfilename(filetypes=[("Dự án Sokolovsky", "*.soko *.json"), ("Tất cả", "*.*")])
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as fh:
                data = json.load(fh)
            if data.get("version") == 2:
                vars_, rows, sample = data["vars"], data["rows"], data.get("sample")
            else:  # định dạng .soko của bản gốc
                vars_ = {("project" if k == "project_name" else k): v for k, v in data.items()
                         if k not in ("rain_data", "cc_data")}
                vars_["mean_x"] = str(data.get("mean_x", DEFAULTS["mean_x"]))
                rd, cd = data.get("rain_data", {}), data.get("cc_data", {})
                rows = [[float(re.sub(r"[^0-9.]", "", lab)), rd[lab], cd.get(lab, "1.0")] for lab in rd]
                sample = None
                vars_["vung"] = lab_zone(int(parse_key(str(vars_.get("vung", "Vùng 13")).replace("Vùng ", ""))))
                b = str(vars_.get("b210", "")).replace("Vùng ", "")
                if b.isdigit(): vars_["b210"] = lab_b210(int(b))
                c = str(vars_.get("c13", "")).replace("Vùng ", "")
                if c in BANG_C13: vars_["c13"] = lab_c13(c)
                vars_.setdefault("km_sel", KP_MANUAL)
                if vars_.get("km_sel") not in [KP_MANUAL] + list(BANG_KM): vars_["km_sel"] = KP_MANUAL
            vars_.setdefault("Qng", "0.0")
            self._internal = True
            try:
                for k, val in vars_.items():
                    if k in self.v:
                        self.v[k].set(val)
            finally:
                self._internal = False
            self.sample = None if sample is None else np.array(sample, float)
            self.lbl_sample.config(text="  (chưa nạp chuỗi số liệu)" if self.sample is None
                                   else f"  Đã nạp n = {len(self.sample)} năm (từ dự án)")
            self.apply_p_list([float(r[0]) for r in rows], [r[1] for r in rows], [r[2] for r in rows])
            self.calculate()
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không mở được dự án: {e}")

    def on_close(self):
        if self._after:
            self.root.after_cancel(self._after)
        self.root.destroy()


def main():
    if not HAS_TK:
        raise SystemExit("Python này chưa có tkinter (Linux: sudo apt install python3-tk).")
    root = tk.Tk()
    SokolovskyApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
